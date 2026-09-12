# -*- coding: utf-8 -*-
"""Génère les livrables DOCX PFA LexHub (Douae EL ASSAL / WEBVOICE)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
ARABIC_FONTS = ("Traditional Arabic", "Segoe UI", "Arial", "Noto Naskh Arabic")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def set_run_font(run, name: str = "Calibri", size: int = 11, bold: bool = False, italic: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text: str,
    *,
    style: str | None = None,
    bold: bool = False,
    italic: bool = False,
    size: int = 11,
    align=None,
    space_after: float = 6,
    space_before: float = 0,
    font: str = "Calibri",
):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, name=font, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return p


def add_bullets(doc, items: list[str], *, bold_prefix: bool = False):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix and " :" in item:
            left, right = item.split(" :", 1)
            r1 = p.add_run(left + " :")
            set_run_font(r1, bold=True, size=11)
            r2 = p.add_run(right)
            set_run_font(r2, size=11)
        else:
            r = p.add_run(item)
            set_run_font(r, size=11)


def add_numbered(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=11)


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_rtl_paragraph(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_arabic_para(doc, text: str, *, size: int = 13, space_after: float = 8):
    p = doc.add_paragraph()
    set_rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    # Prefer Traditional Arabic / Segoe UI for Windows
    font_name = ARABIC_FONTS[0]
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:cs"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run.font.size = Pt(size)
    # Complex script size
    rPr = run._element.get_or_add_rPr()
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(size * 2))
    rPr.append(szCs)
    return p


def make_table(doc, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, bold=True, size=10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(hdr[i], "1A365D")
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=9)
            if ri % 2 == 1:
                set_cell_shading(cells[ci], "F0F4F8")
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def setup_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("LexHub — PFA Douae EL ASSAL — WEBVOICE")
    set_run_font(hr, size=8, italic=True)
    hr.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Document confidentiel — usage encadrement PFA / démonstration")
    set_run_font(fr, size=8, italic=True)

    add_para(doc, "LexHub / Avocat AI Assistance", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Stage PFA — WEBVOICE — Encadrant entreprise : Hatim Benhaddou", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Stagiaire : Douae EL ASSAL", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_para(doc, title, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    if subtitle:
        add_para(doc, subtitle, italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    return doc


# ===========================================================================
# FILE 1 — Liste diagrammes UML
# ===========================================================================

def build_file1() -> Path:
    doc = setup_doc(
        "Liste exhaustive des diagrammes UML nécessaires au rapport PFA LexHub",
        "Cartographie UML 2.x — priorisation et suggestions de contenu pour LexHub",
    )

    add_heading_custom(doc, "1. Objet du document", 1)
    add_para(
        doc,
        "Ce document dresse la liste des types de diagrammes UML pertinents pour le rapport "
        "de Projet de Fin d'Année (PFA) LexHub. Pour chaque type : définition, utilité générale, "
        "mapping concret LexHub, priorité pour le rapport, et suggestion de contenu.",
    )
    add_para(
        doc,
        "Référentiel : UML 2.x (OMG). LexHub = SaaS cabinet d'avocats (Next.js, FastAPI, Postgres, "
        "Redis, MinIO/S3, STT cloud, déploiement VPS lexhub.space).",
    )

    add_heading_custom(doc, "2. Tableau de synthèse des types UML", 1)

    uml_rows = [
        [
            "Cas d'utilisation\n(Use Case)",
            "Capturer les fonctionnalités vues des acteurs et les frontières du système.",
            "Acteurs : avocat, assistant, admin cabinet, admin plateforme. UC : inscription, dictée, validation, assignation, export.",
            "Obligatoire",
            "UC global + zoom atelier plaidoirie",
        ],
        [
            "Classes\n(Class)",
            "Modéliser entités, attributs, associations et cardinalités.",
            "Case, Document, PleaDraft, User, Cabinet, Notification, Invite, ActivityArchive.",
            "Obligatoire",
            "Modèle de données métier + liens tenant",
        ],
        [
            "Séquence\n(Sequence)",
            "Ordonner les messages entre objets/lifelines dans le temps.",
            "OTP, dictée STT+fusḥā, validation/assignation, join cabinet, OCR document.",
            "Obligatoire",
            "3–5 séquences cœur métier",
        ],
        [
            "Activité\n(Activity)",
            "Flux de contrôle / workflow (décisions, parallèles, fin).",
            "Pipeline plaidoirie 5 étapes ; pipeline document (upload→OCR→synthèse).",
            "Obligatoire",
            "1 activité plaidoirie + 1 OCR",
        ],
        [
            "États\n(State Machine)",
            "Cycle de vie d'un objet via états et transitions.",
            "Statuts PleaDraft : brouillon → en revue → validée → assignée (évent. archivée).",
            "Fortement recommandé",
            "Machine d'états plaidoirie",
        ],
        [
            "Composants\n(Component)",
            "Architecture logicielle : modules remplaçables et interfaces.",
            "Frontend Next.js, Backend FastAPI, Postgres, Redis, MinIO/S3, STT, LLM, Caddy.",
            "Fortement recommandé",
            "Vue composants + dépendances",
        ],
        [
            "Déploiement\n(Deployment)",
            "Nœuds physiques/VM, artéfacts et connexions réseau.",
            "VPS 51.x / lexhub.space, Docker Compose, Caddy TLS, volumes.",
            "Obligatoire",
            "Prod lexhub.space",
        ],
        [
            "Paquetages\n(Package)",
            "Organisation logique du code / namespaces.",
            "api/, services/, models/, schemas/ ; frontend app/(routes).",
            "Fortement recommandé",
            "Packages backend Python",
        ],
        [
            "Communication\n(Collaboration)",
            "Interactions centrées sur les liens structurels (pas le temps).",
            "Avocat ↔ Backend ↔ STT ↔ Redis pub/sub (notifications).",
            "Optionnel",
            "Collaboration notifications temps réel",
        ],
        [
            "Timing\n(Timing)",
            "Contraintes temporelles sur états / messages.",
            "TTL OTP e-mail ; timeout chunk audio STT ; debounce WebSocket.",
            "Optionnel",
            "Uniquement si focus non-fonc. temps",
        ],
        [
            "Objets / Instances\n(Object)",
            "Instantané d'objets concrets et liens à un moment T.",
            "Exemple : Cabinet « Cabinet Rabat » + 2 User + 1 PleaDraft validée.",
            "Optionnel",
            "Illustration multi-tenant",
        ],
        [
            "Structure composite\n(Composite Structure)",
            "Parties internes d'un classifieur et ports/connectors.",
            "Peu critique pour LexHub SaaS classique (pas de ports hardware).",
            "Optionnel",
            "Mention brève / hors scope rapport",
        ],
        [
            "Vue d'ensemble d'interaction\n(Interaction Overview)",
            "Activité dont les nœuds sont des interactions (séquence).",
            "Chaîner OTP → login → dictée → validation → export.",
            "Optionnel",
            "Si besoin de « big picture »",
        ],
        [
            "Profil\n(Profile)",
            "Extension UML (stéréotypes, tagged values) pour un domaine.",
            "Avancé (stéréotypes « SaaS Tenant », « STT Service ») — hors besoin PFA.",
            "Optionnel (avancé)",
            "Ne pas produire sauf annexe",
        ],
    ]

    make_table(
        doc,
        ["Nom UML (FR / EN)", "À quoi ça sert", "Pourquoi utile pour LexHub", "Priorité", "Suggestion contenu"],
        uml_rows,
        col_widths=[3.2, 3.5, 4.0, 2.2, 3.0],
    )

    # Detailed sections
    add_heading_custom(doc, "3. Fiches détaillées par type", 1)

    details = [
        (
            "3.1 Cas d'utilisation (Use Case Diagram)",
            "Représente les acteurs externes et les fonctionnalités du système sans entrer dans l'implémentation.",
            [
                "Acteurs : Avocat, Assistant (stagiaire), Admin cabinet, Admin plateforme (WebVoice).",
                "Système : LexHub.",
                "UC prioritaires : S'inscrire (OTP), Se connecter, Créer dossier, Uploader document, "
                "Dicter plaidoirie (AR), Corriger fusḥā, Détecter contradictions, Générer structure, "
                "Valider plaidoirie, Assigner à assistant (PIN), Exporter DOCX/PDF, Gérer équipe, "
                "Configurer branding, Suspendre cabinet (admin plateforme).",
            ],
            "Obligatoire — déjà présent (TikZ + PlantUML uc-lexhub.puml).",
        ),
        (
            "3.2 Classes (Class Diagram)",
            "Décrit le modèle conceptuel / logique des données et leurs relations.",
            [
                "Classes clés : User, Cabinet, Case, Document, PleaDraft, Notification, CabinetInvite, "
                "ActivityArchive, JoinCode.",
                "Relations : User ∈ Cabinet ; Case → Document* ; Case → PleaDraft* ; PleaDraft.status ; "
                "Notification → User.",
                "Cardinalités multi-tenant (isolation par cabinet_id).",
            ],
            "Obligatoire — déjà présent (class-lexhub.puml + tableau LaTeX).",
        ),
        (
            "3.3 Séquence (Sequence Diagram)",
            "Montre l'ordre chronologique des échanges entre participants.",
            [
                "Inscription OTP (email → code → cabinet ACTIVE).",
                "Dictée arabe : micro → frontend → API → STT → correction fusḥā → marques orange.",
                "Validation plaidoirie → assignation assistant (+ PIN si applicable).",
                "Rejoindre cabinet (code / invite).",
                "Optionnel : upload document + OCR / synthèse.",
            ],
            "Obligatoire — séquences dictée, assign, OTP déjà dans le rapport / PlantUML.",
        ),
        (
            "3.4 Activité (Activity Diagram)",
            "Formalise un processus métier avec décisions et étapes.",
            [
                "Pipeline plaidoirie 5 étapes : Contexte → Transcription → Correction forme → "
                "Contradictions → Génération/structure → Validation.",
                "Pipeline document : Upload → sniff MIME → OCR cloud → extraction → synthèse / ask.",
            ],
            "Obligatoire pour plaidoirie (activity-plea.puml) ; OCR à produire/compléter.",
        ),
        (
            "3.5 Machine à états (State Machine)",
            "Cycle de vie d'une instance (ici PleaDraft / dossier).",
            [
                "États : brouillon, en_revue, validée, assignée, (archivée).",
                "Transitions : avocat édite ; avocat valide ; avocat assigne ; assistant consulte pièce validée.",
                "Utile pour expliquer les droits UI selon le statut.",
            ],
            "Fortement recommandé — à produire (pas encore central dans LaTeX).",
        ),
        (
            "3.6 Composants (Component Diagram)",
            "Découpe logicielle et interfaces fournies/requises.",
            [
                "Composants : frontend (Next.js), backend (FastAPI), db (Postgres), cache/broker (Redis), "
                "object storage (MinIO/S3), services STT/LLM externes, reverse-proxy Caddy.",
                "Interfaces : REST/JSON, WebSocket realtime, SMTP OTP.",
            ],
            "Fortement recommandé — partiellement couvert par « architecture » ; formaliser en UML composants.",
        ),
        (
            "3.7 Déploiement (Deployment Diagram)",
            "Cartographie des nœuds d'exécution et artéfacts déployés.",
            [
                "Nœud VPS, conteneurs Docker Compose (frontend, backend, postgres, redis, minio), "
                "Caddy (TLS lexhub.space), volumes persistants, DNS.",
            ],
            "Obligatoire — deploy-lexhub.puml déjà présent.",
        ),
        (
            "3.8 Paquetages (Package Diagram)",
            "Organisation des modules et dépendances entre packages.",
            [
                "Backend : app.api, app.services, app.models, app.schemas, app.core, alembic.",
                "Frontend : app routes (dashboard, cases, documents, plea, team, settings).",
            ],
            "Fortement recommandé — à produire (vue code).",
        ),
        (
            "3.9 Communication / Collaboration",
            "Même information qu'une séquence mais centrée sur les liens structurels.",
            [
                "Peu de valeur ajoutée si les séquences sont déjà claires.",
                "Éventuel focus : graphe avocat–API–Redis–clients pour notifications.",
            ],
            "Optionnel.",
        ),
        (
            "3.10 Timing",
            "Axes de temps et contraintes (durée, période).",
            [
                "OTP expiration, rate-limit, latence STT chunks — utile seulement en annexe non-fonctionnelle.",
            ],
            "Optionnel (si pertinent).",
        ),
        (
            "3.11 Objets / Instances",
            "Exemple concret d'instances liées.",
            [
                "Scénario démo : cabinet demo + avocat@sanad.ma + 1 dossier divorce + 1 PleaDraft validée.",
            ],
            "Optionnel.",
        ),
        (
            "3.12 Structure composite",
            "Détail interne d'un composant (parts, ports).",
            [
                "Peu adapté au niveau d'abstraction du rapport PFA LexHub (architecture SaaS standard).",
            ],
            "Optionnel — non prioritaire.",
        ),
        (
            "3.13 Vue d'ensemble d'interaction (Interaction Overview)",
            "Combine activité et fragments d'interaction.",
            [
                "Peut servir de « carte » reliant OTP → dictée → validation → export.",
            ],
            "Optionnel.",
        ),
        (
            "3.14 Profil UML (Profile)",
            "Mécanisme avancé d'extension du métamodèle UML.",
            [
                "Hors scope du rapport PFA sauf mention théorique en état de l'art UML.",
            ],
            "Optionnel (avancé) — mention brève uniquement.",
        ),
    ]

    for title, purpose, bullets, prio in details:
        add_heading_custom(doc, title, 2)
        add_para(doc, purpose)
        add_bullets(doc, bullets)
        add_para(doc, f"Priorité rapport : {prio}", bold=True, size=10, space_before=4)

    add_heading_custom(doc, "4. Diagrammes déjà présents vs à produire", 1)

    add_heading_custom(doc, "4.1 Déjà présents dans le rapport LaTeX / PlantUML", 2)
    make_table(
        doc,
        ["Diagramme", "Source", "Statut"],
        [
            ["Cas d'utilisation (stick figures TikZ)", "chapitres/04-analyse.tex + uc-lexhub.puml", "Présent"],
            ["Séquence dictée (voix → texte corrigé)", "05-conception.tex + seq-dictation.puml", "Présent"],
            ["Séquence validation / assignation", "05-conception.tex + seq-assign.puml", "Présent"],
            ["Séquence inscription OTP", "05-conception.tex + seq-otp.puml", "Présent (version locale)"],
            ["Classes / modèle de données", "05-conception.tex + class-lexhub.puml", "Présent"],
            ["Activité atelier plaidoirie", "05-conception.tex + activity-plea.puml", "Présent"],
            ["Déploiement lexhub.space", "05-conception.tex + deploy-lexhub.puml", "Présent"],
            ["Architecture (vue composants informelle)", "Chapitre conception / réalisation", "Présent (à formaliser UML)"],
        ],
        col_widths=[5.5, 7.0, 3.0],
    )

    add_heading_custom(doc, "4.2 À produire ou compléter", 2)
    make_table(
        doc,
        ["Diagramme recommandé", "Priorité", "Action"],
        [
            ["Machine à états PleaDraft", "Fortement recommandé", "Créer state-plea.puml + figure LaTeX"],
            ["Composants UML formels", "Fortement recommandé", "Formaliser architecture en Component"],
            ["Paquetages backend", "Fortement recommandé", "package-backend.puml"],
            ["Activité pipeline OCR documents", "Fortement recommandé", "activity-ocr.puml"],
            ["Séquence join cabinet / invite", "Recommandé", "seq-join.puml"],
            ["UC zoom admin plateforme", "Optionnel", "Compléter acteurs WebVoice"],
            ["Communication / Timing / Objects", "Optionnel", "Uniquement si place en annexe"],
            ["Composite Structure / Profile", "Optionnel", "Ne pas produire (mention théorique)"],
        ],
        col_widths=[5.5, 3.5, 6.5],
    )

    add_heading_custom(doc, "5. Recommandation concrète pour le rapport", 1)
    add_para(doc, "Jeu minimal « rapport PFA solide » :")
    add_numbered(
        doc,
        [
            "1× Cas d'utilisation (acteurs complets).",
            "1× Classes (domaine).",
            "3× Séquences (OTP, dictée, validation/assign).",
            "1× Activité (pipeline plaidoirie).",
            "1× Déploiement (prod).",
            "1× États (PleaDraft) — fort impact pédagogique.",
            "1× Composants ou Paquetages — pour lier conception ↔ code.",
        ],
    )
    add_para(
        doc,
        "Sources PlantUML déjà disponibles : docs/rapport-pfa/diagrams/*.puml "
        "(uc, class, seq-dictation, seq-assign, seq-otp, activity-plea, deploy-lexhub).",
        italic=True,
        size=10,
    )

    add_heading_custom(doc, "6. Conclusion", 1)
    add_para(
        doc,
        "Le rapport LexHub couvre déjà le cœur UML (UC, classes, séquences, activité, déploiement). "
        "Les ajouts à plus forte valeur sont la machine à états des plaidoiries, un diagramme de "
        "composants/paquetages, et l'activité OCR. Les types Timing, Composite Structure et Profile "
        "restent optionnels et non indispensables au PFA.",
    )

    path = OUT_DIR / "01_Liste_diagrammes_UML_LexHub.docx"
    doc.save(path)
    return path


# ===========================================================================
# FILE 2 — Guide utilisation
# ===========================================================================

def build_file2() -> Path:
    doc = setup_doc(
        "Guide d'utilisation LexHub",
        "Destinataire : encadrant entreprise (Hatim Benhaddou / WEBVOICE) — Démo https://lexhub.space — Version août 2026",
    )

    add_heading_custom(doc, "1. Présentation LexHub", 1)
    add_para(
        doc,
        "LexHub (https://lexhub.space) est une plateforme SaaS destinée aux cabinets d'avocats marocains. "
        "Elle centralise dossiers, documents, planning, équipe et — cœur du produit — un atelier de "
        "rédaction de plaidoirie par dictée vocale en arabe, avec correction de forme (fusḥā), "
        "détection de contradictions, génération de structure, validation avocat et assignation à l'assistant.",
    )
    add_para(
        doc,
        "Interface web responsive, authentification e-mail (OTP), isolation par cabinet, exports Word/PDF. "
        "Le présent guide décrit le parcours démo pour l'encadrant entreprise.",
    )

    add_heading_custom(doc, "2. Accès et comptes de test", 1)
    add_para(doc, "URL de production : https://lexhub.space", bold=True)
    add_para(
        doc,
        "Avertissement : mots de passe réservés aux tests PFA / démonstration uniquement.",
        italic=True,
        size=10,
    )
    make_table(
        doc,
        ["E-mail", "Mot de passe", "Rôle"],
        [
            ["demo@sanad.ma", "Sanad2026!", "Admin cabinet (coordonnateur)"],
            ["avocat@sanad.ma", "Sanad2026!", "Avocat (membre)"],
            ["assistant@sanad.ma", "Sanad2026!", "Assistant / stagiaire"],
            ["admin@webvoice.ma", "WebVoice2026!", "Admin plateforme WebVoice"],
        ],
        col_widths=[5.5, 4.0, 5.5],
    )

    add_heading_custom(doc, "3. Différenciation des rôles (important)", 1)
    add_para(
        doc,
        "Les permissions dépendent uniquement du rôle attribué. "
        "Un avocat invité par e-mail a le rôle avocat : il travaille, mais ne gère pas l'équipe ni l'abonnement.",
    )
    make_table(
        doc,
        ["Capacité", "Avocat (invité)", "Admin cabinet", "Assistant"],
        [
            ["Dossiers, documents, planning", "Oui", "Oui", "Oui (lecture + upload docs)"],
            ["Dictée / créer plaidoirie", "Oui", "Oui", "Non"],
            ["Valider / assigner plaidoirie", "Oui", "Oui", "Non (voit les validées)"],
            ["Synthèse IA / questions doc", "Oui", "Oui", "Non"],
            ["Télécharger / imprimer docs", "Oui", "Oui", "Oui"],
            ["Inviter membres (e-mail)", "Non", "Oui", "Non"],
            ["Abonnement / tarifs cabinet", "Voir", "Gérer", "Non"],
            ["Branding cabinet", "Non", "Oui", "Non"],
            ["2FA / zone sensible paramètres", "Oui", "Oui", "Masqué"],
        ],
        col_widths=[4.5, 3.0, 3.5, 4.0],
    )

    add_heading_custom(doc, "4. Inscription cabinet (OTP + essai 14 j)", 1)
    add_numbered(
        doc,
        [
            "https://lexhub.space → Créer un bureau (inscription).",
            "Étape e-mail → code OTP → infos cabinet → profil avocat → mot de passe.",
            "Après validation : cabinet ACTIVE, essai plan Cabinet (14 jours), e-mail de bienvenue.",
            "Connexion immédiate au tableau de bord ; bannière d'essai + page Paramètres → Abonnement.",
        ],
    )

    add_heading_custom(doc, "5. Connexion", 1)
    add_numbered(
        doc,
        [
            "Connexion avec e-mail + mot de passe (ex. avocat@sanad.ma).",
            "Vérifier le menu selon le rôle (avocat / assistant / admin cabinet / admin plateforme).",
        ],
    )

    add_heading_custom(doc, "6. Inviter un collègue (admin cabinet uniquement)", 1)
    add_para(
        doc,
        "Il n'y a plus de code d'équipe public. L'admin cabinet invite exclusivement par e-mail "
        "(Équipe → Inviter). Le destinataire reçoit un lien sécurisé /join/accept?code=…, "
        "choisit son mot de passe, et son compte est actif immédiatement (rôle avocat ou assistant).",
    )
    add_numbered(
        doc,
        [
            "Admin → Équipe → Inviter par e-mail (nom, adresse, rôle).",
            "Le collègue ouvre le lien reçu, définit un mot de passe, se connecte.",
            "E-mails : invitation au membre + confirmation « bienvenue » + notification à l'admin "
            "quand l'invitation est acceptée.",
            "Les quotas sièges (avocats / assistants) dépendent du plan d'abonnement.",
        ],
    )
    add_para(
        doc,
        "Ancien parcours « code join + demande d'approbation » : désactivé côté produit. "
        "Seules les invitations e-mail admin sont valides.",
        italic=True,
        size=10,
    )

    add_heading_custom(doc, "7. Tableau de bord", 1)
    add_heading_custom(doc, "7.1 Avocat / admin cabinet", 2)
    add_bullets(
        doc,
        [
            "Compteurs dossiers actifs, documents, événements à venir.",
            "Raccourcis : Nouveau dossier, Nouvelle plaidoirie, Documents.",
            "Événements dans les 48 h : visibles ; rappel automatique (notif + e-mail) à toute l'équipe.",
        ],
    )
    add_heading_custom(doc, "7.2 Assistant / stagiaire", 2)
    add_para(
        doc,
        "Espace dédié (pas un clone avocat) : plaidoiries assignées, plaidoiries validées, "
        "événements 48 h. Pas de bouton « Nouveau dossier » ni micro. Pas de widget « Tâches ouvertes ».",
    )

    add_heading_custom(doc, "8. Dossiers", 1)
    add_numbered(
        doc,
        [
            "Créer un dossier (référence auto WV-AAAA-…), domaine, responsable, statut.",
            "À la suppression : aperçu du contenu (client, docs, tâches) + confirmation « supprimer ».",
            "Les documents / notes / agenda / plaidoiries liés sont purgés avec le dossier.",
        ],
    )

    add_heading_custom(doc, "9. Documents", 1)
    add_heading_custom(doc, "9.1 Avocat", 2)
    add_numbered(
        doc,
        [
            "Upload PDF / Word / image ; OCR / extraction ; synthèse IA ; questions sur le document.",
            "Télécharger, imprimer, supprimer.",
        ],
    )
    add_heading_custom(doc, "9.2 Assistant", 2)
    add_bullets(
        doc,
        [
            "Peut uploader et voir tous les documents du cabinet (avocat ou assistant).",
            "Peut ouvrir, prévisualiser (PDF/image), télécharger, imprimer, supprimer.",
            "Pas d'aperçu « texte extrait », pas de synthèse IA, pas de questions IA.",
        ],
    )

    add_heading_custom(doc, "10. Plaidoirie — parcours avocat", 1)
    add_para(
        doc,
        "Dictée micro arabe uniquement (pas de choix FR/AR, pas d'upload fichier audio).",
        bold=True,
    )
    add_numbered(
        doc,
        [
            "Contexte → Transcription micro → Correction fusḥā + marques orange STT → "
            "Contradictions → Structure → Validation → Assignation assistant.",
            "Export Word / PDF (aperçu document officiel).",
            "Plaidoirie sans dossier : possible. Ensuite, bouton « Lier à un dossier » : "
            "sans dossier / dossier existant / créer un nouveau dossier.",
        ],
    )

    add_heading_custom(doc, "11. Planning et rappels 48 h", 1)
    add_para(
        doc,
        "Créer audiences, délais, RDV. Environ 48 h avant un événement, LexHub envoie "
        "une notification in-app et un e-mail à tous les membres actifs du cabinet "
        "(avocats, assistants, coordonnateur).",
    )

    add_heading_custom(doc, "12. Équipe (admin cabinet uniquement)", 1)
    add_bullets(
        doc,
        [
            "Inviter par e-mail uniquement (pas de code public).",
            "Voir / modifier membres, révoquer un accès si besoin.",
            "Un avocat invité : voit l'équipe en lecture, sans actions de gestion.",
            "Respect des quotas sièges du plan (Solo / Cabinet / Entreprise).",
        ],
    )

    add_heading_custom(doc, "13. Paramètres & abonnement", 1)
    add_bullets(
        doc,
        [
            "Admin cabinet : branding (logo, nom), digest e-mail, densités, section Abonnement.",
            "Avocat : profil, nom arabe (signature plaidoiries), 2FA, mot de passe, zone sensible compte ; "
            "peut consulter l'abonnement (upgrade via admin).",
            "Assistant : profil + changement de mot de passe uniquement.",
            "Page publique /tarifs + Paramètres → Abonnement : plans Découverte / Solo / Cabinet / Entreprise.",
        ],
    )

    add_heading_custom(doc, "13.1 Plans et essai (commercial)", 2)
    add_para(
        doc,
        "À l'inscription : essai 14 jours du plan Cabinet (recommandé). Ensuite, bascule automatique "
        "sur Découverte (gratuit, quotas réduits) sauf choix d'un plan payant. "
        "Paiement CMI / PayZone en préparation : sélection de plan active déjà l'accès (mode mock).",
    )
    make_table(
        doc,
        ["Plan", "Prix (MAD/mois)", "À retenir"],
        [
            ["Découverte", "0", "1 avocat, quotas bas — freemium permanent"],
            ["Solo", "299", "1 avocat + 1 assistant, volume plaidoiries"],
            ["Cabinet", "699", "Équipe, clé de barre IA, plan recommandé"],
            ["Entreprise", "1890", "Gros cabinets, quotas élevés"],
        ],
        col_widths=[3.5, 4.0, 7.5],
    )
    add_bullets(
        doc,
        [
            "E-mails de rétention : bienvenue, rappels J-7 / J-3 / J-1, fin d'essai, upgrade.",
            "Bannière in-app quand l'essai touche à sa fin ou le quota plaidoiries ≥ 80 %.",
            "IA : quotas free (Groq / routeurs) tant que la clé OpenAI WEBVOICE n'est pas fournie.",
        ],
    )

    add_heading_custom(doc, "14. Notifications", 1)
    add_para(
        doc,
        "Cloche : plaidoirie validée / assignée, rappels planning 48 h, activité équipe. "
        "Marquer comme lu. Complété par les e-mails lifecycle (catalogue séparé pour Hatim).",
    )

    add_heading_custom(doc, "15. Rôle Assistant — résumé", 1)
    add_bullets(
        doc,
        [
            "Dashboard stagiaire : assignées / validées / événements 48 h.",
            "Documents : upload + téléchargement + impression (sans IA).",
            "Plaidoiries : uniquement les validées (aperçu, Word, PDF, impression).",
            "Compte test : assistant@sanad.ma / Sanad2026!",
        ],
    )

    add_heading_custom(doc, "16. Admin plateforme (WebVoice)", 1)
    add_para(
        doc,
        "Compte admin@webvoice.ma : tableau de bord, liste cabinets (activer / suspendre), "
        "détail membres (code équipe masqué), journal d'audit.",
    )
    add_heading_custom(doc, "16.1 Journal d'audit — à quoi ça sert ?", 2)
    add_para(
        doc,
        "Ce n'est pas une surveillance du contenu des dossiers. C'est la traçabilité technique "
        "standard des applications SaaS (comme Google Workspace, Salesforce, banques) : "
        "qui s'est connecté, qui a exporté un fichier, qui a supprimé un document, avec date et IP. "
        "Aucun texte de plaidoirie ni contenu PDF n'y figure. Conservation indicatif : 5 ans.",
    )
    add_para(
        doc,
        "Au-delà de 90 jours, les entrées anciennes sont archivées en Excel. "
        "Dans Journal d'audit → section « Archives Excel » : bouton pour forcer l'archivage "
        "et télécharger les fichiers.",
    )

    add_heading_custom(doc, "17. Activité du cabinet (équipe)", 1)
    add_para(
        doc,
        "Fil visible à tout le cabinet (connexions, dictées, validations…). "
        "Au-delà de 90 jours : archive Excel automatique, téléchargeable depuis la page Activité.",
    )

    add_heading_custom(doc, "18. Conseils démo devant encadrant", 1)
    add_numbered(
        doc,
        [
            "Scénario unique (ex. divorce) : contexte → dictée → orange → contradictions → validation → vue assistant.",
            "Montrer le dashboard assistant (différent de l'avocat) et le téléchargement document.",
            "Montrer « Lier à un dossier » sur une plaidoirie libre.",
            "Mentionner rappels 48 h + journal d'audit (traçabilité, pas contenu).",
            "Durée cible : 8–12 minutes.",
        ],
    )

    add_heading_custom(doc, "19. Limitations / bonnes pratiques micro", 1)
    add_bullets(
        doc,
        [
            "Dictée = arabe oral clair ; bruit ou débit trop rapide dégrade le STT.",
            "Marques orange = incertitude : toujours relire avant validation.",
            "Chrome / Edge, permission micro, réseau stable.",
            "Démo = micro live uniquement (pas d'upload audio).",
        ],
    )

    add_heading_custom(doc, "Annexe — checklist démo rapide", 1)
    make_table(
        doc,
        ["#", "Action", "OK"],
        [
            ["1", "Connexion avocat@sanad.ma", "☐"],
            ["2", "Créer dossier / upload doc", "☐"],
            ["3", "Dictée micro arabe + fusḥā + orange", "☐"],
            ["4", "Contradictions + structure", "☐"],
            ["5", "Lier plaidoirie à un dossier", "☐"],
            ["6", "Validation + assignation", "☐"],
            ["7", "Export Word/PDF", "☐"],
            ["8", "Vue assistant (docs + plaidoiries)", "☐"],
            ["9", "Planning / rappel 48 h (mention)", "☐"],
            ["10", "Admin : journal d'audit + archives", "☐"],
        ],
        col_widths=[1.5, 12.0, 2.0],
    )

    path = OUT_DIR / "02_Guide_utilisation_LexHub.docx"
    doc.save(path)
    return path


# ===========================================================================
# FILE 3 — Scripts arabes + comptes
# ===========================================================================

SCRIPTS = [
    {
        "titre_fr": "Divorce et garde d'enfants (حضانة)",
        "domaine": "Droit de la famille",
        "contexte": (
            "L'avocat représente la mère dans une demande de divorce pour discorde et de fixation "
            "de la garde. Le père contestait la résidence habituelle des enfants. Script destiné "
            "à tester STT + mots-clés familiaux."
        ),
        "arabe": (
            "بسم الله الرحمن الرحيم، السيد الرئيس، السادة المستشارون، "
            "أتشرف بالدفاع عن موكلتي السيدة فاطمة بنعلي في ملف الطلاق للشقاق وطلب الحضانة. "
            "تزوج الطرفان بتاريخ الخامس عشر من مارس سنة ألفين وثمانية عشر، وأنجبا طفلين قاصرين: "
            "يوسف في السابعة من عمره، وسارة في الخامسة. "
            "وقد استحال استمرار العشرة الزوجية بسبب الخلاف المستمر وإهمال الزوج لواجباته المادية والمعنوية. "
            "إن موكلتي هي من يتولى يوميا تربية الأطفال ومتابعة دراستهم وعلاجهم، "
            "بينما يقيم المدعى عليه في مدينة أخرى ولا يزورهم إلا نادرا. "
            "وطبقا لمقتضيات مدونة الأسرة، فإن مصلحة المحضون هي المعيار الأساسي في إسناد الحضانة. "
            "لذلك نلتمس الحكم بتطليق موكلتي للشقاق، وإسناد حضانة الطفلين إليها، "
            "مع تحديد زيارة منظمة للأب لا تضر باستقرارهما المدرسي، "
            "والحكم بنفقة شهرية كافية لتغطية السكن والغذاء والتعليم، "
            "مع شمول الحكم بالنفاذ المعجل وتحميل المدعى عليه الصائر. "
            "حفظ الله عدالتكم والسلام عليكم."
        ),
        "mots_cles": "طلاق، شقاق، حضانة، مدونة الأسرة، نفقة، مصلحة المحضون، زيارة، طفلين",
        "contradiction": (
            "Variante avec contradiction : remplacer « يوسف في السابعة » par « يوسف في التاسعة » "
            "tout en gardant plus loin « الطفلان دون السادسة » — incohérence d'âge volontaire "
            "pour tester le détecteur."
        ),
    },
    {
        "titre_fr": "Contrat de travail / licenciement abusif (فصل تعسفي)",
        "domaine": "Droit du travail",
        "contexte": (
            "Salarié licencié sans procédure ni motif réel. Demande de dommages-intérêts et "
            "rappel de salaires. Domaine labor pour STT juridique social."
        ),
        "arabe": (
            "السيد الرئيس، حضرات المستشارين، "
            "أدافع عن السيد كريم العلوي الذي اشتغل لدى شركة النور للخدمات بعقد شغل غير محدد المدة "
            "منذ فاتح يناير سنة ألفين وعشرين، بصفته تقنيا متخصصا، وبأجر شهري صاف قدره ستة آلاف درهم. "
            "وفي يوم العاشر من شتنبر سنة ألفين وأربعة وعشرين، فوجئ بقرار الفصل الشفوي دون استفسار كتابي "
            "ودون احترام مسطرة التأديب المنصوص عليها في مدونة الشغل. "
            "لم يثبت المشغل أي خطأ جسيم، ولم يمكن الأجير من الدفاع عن نفسه، "
            "كما لم يؤد له مستحقاته من أجرة غشت وتعويض عن الفصل والعطلة السنوية. "
            "إن هذا الفصل يعد فصلا تعسفيا موجبا للتعويض. "
            "لذلك نلتمس الحكم بإفراغ طابعه التعسفي، وأداء أجرة الشهر غير المؤدى، "
            "والتعويض عن الإخطار، والتعويض عن الضرر المادي والمعنوي بما لا يقل عن أجرة ستة أشهر، "
            "مع تسليم شهادة العمل وكشف الحساب، وشمول الحكم بالنفاذ المعجل وتحميل المدعى عليها الصائر. "
            "والسلام عليكم ورحمة الله."
        ),
        "mots_cles": "مدونة الشغل، فصل تعسفي، أجير، أجرة، إخطار، شهادة العمل، خطأ جسيم",
        "contradiction": (
            "Variante avec contradiction : indiquer d'abord « أجر شهري ستة آلاف درهم » puis "
            "« أجره لم يتجاوز أربعة آلاف درهم طيلة مدة العمل » — contradiction de montant."
        ),
    },
    {
        "titre_fr": "Accident de circulation / indemnisation (تعويض)",
        "domaine": "Responsabilité civile / accidents",
        "contexte": (
            "Victime piétonne heurtée par un véhicule. Demande d'expertise et d'indemnisation "
            "corporelle. Test STT médical-juridique."
        ),
        "arabe": (
            "السيد الرئيس، السادة المستشارون المحترمون، "
            "يمثل هذا الملف السيدة نادية المرابط ضحية حادثة سير وقعت بتاريخ الثاني من ماي سنة ألفين وخمسة وعشرين "
            "على مستوى شارع محمد الخامس بمدينة فاس، حين دهسها سائق السيارة المسجلة تحت عدد أرقام وهمية للاختبار "
            "وهو في حالة عدم انتباه عند قطع الممر المخصص للراجلين. "
            "نقلت الضحية إلى المستشفى وأثبت التقرير الطبي كسرا في الساق اليمنى ورضوضا متعددة ومدة عجز مؤقت "
            "حددت في تسعين يوما، مع احتمال عجز جزئي دائم. "
            "تقع المسؤولية كاملة على السائق المؤمَّن لدى شركة التأمين المدعى عليها، "
            "طبقا لقواعد المسؤولية عن حوادث السير ومقتضيات ظهير الالتزامات والعقود فيما يخص الضرر. "
            "نلتمس الأمر بإجراء خبرة طبية حسابية لتقدير الأضرار البدنية والمادية، "
            "والحكم بأداء تعويض مسبق، ثم التعويض النهائي عن العجز والآلام وضرر فقدان الأجر، "
            "مع الفوائد القانونية وتحميل المدعى عليهم الصائر. "
            "حفظ الله مجلسكم الموقر."
        ),
        "mots_cles": "حادثة سير، راجل، عجز مؤقت، خبرة طبية، تعويض، تأمين، مسؤولية",
        "contradiction": (
            "Variante avec contradiction : dire « مدة العجز تسعين يوما » puis « عادت إلى عملها بعد أسبوع واحد دون انقطاع »."
        ),
    },
    {
        "titre_fr": "Bail commercial / expulsion du local (إخلاء محل)",
        "domaine": "Droit commercial / baux",
        "contexte": (
            "Bailleur commercial demande résiliation pour défaut de paiement de loyers et expulsion. "
            "Vocabulaire bail commercial marocain."
        ),
        "arabe": (
            "بسم الله، السيد الرئيس، "
            "أدافع عن السيد عبد الرحيم التازي مالك المحل التجاري الكائن بزنقة السوق القديم بالرباط، "
            "والمكترى للسيد هشام بنعلي بمقتضى عقد كراء تجاري مؤرخ في الأول من يوليوز سنة ألفين وتسعة عشر. "
            "التزم المكتري بأداء سومة كرائية شهرية قدرها أربعة آلاف وخمسمائة درهم، "
            "غير أنه توقف عن الأداء منذ شهر مارس سنة ألفين وخمسة وعشرين، رغم الإنذارات الموجهة إليه "
            "بصفة قانونية مع منحه أجل الأداء. "
            "وقد راكم دين كراء يفوق مجموع ستة أشهر، مما يبرر فسخ العقد والإفراغ طبقا للقواعد المنظمة "
            "للكراء التجاري وحماية حقوق المكري. "
            "لم يدل المكتري بأي مبرر مشروع ولا بعرض عيني حقيقي للدين. "
            "لذلك نلتمس الحكم بفسخ عقد الكراء، وأداء جميع الوجيبة الكرائية المتأخرة والضرائب، "
            "والحكم بإفراغ المحل هو ومن يقوم مقامه، تحت طائلة الغرامة التهديدية عن كل يوم تأخير، "
            "مع النفاذ المعجل وتحميل المدعى عليه الصائر. "
            "والسلام عليكم."
        ),
        "mots_cles": "كراء تجاري، سومة كرائية، إنذار، فسخ، إفراغ، وجيبة، محل تجاري",
        "contradiction": (
            "Variante avec contradiction : affirmer « توقف عن الأداء منذ مارس » puis "
            "« وقد أدى جميع الكراء إلى غاية غشت »."
        ),
    },
    {
        "titre_fr": "Vol / délit (جنحة سرقة)",
        "domaine": "Droit pénal (délit)",
        "contexte": (
            "Défense pénale légère : prévenu poursuivi pour vol simple. Plaidoirie axée sur "
            "absence de preuves suffisantes et circonstances atténuantes — test vocabulaire pénal."
        ),
        "arabe": (
            "السيد الرئيس، حضرات القضاة الأفاضل، "
            "أدافع عن المتهم ياسين الإدريسي المتابع من أجل جنحة السرقة طبقا للقانون الجنائي. "
            "تلخص الوقائع في ادعاء تعرض محل تجاري لكسر ليلي وسرقة بضائع قدرت قيمتها بخمسة آلاف درهم. "
            "غير أن محضر الضابطة القضائية لا يتضمن أي معاينة مباشرة لتلبس، "
            "ولا ضبطا للمتاع المسروق بحوزة موكلي، ولا اعترافا قضائيا صريحا، "
            "وإنما مجرد تصريحات متناقضة لشاهد واحد لم يحضر أثناء البحث التكميلي. "
            "إن قرينة البراءة أصل، والشك يفسر لصالح المتهم. "
            "موكلي شاب بدون سوابق قضائية، وله عمل قار، وحضر جميع جلسات التحقيق. "
            "نلتمس أساسا التصريح بالبراءة لعدم كفاية الدليل، "
            "واحتياطيا تمتيعه بظروف التخفيف في حال الإدانة مع إيقاف التنفيذ، "
            "مراعاة لسنه ولظروفه الاجتماعية. "
            "حفظ الله عدالتكم."
        ),
        "mots_cles": "جنحة سرقة، قانون جنائي، قرينة البراءة، محضر، تلبس، ظروف التخفيف، إيقاف التنفيذ",
        "contradiction": (
            "Variante avec contradiction : dire « بدون سوابق قضائية » puis « سبق إدانته من أجل السرقة سنة ألفين واثنين وعشرين »."
        ),
    },
    {
        "titre_fr": "Litige foncier / propriété (ملكية عقار)",
        "domaine": "Droit immobilier / foncier",
        "contexte": (
            "Action en revendication de propriété et annulation d'acte. Vocabulaire foncier "
            "(titre, conservation foncière, bornage)."
        ),
        "arabe": (
            "السيد الرئيس، السادة المستشارون، "
            "أترافع لفائدة السيد محمد الأموي طالب استحقاق العقار الفلاحي الكائن بجماعة عين الشقف عمالة مولاي يعقوب، "
            "ذي الرسم العقاري عدد اختبار للعرض، والمحددة مساحته بأربعة هكتارات تقريبا. "
            "يثبت أصل الملك لموكلي بمقتضى رسم ملكية قديم مدعم بشهادة شهود مستمرة الحيازة الهادئة العلنية "
            "منذ أكثر من أربعين سنة، وبتقييدات جزئية بالمحافظة العقارية. "
            "غير أن المدعى عليه استظهر عقد بيع حديث صادر عمن لا صفة له في التفويت، "
            "محاولا الاستيلاء على جزء من الأرض وحرثه دون إذن. "
            "إن الحيازة المستوفية للشروط القانونية مع السند الأقوى ترجح حق موكلي. "
            "نلتمس لذلك الحكم باستحقاق العقار موضوع النزاع لفائدته، "
            "وببطلان عقد البيع المدلى به فيما يخص الجزء المتنازع عليه، "
            "وبأمر بإجراء خبرة طبوغرافية لتحديد الحدود وإزالة التعدي، "
            "مع تعويض عن الحرمان من الاستغلال وتحميل المدعى عليه الصائر. "
            "والسلام عليكم ورحمة الله تعالى وبركاته."
        ),
        "mots_cles": "رسم عقاري، محافظة عقارية، حيازة، استحقاق، خبرة طبوغرافية، عقد بيع، ملكية",
        "contradiction": (
            "Variante avec contradiction : indiquer « أربعة هكتارات » puis « المساحة لا تتجاوز هكتارا واحدا حسب الخبرة السابقة »."
        ),
    },
]


def build_file3() -> Path:
    doc = setup_doc(
        "Scripts de plaidoirie arabes + comptes de test LexHub",
        "Livrable démo STT / détecteur de contradictions — https://lexhub.space",
    )

    # Part A
    add_heading_custom(doc, "Partie A — Comptes de test", 1)
    add_para(doc, "URL : https://lexhub.space", bold=True)
    add_para(
        doc,
        "Avertissement important : ces identifiants sont destinés exclusivement aux tests PFA "
        "et aux démonstrations. Ne pas les publier hors cadre académique / entreprise, ni les "
        "réutiliser comme mots de passe réels.",
        italic=True,
        size=10,
    )
    make_table(
        doc,
        ["E-mail", "Mot de passe", "Usage démo"],
        [
            ["demo@sanad.ma", "Sanad2026!", "Compte démo général"],
            ["avocat@sanad.ma", "Sanad2026!", "Parcours avocat (dictée, validation)"],
            ["assistant@sanad.ma", "Sanad2026!", "Vue assistant (pièces validées)"],
            ["admin@webvoice.ma", "WebVoice2026!", "Admin plateforme WebVoice"],
            ["test@cabinet-rabat.ma", "Test2026!", "Compte test (si encore valide)"],
        ],
        col_widths=[5.5, 4.0, 6.0],
    )

    # Protocole
    add_heading_custom(doc, "Protocole de test micro (STT arabe)", 1)
    add_numbered(
        doc,
        [
            "Casque micro ou micro intégré de qualité ; pièce calme.",
            "Navigateur Chrome/Edge ; autoriser le micro sur https://lexhub.space.",
            "Ouvrir l'atelier Plaidoirie → étape Transcription.",
            "Choisir un script ci-dessous ; lire à voix haute, débit modéré, articuler les termes juridiques.",
            "Éviter les digressions ; une pause courte entre les phrases longues.",
            "Arrêter l'enregistrement ; attendre la transcription.",
            "Vérifier les marques orange (incertitude STT) et corriger si besoin.",
            "Lancer la correction de forme fusḥā ; comparer au texte attendu.",
            "Optionnel : réutiliser la « variante avec contradiction » pour tester le détecteur.",
            "Contrôler les mots-clés listés sous chaque script.",
        ],
    )

    # Part B
    add_heading_custom(doc, "Partie B — Scripts de plaidoirie à lire (arabe)", 1)
    add_para(
        doc,
        "Six scripts distincts (~150–350 mots), style judiciaire clair (fusḥā / arabe de prétoire). "
        "Prêts à être dictés au micro. Pour chaque script : titre, contexte FR, texte arabe RTL, "
        "mots-clés attendus, variante contradiction optionnelle.",
    )

    for i, s in enumerate(SCRIPTS, start=1):
        add_heading_custom(doc, f"Script {i} — {s['titre_fr']}", 2)
        add_para(doc, f"Domaine : {s['domaine']}", bold=True, size=11)
        add_para(doc, "Contexte :", bold=True, size=10, space_after=2)
        add_para(doc, s["contexte"], size=10)
        add_para(doc, "Texte arabe à lire :", bold=True, size=10, space_after=4)
        add_arabic_para(doc, s["arabe"], size=13)
        add_para(doc, "Mots-clés attendus après transcription :", bold=True, size=10, space_after=2)
        add_arabic_para(doc, s["mots_cles"], size=12, space_after=6)
        add_para(doc, "Variante avec contradiction (test détecteur) :", bold=True, size=10, space_after=2)
        add_para(doc, s["contradiction"], size=10, italic=True)

    add_heading_custom(doc, "Annexe — grille de score rapide STT", 1)
    make_table(
        doc,
        ["Critère", "OK", "Partiel", "KO"],
        [
            ["Termes juridiques majeurs reconnus", "☐", "☐", "☐"],
            ["Noms propres / dates globalement corrects", "☐", "☐", "☐"],
            ["Marques orange pertinentes", "☐", "☐", "☐"],
            ["Correction fusḥā utile", "☐", "☐", "☐"],
            ["Contradiction volontaire détectée (si testée)", "☐", "☐", "☐"],
        ],
        col_widths=[8.0, 2.5, 2.5, 2.5],
    )

    path = OUT_DIR / "03_Plaidoiries_arabes_et_comptes_test.docx"
    doc.save(path)
    return path


# ===========================================================================
# FILE 4 — Guide accès base de données
# ===========================================================================

def build_file4() -> Path:
    doc = setup_doc(
        "Guide d'accès à la base de données LexHub",
        "Destinataire : encadrant entreprise (Hatim Benhaddou / WEBVOICE) — PostgreSQL 16 — Version août 2026",
    )

    add_heading_custom(doc, "1. Objet", 1)
    add_para(
        doc,
        "Ce document explique comment visualiser et ouvrir la base PostgreSQL de LexHub sur un poste "
        "personnel (Windows, macOS ou Linux). La base de production (https://lexhub.space) n'est pas "
        "exposée sur Internet pour des raisons de sécurité : la méthode recommandée est de lancer une "
        "copie locale identique via Docker, puis de s'y connecter avec DBeaver ou pgAdmin.",
    )

    add_heading_custom(doc, "2. Prérequis", 1)
    add_bullets(
        doc,
        [
            "Docker Desktop installé : https://www.docker.com/products/docker-desktop/",
            "Le dossier du projet LexHub (archive ZIP ou dépôt Git)",
            "Optionnel : DBeaver (https://dbeaver.io/download/) ou pgAdmin",
        ],
    )

    add_heading_custom(doc, "3. Démarrer la base en local (recommandé)", 1)
    add_para(doc, "Dans un terminal, à la racine du projet :", bold=True, size=10, space_after=4)
    add_para(doc, "docker compose up -d --build", size=10, space_after=6)
    add_para(
        doc,
        "Attendre 2 à 3 minutes que PostgreSQL démarre. Pour insérer des données de démonstration :",
        size=10,
        space_after=4,
    )
    add_para(
        doc,
        "docker compose exec backend python scripts/seed_test_accounts.py",
        size=10,
        space_after=6,
    )
    add_para(
        doc,
        "Comptes applicatifs créés : demo@sanad.ma / Sanad2026! (admin cabinet), avocat@sanad.ma, "
        "assistant@sanad.ma, admin@webvoice.ma / WebVoice2026! — voir aussi docs/COMPTES_ET_TESTS.md.",
        size=10,
    )

    add_heading_custom(doc, "4. Paramètres de connexion PostgreSQL", 1)
    make_table(
        doc,
        ["Paramètre", "Valeur"],
        [
            ["Hôte", "localhost"],
            ["Port", "15432 (port hôte Docker — pas 5432)"],
            ["Base de données", "sanad"],
            ["Utilisateur", "sanad"],
            ["Mot de passe", "sanad_secret"],
            ["URL complète", "postgresql://sanad:sanad_secret@localhost:15432/sanad"],
        ],
        col_widths=[5.0, 10.5],
    )

    add_heading_custom(doc, "5. Ouvrir avec DBeaver (interface graphique)", 1)
    add_numbered(
        doc,
        [
            "Installer DBeaver → Nouvelle connexion → PostgreSQL.",
            "Renseigner host localhost, port 15432, database sanad, user sanad, password sanad_secret.",
            "Tester la connexion → Finish.",
            "Explorer les tables (users, cabinets, cases, documents, agenda_events…), afficher les lignes, "
            "ou générer un diagramme ER (clic droit sur la base → View Diagram).",
        ],
    )

    add_heading_custom(doc, "6. Ligne de commande (sans interface)", 1)
    add_para(doc, "docker compose exec postgres psql -U sanad -d sanad", size=10, space_after=4)
    add_para(doc, "Commandes utiles dans psql :", bold=True, size=10, space_after=4)
    add_bullets(
        doc,
        [
            "\\dt — liste des tables",
            "\\d users — structure de la table users",
            "SELECT email, role FROM users LIMIT 10; — aperçu des données",
            "\\q — quitter",
        ],
    )

    add_heading_custom(doc, "7. Comprendre le schéma sans installer", 1)
    add_bullets(
        doc,
        [
            "Documentation texte : docs/03-base-de-donnees.md (toutes les tables expliquées)",
            "Diagramme MCD PlantUML : docs/diagrammes/47-mcd-entite-relation.puml",
            "Rendu image : coller le fichier sur https://www.plantuml.com/plantuml → télécharger le PNG",
        ],
    )
    add_para(
        doc,
        "Principe multi-tenant : chaque cabinet est isolé via cabinet_id. Les fichiers binaires (PDF, Word) "
        "sont stockés dans MinIO ; PostgreSQL conserve métadonnées, texte extrait et chemins storage_path.",
        size=10,
    )

    add_heading_custom(doc, "8. Base de production (lexhub.space)", 1)
    add_para(
        doc,
        "La base PostgreSQL du VPS n'est pas accessible directement depuis l'extérieur. Pour une revue "
        "technique en production, contacter la stagiaire (export pg_dump anonymisé ou accès SSH tunnel "
        "sur autorisation WebVoice). En pratique, la copie locale Docker suffit pour examiner la structure "
        "et des données de test représentatives.",
        size=10,
    )

    add_heading_custom(doc, "9. Tables principales (aperçu)", 1)
    make_table(
        doc,
        ["Table", "Rôle"],
        [
            ["cabinets", "Cabinet d'avocats (tenant)"],
            ["users", "Comptes (avocat, assistant, admin)"],
            ["cases", "Dossiers / affaires"],
            ["documents", "Métadonnées des pièces jointes"],
            ["agenda_events", "Planning lié aux dossiers"],
            ["audit_logs", "Journal d'audit (traçabilité)"],
        ],
        col_widths=[4.5, 11.0],
    )

    path = OUT_DIR / "04_Guide_acces_base_de_donnees_LexHub.docx"
    doc.save(path)
    return path


def build_file5_emails() -> Path:
    """Catalogue e-mails LexHub (rétention type Cursor / Notion / Clio) pour Hatim."""
    doc = setup_doc(
        "Catalogue e-mails LexHub",
        "Destinataire : Hatim Benhaddou / WEBVOICE — Version août 2026 — https://lexhub.space",
    )

    add_heading_custom(doc, "1. Objectif", 1)
    add_para(
        doc,
        "LexHub envoie des e-mails transactionnels et de rétention commerciale, sur le modèle des "
        "grandes apps SaaS (Cursor, Notion, Clio, Stripe) : onboarding, essai, upgrade, équipe, sécurité. "
        "L'IA reste sur quotas free (routeurs / Groq) tant que la clé OpenAI n'est pas fournie.",
    )

    add_heading_custom(doc, "2. E-mails déjà envoyés (implémentés)", 1)
    make_table(
        doc,
        ["#", "Événement", "Destinataire", "Déclencheur"],
        [
            ["1", "OTP inscription / vérif e-mail", "Prospect", "Demande OTP"],
            ["2", "Bienvenue + essai Cabinet 14 j", "Admin cabinet", "Inscription réussie"],
            ["3", "Rappel essai J-7", "Admin cabinet", "Celery (horaire)"],
            ["4", "Rappel essai J-3", "Admin cabinet", "Celery (horaire)"],
            ["5", "Rappel essai J-1 (dernier jour)", "Admin cabinet", "Celery (horaire)"],
            ["6", "Essai terminé → Découverte", "Admin cabinet", "Fin trial + bascule Free"],
            ["7", "Plan activé / changé", "Admin cabinet", "Sélection plan (mock)"],
            ["8", "Résiliation demandée", "Admin cabinet", "POST /billing/cancel"],
            ["9", "Résiliation annulée", "Admin cabinet", "Undo cancel"],
            ["10", "Facture prête (provisoire)", "Admin cabinet", "Génération facture"],
            ["11", "Alerte quota plaidoiries ~80 %", "Admin cabinet", "Incrément usage"],
            ["12", "Invitation à rejoindre le cabinet", "Invité", "Admin → Inviter"],
            ["13", "Bienvenue après acceptation invite", "Nouveau membre", "Accept invite"],
            ["14", "Membre a rejoint (notif admin)", "Admin cabinet", "Accept invite"],
            ["15", "Demande join rejetée (legacy)", "Candidat", "Si flux rejet encore actif"],
            ["16", "Reset mot de passe", "Utilisateur", "Forgot password"],
            ["17", "Mot de passe modifié", "Utilisateur", "Changement MDP"],
            ["18", "Compte verrouillé (bruteforce)", "Utilisateur", "Trop d'échecs login"],
            ["19", "Compte supprimé", "Utilisateur", "Suppression compte"],
            ["20", "Plaidoirie validée", "Assistant assigné", "Validation avocat"],
            ["21", "Digest plaidoiries validées", "Abonnés digest", "Job digest"],
            ["22", "Rappel agenda ~48 h", "Équipe cabinet", "Celery 15 min"],
            ["23", "Cabinet approuvé (legacy)", "Manager", "Si flux activation manuelle"],
            ["24", "Paiement échoué (past_due)", "Admin cabinet", "Webhook / simulate payment-failed"],
            ["25", "Renouvellement J-7 / J-1", "Admin cabinet", "Celery daily 9h"],
            ["26", "Reçu / facture (fiscal-ready)", "Admin cabinet", "Génération facture"],
            ["27", "Tip onboarding J+1 (1ʳᵉ plaidoirie)", "Admin cabinet", "Celery daily"],
            ["28", "Tip onboarding J+2 (équipe / dossier)", "Admin cabinet", "Celery daily"],
            ["29", "Win-back inactivité 14 j", "Admin cabinet", "Celery daily"],
            ["30", "Checklist post-upgrade", "Admin cabinet", "Sélection plan payant"],
            ["31", "Digest hebdo activité cabinet", "Admin cabinet", "Celery lundi 9h30"],
            ["32", "Changelog soft (mensuel)", "Admin paid/trial", "1er du mois"],
            ["33", "Sièges presque pleins", "Admin cabinet", "Invite / usage ≥80 %"],
            ["34", "Nudge 2FA", "Admin sans 2FA", "Celery daily (1×/mois)"],
        ],
        col_widths=[1.0, 5.5, 3.5, 5.0],
    )

    add_heading_custom(doc, "3. Idées intelligentes in-app (complément e-mail)", 1)
    make_table(
        doc,
        ["Insight", "Quand", "CTA"],
        [
            ["1ʳᵉ plaidoirie", "0 plaidoirie, cabinet < 21 j", "Atelier plaidoirie"],
            ["Invite assistant", "Plan Solo/Cabinet, 0 assistant", "Équipe"],
            ["2FA", "Admin sans TOTP, ≥ 2 j", "Paramètres sécurité"],
            ["Sièges saturés", "≥ 80 % avocats/assistants", "Upgrade"],
            ["Renouvellement", "≤ 7 j avant fin de période", "Abonnement"],
            ["Paiement échoué", "status past_due", "Régulariser"],
        ],
        col_widths=[4.0, 6.5, 4.5],
    )
    add_para(
        doc,
        "Prochaines idées (backlog) : A/B copy upgrade, NPS post-validation plaidoirie, "
        "rappel soft si audience demain sans plaidoirie liée, pack annuel -2 mois mis en avant.",
        italic=True,
        size=10,
    )

    add_heading_custom(doc, "4. UX commerciale in-app", 1)
    add_bullets(
        doc,
        [
            "Bannière essai (urgent ≤ 3 j) + CTA « Choisir mon plan ».",
            "Bannière Découverte post-essai → Solo / Cabinet.",
            "Bannière quota ≥ 80 % plaidoiries / renouvellement / past_due.",
            "SmartInsightsStrip : tips contextuels dismissables (session).",
            "Pages /tarifs + Paramètres → Abonnement (mock paiement MA).",
        ],
    )

    add_heading_custom(doc, "5. Technique", 1)
    add_para(
        doc,
        "SMTP/Resend via app.services.email ; dédup SubscriptionEvent ; "
        "jobs Celery : trial_lifecycle_hourly, retention_lifecycle_daily (9h), "
        "retention_weekly_digest (lundi). Insights : GET /api/billing/insights. "
        "IA : free quotas jusqu'à clé OpenAI.",
        size=10,
    )

    path = OUT_DIR / "07_Catalogue_emails_LexHub.docx"
    doc.save(path)
    return path


# ===========================================================================
# Main
# ===========================================================================

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [
        build_file1(),
        build_file2(),
        build_file3(),
        build_file4(),
        build_file5_emails(),
    ]
    print("Generated:")
    for p in paths:
        size = p.stat().st_size
        print(f"  {p}  ({size} bytes, {size/1024:.1f} KiB)")


if __name__ == "__main__":
    main()
