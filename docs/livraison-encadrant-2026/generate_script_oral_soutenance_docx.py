# -*- coding: utf-8 -*-
"""Génère le script oral PPT + annexe détails + questions jury (DOCX)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parent / "SCRIPT_ORAL_SOUTENANCE_LexHub_COMPLET.docx"
DESKTOP = Path.home() / "Desktop" / "SCRIPT_ORAL_SOUTENANCE_LexHub_COMPLET.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))
    return p


def add_para(doc, text, bold=False, italic=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=11)
    run.italic = True
    run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10)
    run.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p.paragraph_format.space_after = Pt(10)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C4A35A")
    pBdr.append(bottom)
    pPr.append(pBdr)


def slide(doc, num, title, duree, script, tip=None):
    add_heading_styled(doc, f"Slide {num} — {title}", level=2)
    add_note(doc, f"Durée indicative : {duree}")
    add_para(doc, "À dire :", bold=True, size=10)
    add_quote(doc, script)
    if tip:
        add_note(doc, f"Astuce : {tip}")
    add_hr(doc)


def qa(doc, q, a):
    add_para(doc, f"Q. {q}", bold=True, size=11)
    add_quote(doc, f"R. {a}")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("SCRIPT ORAL — SOUTENANCE LexHub")
    set_run_font(r, size=20, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("PPT · Ce que tu dis · Annexe détails · Questions jury")
    set_run_font(r, size=12, color=RGBColor(0xC4, 0xA3, 0x5A))

    add_para(
        doc,
        "Douae El Assal — ENSA Tétouan — PFA WEBVOICE — https://lexhub.space\n"
        "Durée orale cible : 10–12 min (+ démo 2–3 min + questions).",
        size=10,
    )
    add_note(
        doc,
        "Comment utiliser : lis à voix haute 1–2 fois. Tu n’apprends pas par cœur — "
        "tu suis le fil. Les blocs « À dire » sont prêts à être lus.",
    )

    # ========== PARTIE A ==========
    add_heading_styled(doc, "PARTIE A — Script oral slide par slide", level=1)

    slide(
        doc,
        1,
        "Couverture",
        "20 s",
        "Bonjour. Je m’appelle Douae El Assal, étudiante à l’ENSA Tétouan. "
        "Mon Projet de Fin d’Année a été réalisé chez WEBVOICE, à Tanger. "
        "Encadrant entreprise : Monsieur Hatim Benhaddou. "
        "Tuteur pédagogique : Professeur Mohammed Al Achhab. "
        "Le produit livré s’appelle LexHub, et il est en production sur lexhub.space.",
    )

    slide(
        doc,
        2,
        "Plan",
        "15 s",
        "Je vais présenter le problème et le marché, les fonctionnalités — surtout l’atelier de plaidoirie —, "
        "la conception technique, la rentabilisation, puis le bilan et la démonstration.",
    )

    slide(
        doc,
        3,
        "Problème / contexte",
        "45–60 s",
        "Le problème de départ est concret : pour préparer une plaidoirie, beaucoup d’avocats dictent, "
        "un collaborateur retape, puis on met en page Word juste avant l’audience. C’est lent, peu traçable, "
        "et le secret professionnel est une contrainte forte. "
        "LexHub répond par un SaaS multi-cabinets : atelier en cinq étapes, rôles clairs, "
        "et une IA qui assiste sans inventer des faits.",
    )

    slide(
        doc,
        4,
        "Marché / positionnement",
        "40 s",
        "Les outils de gestion cabinet gèrent dossiers et agenda, mais sans atelier voix vers plaidoirie. "
        "Les outils d’IA juridique font surtout de la recherche ou du Q&A, sans SaaS multi-rôles cabinet. "
        "LexHub se place au croisement : atelier guidé plus équipe plus abonnement.",
        tip="Ne cite pas « sur mesure / devis » comme offre principale.",
    )

    slide(
        doc,
        5,
        "Objectifs",
        "35 s",
        "Objectifs du PFA : livrer un atelier jusqu’au Word, un espace cabinet multi-rôles, "
        "la règle « seul le validé est visible de l’assistant », et un modèle d’abonnement opérationnel. "
        "Règle d’ingénierie : PostgreSQL est la vérité métier ; le LLM propose ; l’avocat reste responsable.",
    )

    slide(
        doc,
        6,
        "Rôles / RBAC",
        "40 s",
        "Quatre rôles. L’avocat dicte, rédige, valide, exporte. "
        "L’admin cabinet fait la même chose et gère aussi l’équipe et l’abonnement. "
        "L’assistant ne voit que les plaidoiries validées. "
        "L’admin plateforme — WebVoice — supervise cabinets et plans, sans lire le contenu des dossiers. "
        "Le cabinet n’est pas un acteur UML : c’est le locataire, le périmètre isolé par cabinet_id.",
    )

    slide(
        doc,
        7,
        "Pilotage cabinet / plateforme (si présente)",
        "35 s",
        "Deux niveaux de pilotage. Le cabinet pilote son équipe et son abonnement. "
        "WebVoice pilote la plateforme — indicateurs globaux — sans accéder au fond des affaires.",
    )

    slide(
        doc,
        8,
        "Atelier plaidoirie — vue produit",
        "45 s",
        "Le cœur du produit, c’est l’atelier en cinq étapes. "
        "Un : dossier et pièces, avec OCR. "
        "Deux : dictée — Web Speech sur Chrome ou Edge, Whisper sinon ou pour un fichier audio. "
        "Trois : transcription éditable, bouton Corriger. "
        "Quatre : rédaction IA et chat de révision. "
        "Cinq : export Word avec en-tête cabinet. "
        "Trois garde-fous : pas de faits inventés, relecture humaine, identités pseudonymisées avant l’IA.",
        tip="Ici tu montres le PRODUIT. Pas l’architecture.",
    )

    slide(
        doc,
        9,
        "Dossier / données qui alimentent l’atelier",
        "30 s",
        "L’atelier s’appuie sur le dossier : fiche, parties, pièces indexées. "
        "C’est ce contexte qui nourrit la rédaction et le masquage des noms.",
    )

    slide(
        doc,
        10,
        "Fonctionnalités annexes (agenda, équipe…)",
        "25–40 s",
        "Autour de l’atelier : planning et rappels, équipe par invitation, paramètres cabinet. "
        "Je n’entre pas dans chaque écran — le fil conducteur reste la plaidoirie.",
        tip="Adapte si ta numérotation PPT diffère — garde l’idée.",
    )

    slide(
        doc,
        11,
        "Pilotage / plans (si slide plans)",
        "35 s",
        "Côté cabinet : l’admin choisit le plan, suit les quotas ; les avocats consultent sans payer. "
        "Côté plateforme : supervision globale, pas de lecture des dossiers.",
    )

    slide(
        doc,
        12,
        "Transition conception",
        "15 s",
        "Je passe maintenant à la conception : comment c’est câblé, puis le pipeline métier, "
        "la protection des données, et le stockage.",
    )

    slide(
        doc,
        13,
        "Architecture 3-tiers",
        "45 s",
        "Architecture client-serveur en trois tiers. "
        "Présentation : navigateur Next.js, entrée HTTPS via Caddy. "
        "Métier : API FastAPI — seul point d’entrée vers les données — avec Celery pour les tâches longues. "
        "Données et services : PostgreSQL, Redis, MinIO en privé ; à l’extérieur OpenAI, Whisper, OCR, Resend, Turnstile. "
        "Le navigateur ne touche jamais la base. L’API contrôle rôles, cabinet et règles métier.",
        tip="Transition : « Que fait le système quand l’avocat travaille ? Le pipeline. »",
    )

    slide(
        doc,
        14,
        "Pipeline métier — humain dans la boucle",
        "45 s",
        "Même parcours que l’atelier, mais vue conception : à chaque étape, l’avocat agit, "
        "la tech exécute — OCR, Whisper, LLM — et on a un livrable. "
        "L’humain reste dans la boucle, surtout avant l’export Word. "
        "Ce n’est pas une deuxième liste produit : c’est qui fait quoi.",
        tip="Ne pas refaire le slide 8 mot pour mot.",
    )

    slide(
        doc,
        15,
        "Pseudonymisation (coulisse backend)",
        "60–75 s",
        "Cette slide n’est pas une étape de l’atelier : c’est une coulisse backend, avant l’appel à l’API OpenAI. "
        "Entrée : texte encore en clair — dictée, pièces OCR, fiche dossier. "
        "Masquage : le module pii_mask remplace les identités par des codes — PARTIE_1, CIN_1 — "
        "et pii_store enregistre la correspondance sur le dossier, en PostgreSQL, table cases, champ pii_mapping_json. "
        "OpenAI reçoit uniquement les codes. "
        "Restauration : le même module remet les vrais noms grâce au dictionnaire. "
        "Puis l’avocat voit le Word. "
        "Marge d’erreur : une donnée non reconnue peut partir en clair vers l’API — "
        "ce n’est pas un NER entraîné. Perspective : renforcer plus tard avec NER ou DLP.",
        tip="Dis « codes » ; « jetons » = même chose. Un module, deux sens : aller et retour.",
    )

    slide(
        doc,
        16,
        "Trois espaces de stockage",
        "45 s",
        "Règle : isolation multi-tenant. Base et schéma partagés, chaque table porte cabinet_id, "
        "l’API filtre toujours avec le JWT. "
        "Trois espaces : PostgreSQL = vérité métier — comptes, dossiers, brouillons, abonnements, "
        "et aussi le dictionnaire PII. "
        "MinIO = fichiers chiffrés — pièces, logos, Word. "
        "Redis = files d’attente — OCR, e-mails — et notifications. "
        "Exemple pièce : fichier dans MinIO, métadonnées dans Postgres, OCR via Redis, texte OCR revenu en Postgres.",
    )

    slide(
        doc,
        17,
        "Déploiement / CI-CD (si présente)",
        "30 s",
        "Production sur VPS, Docker Compose, Caddy pour le TLS. "
        "Mise à jour : images via GHCR, pull sur le serveur, redémarrage des services. "
        "URL : lexhub.space.",
    )

    slide(
        doc,
        18,
        "Rentabilisation — logique ingénieur",
        "50 s",
        "On facture de la valeur métier, pas des tokens. Une rédaction plaidoirie = dix crédits. "
        "Le coût API estimé par rédaction est très inférieur au prix packagé du crédit dans les plans. "
        "Plans au-dessus du coût, packs, sièges, essai hybride. "
        "Le détail des calculs est dans le rapport — ici je montre la logique produit.",
        tip="Si on t’embrouille sur les tokens : ramène à « 1 rédaction = 10 crédits ».",
    )

    slide(
        doc,
        19,
        "Plans / packs / croissance (si présente)",
        "40 s",
        "Découverte, Standard, Expert — l’admin cabinet choisit et paie. "
        "Packs de crédits pour les pics d’usage. "
        "Croissance : sièges supplémentaires, parrainage, e-mails lifecycle.",
    )

    slide(
        doc,
        20,
        "Tests / qualité (si présente)",
        "30 s",
        "Parcours testés de bout en bout : inscription, atelier, validation assistant, isolation cabinet, "
        "masquage PII visible dans le flux, export Word.",
    )

    slide(
        doc,
        21,
        "Démo — intro",
        "15 s",
        "Je vais maintenant montrer LexHub en production. "
        "Je me concentre sur l’atelier et, si le temps le permet, la règle assistant / validé.",
    )

    slide(
        doc,
        22,
        "Démo (ou captures)",
        "2–3 min",
        "Pendant la démo je commente à voix haute : création ou ouverture dossier, dépôt pièce, "
        "dictée ou texte, correction, rédaction, validation, export. "
        "J’insiste : tant que ce n’est pas validé, l’assistant ne voit pas le brouillon.",
        tip="Prépare un compte démo + kit fictif avant la soutenance.",
    )

    slide(
        doc,
        23,
        "Résultats, limites, perspectives",
        "50 s",
        "Résultats : plateforme en production, atelier bout en bout, isolation multi-tenant et RBAC, "
        "grille tarifaire opérationnelle. "
        "Limites : pas de comptabilité cabinet, pas de recherche juridique intégrée, pas d’app mobile, "
        "pas de lien Mahakim, et un masquage PII par règles — pas NER ou DLP industriel. "
        "Perspectives : paiement carte réel, mobile, recherche juridique, passerelle tribunaux, "
        "renforcement PII type Presidio, extension géographique.",
    )

    slide(
        doc,
        24,
        "Conclusion",
        "30 s",
        "LexHub n’est pas un chatbot qui invente le droit. "
        "C’est un outil de cabinet : de la voix au Word, avec rôles, isolation, et un produit réellement déployé. "
        "Je vous remercie. Je suis prête pour vos questions.",
    )

    slide(
        doc,
        25,
        "Questions / merci",
        "—",
        "Je vous remercie pour votre attention. Je suis à votre disposition pour les questions.",
        tip="Si ta pagination PPT diffère de 1–25, mappe mentalement par TITRE de slide, pas par numéro.",
    )

    # ========== PARTIE B ==========
    doc.add_page_break()
    add_heading_styled(doc, "PARTIE B — Annexe détails (pour maîtriser)", level=1)
    add_note(
        doc,
        "Tu ne lis pas ça à l’oral sauf si on te questionne. C’est ta fiche de secours.",
    )

    add_heading_styled(doc, "B1. Pseudonymisation — film exact", level=2)
    add_para(
        doc,
        "Ce n’est PAS une étape visible de l’atelier. C’est une coulisse : juste avant chaque appel LLM "
        "(correction, rédaction, révision…).",
    )
    add_para(doc, "1) Entrée : texte clair = dictée + OCR pièces + fiche dossier.", bold=False)
    add_para(doc, "2) Masquage : pii_mask.py remplace → PARTIE_1, CIN_1, TEL_*, EMAIL_*.")
    add_para(doc, "3) Stockage du dictionnaire : pii_store.py écrit cases.pii_mapping_json (PostgreSQL).")
    add_para(doc, "4) OpenAI API reçoit uniquement les codes (pas ChatGPT grand public / abonnement Pro).")
    add_para(doc, "5) Restauration : unmask_text dans pii_mask.py lit le dictionnaire et remet les noms.")
    add_para(doc, "6) UI / Word : l’avocat voit le texte en clair.")

    add_heading_styled(doc, "B2. Deux types de données masquées (un seul module)", level=2)
    add_para(
        doc,
        "Un seul programme pii_mask.py. À l’intérieur :",
    )
    add_para(doc, "• Structuré (CIN, tél, mail) → regex (forme stable).")
    add_para(doc, "• Noms → parties du dossier + heuristiques sur la dictée (plus fragile).")
    add_para(
        doc,
        "Ne dis PAS « deux moteurs » = deux applications. Dis : un module, deux types de remplacements.",
    )

    add_heading_styled(doc, "B3. Où est stocké le dictionnaire ?", level=2)
    add_para(doc, "Base : PostgreSQL")
    add_para(doc, "Table : cases")
    add_para(doc, "Colonne : pii_mapping_json")
    add_para(doc, "Module Python qui lit/écrit : pii_store.py")
    add_para(
        doc,
        "Exemple JSON : {\"PARTIE_1\": \"Ahmed Kadiri\", \"CIN_1\": \"BE123456\"}",
        italic=True,
    )

    add_heading_styled(doc, "B4. Marge d’erreur — formulation honnête", level=2)
    add_para(
        doc,
        "Si un nom n’est ni dans la fiche ni bien détecté dans la dictée, il peut partir en clair vers OpenAI. "
        "La relecture Word contrôle le LIVRABLE ; elle n’annule pas un envoi API déjà fait. "
        "Mitigation actuelle = règles + mapping dossier. Perspective = NER/DLP.",
    )

    add_heading_styled(doc, "B5. Architecture — une phrase par brique", level=2)
    pairs = [
        ("Caddy", "Portier HTTPS du site."),
        ("Next.js", "Interface dans le navigateur."),
        ("FastAPI", "Cerveau : auth, rôles, règles, appels services."),
        ("PostgreSQL", "Vérité métier (tables + pii_mapping_json)."),
        ("MinIO", "Fichiers (pièces, logos, exports)."),
        ("Redis + Celery", "File d’attente + workers (OCR, mails…)."),
        ("Whisper / Web Speech", "Voix → texte."),
        ("OpenAI API", "LLM pour corriger / rédiger / réviser."),
        ("Resend", "E-mails (OTP, invitations…)."),
        ("Turnstile", "Anti-robot après 3 échecs login."),
    ]
    for k, v in pairs:
        add_para(doc, f"{k} — {v}")

    add_heading_styled(doc, "B6. Stockage — cycle d’une pièce", level=2)
    add_para(doc, "Fichier → MinIO. Métadonnées (nom, chemin, cabinet_id) → PostgreSQL. Tâche OCR → Redis/Celery. Texte OCR → PostgreSQL.")

    add_heading_styled(doc, "B7. Crédits (rappel oral)", level=2)
    add_para(doc, "1 rédaction plaidoirie = 10 crédits. On vend de la valeur métier, pas des tokens. Détail dans le rapport de rentabilisation.")

    add_heading_styled(doc, "B8. Phrases magiques si tu bloques", level=2)
    add_quote(
        doc,
        "« Son rôle dans LexHub, c’est [utilité]. Le détail d’implémentation est dans le rapport ; "
        "ce que je maîtrise, c’est son utilité dans le parcours avocat. »",
    )
    add_quote(
        doc,
        "« Le slide 8 = l’écran. Le 13 = les tuyaux. Le 14 = l’humain + la machine. Le 15 = la coulisse PII. »",
    )

    # ========== PARTIE C ==========
    doc.add_page_break()
    add_heading_styled(doc, "PARTIE C — Questions jury difficiles + réponses", level=1)

    qa(
        doc,
        "Pourquoi ne pas utiliser seulement ChatGPT ?",
        "ChatGPT seul n’isole pas les cabinets, n’a pas les rôles avocat/assistant, "
        "ne masque pas les PII avant l’appel, et n’exporte pas un acte au nom du cabinet avec validation.",
    )
    qa(
        doc,
        "C’est du ChatGPT ou de l’API OpenAI ?",
        "On appelle l’API OpenAI (usage facturé à l’usage), pas l’abonnement ChatGPT grand public. "
        "Le LLM est un service externe appelé par notre backend.",
    )
    qa(
        doc,
        "Comment vous protégez le secret professionnel ?",
        "Isolation multi-tenant par cabinet_id, RBAC, TLS, navigateur sans accès direct à la base, "
        "pseudonymisation avant l’API, admin plateforme sans lecture des dossiers.",
    )
    qa(
        doc,
        "Comment fonctionne la pseudonymisation ?",
        "Coulisse backend avant l’API. pii_mask remplace par des codes ; pii_store stocke la correspondance "
        "dans cases.pii_mapping_json ; après la réponse on restaure chez nous.",
    )
    qa(
        doc,
        "C’est de la NER ? Vous avez entraîné un modèle ?",
        "Non. C’est de la pseudonymisation par règles : regex pour CIN/tél/mail, noms du dossier et heuristiques dictée. "
        "Pas un NER entraîné. Renforcement NER/DLP = perspective.",
    )
    qa(
        doc,
        "Et si la dictée contient un nom qui n’est pas dans la fiche dossier ?",
        "On tente de l’extraire par heuristiques, mais ce n’est pas parfait. "
        "Marge d’erreur : une PII non reconnue peut partir en clair vers l’API. "
        "La relecture Word ne « rollback » pas cet envoi.",
    )
    qa(
        doc,
        "Dans quelle table PostgreSQL est le mapping PII ?",
        "Table cases, colonne pii_mapping_json. Le module pii_store.py lit et écrit ce champ. "
        "Ce n’est pas une table SQL nommée pii_store.",
    )
    qa(
        doc,
        "Pourquoi MinIO et pas les fichiers dans PostgreSQL ?",
        "Postgres reste léger : métadonnées. Les binaires (PDF, audio, Word) vont dans MinIO, stockage objet adapté.",
    )
    qa(
        doc,
        "À quoi sert Redis ?",
        "Broker/file pour Celery : tâches asynchrones OCR, e-mails, etc., et notifications temps réel. "
        "Ce n’est pas la base métier.",
    )
    qa(
        doc,
        "Le cabinet est-il un acteur UML ?",
        "Non. Le cabinet est le locataire multi-tenant. Les acteurs sont avocat, assistant, admin cabinet, admin plateforme.",
    )
    qa(
        doc,
        "Que voit l’assistant ?",
        "Uniquement les plaidoiries validées par l’avocat. Les brouillons restent invisibles.",
    )
    qa(
        doc,
        "Whisper ou Web Speech ?",
        "Web Speech dans Chrome/Edge pour la dictée navigateur. Whisper pour Firefox ou import fichier audio.",
    )
    qa(
        doc,
        "Comment justifiez-vous les prix / les crédits ?",
        "On packagé une valeur métier : une rédaction = 10 crédits. "
        "Le coût API estimé par rédaction est inférieur au prix du crédit dans l’offre. "
        "Détail des hypothèses dans le rapport de rentabilisation.",
    )
    qa(
        doc,
        "Le paiement carte est-il réel ?",
        "Le parcours d’abonnement et de packs est opérationnel côté produit ; "
        "l’encaissement carte marchand réel (CMI/PayZone) est une perspective annoncée.",
    )
    qa(
        doc,
        "Où est hébergé LexHub ?",
        "VPS, Docker Compose, Caddy TLS, domaine lexhub.space. Déploiement via images GHCR.",
    )
    qa(
        doc,
        "Comment l’API empêche un cabinet de voir un autre ?",
        "Chaque entité métier a un cabinet_id. Après auth JWT, toutes les requêtes sont filtrées sur ce cabinet. "
        "Pas d’accès direct navigateur → base.",
    )
    qa(
        doc,
        "Qu’est-ce qui est livré concrètement ?",
        "Une plateforme en production : atelier dictée→Word, équipe, plans/crédits, admin plateforme, "
        "isolation et pseudonymisation déployées.",
    )
    qa(
        doc,
        "Différence slide atelier et slide pipeline ?",
        "Atelier = ce que l’avocat voit dans l’UI. Pipeline = qui fait quoi (avocat / tech / livrable). "
        "Architecture = tuyaux. Pseudonymisation = coulisse sécurité avant l’API.",
    )
    qa(
        doc,
        "Loi 09-08 / conformité ?",
        "On minimise l’exposition : pas d’identités en clair vers le fournisseur IA quand le masquage réussit, "
        "données cabinet isolées, accès contrôlés. Ce n’est pas une certification juridique complète — "
        "c’est une conception alignée secret professionnel et minimisation.",
    )
    qa(
        doc,
        "Que feriez-vous pour renforcer le masquage ?",
        "Ajouter une couche NER/DLP type Presidio, plus de types PII (adresses, IBAN…), "
        "et une alerte si donnée sensible suspecte non masquée avant l’appel LLM.",
    )

    add_heading_styled(doc, "Checklist 2 minutes avant", level=2)
    for item in [
        "PPT ouvert + mode présentateur",
        "lexhub.space prêt (compte démo connecté)",
        "Kit démo fictif / audio sur le Bureau",
        "Ce DOCX ouvert ou imprimé (Partie A + C)",
        "Eau — respirer — commencer par le bonjour Slide 1",
    ]:
        add_para(doc, f"☐ {item}")

    add_para(
        doc,
        "Fin du document — LexHub PFA ENSA Tétouan — WEBVOICE",
        bold=True,
        size=10,
    )

    doc.save(OUT)
    try:
        doc.save(DESKTOP)
    except Exception:
        pass
    return OUT, DESKTOP


if __name__ == "__main__":
    out, desk = build()
    print("OK", out)
    print("DESKTOP", desk, "exists=", desk.exists())
