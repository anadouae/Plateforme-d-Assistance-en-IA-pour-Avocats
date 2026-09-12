# -*- coding: utf-8 -*-
"""Proposition commerciale LexHub — état réel de la production (déjà déployé)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent
LOGO = OUT / "logo-lexhub.png"
NAVY = RGBColor(0x0A, 0x12, 0x28)
GOLD = RGBColor(0xC9, 0xA8, 0x4A)
MUTED = RGBColor(0x4A, 0x55, 0x68)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=6, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "0A1228")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_run_font(r, size=9)
            if ri % 2:
                shade(cell, "EEF1F8")
    doc.add_paragraph()
    return t


def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    for run in list(p.runs):
        run.text = ""
    r = p.add_run(text)
    set_run_font(r, size=size)
    p.paragraph_format.space_after = Pt(3)
    return p


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)

    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Cm(3.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LexHub — Plan de croissance & rentabilité")
    set_run_font(r, size=18, bold=True, color=NAVY)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("État réel de la production déployée sur https://lexhub.space")
    set_run_font(r, size=11, italic=True, color=MUTED)

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run(
        "Tarifs · essais · quotas · usage à la demande · parrainage — déjà en place · WEBVOICE / ENSA · septembre 2026"
    )
    set_run_font(r, size=10, color=GOLD)

    add_para(
        doc,
        "Ce document décrit ce qui tourne aujourd’hui en production. "
        "Ce n’est pas une feuille de route : abonnements, quotas, usage à la demande et parrainage "
        "sont déjà implémentés dans le backend (billing.py, referral.py) et visibles dans l’interface "
        "(Paramètres → Plan & usage, landing, inscription). "
        "Chaque chiffre et chaque règle ci-dessous peut être vérifié sur lexhub.space.",
        size=10,
        italic=True,
        color=MUTED,
    )

    # —— 1 ——
    heading(doc, "1. Ce que LexHub monétise déjà", 1)
    add_para(
        doc,
        "LexHub est un SaaS multi-tenant pour cabinets d’avocats au Maroc : dictée → plaidoirie "
        "(arabe juridique), dossiers, pièces, planning, équipe. "
        "Le revenu vient des abonnements payants (Solo, Cabinet, Entreprise). "
        "L’acquisition repose sur le freemium Découverte, les essais carte, et le programme de parrainage "
        "cabinet → cabinet — déjà actif. "
        "La marge est protégée par des quotas mensuels / journaliers et par l’usage à la demande "
        "(option admin, désactivée par défaut, jamais silencieuse).",
        size=10,
    )
    bullet(doc, "Seul l’admin cabinet paie, change de plan, active l’usage à la demande et gère le code de parrainage.")
    bullet(doc, "Les avocats / assistants du même cabinet voient les quotas en lecture seule ; ils ne paient pas.")
    bullet(doc, "Interface FR / AR. Facturation affichée en MAD.")

    # —— 2 ——
    heading(doc, "2. Grille tarifaire en production", 1)
    add_para(
        doc,
        "Les quatre plans suivants sont ceux servis par l’API /api/billing/plans et affichés sur /tarifs. "
        "Ils sont seedés / synchronisés au démarrage du backend (DEFAULT_PLANS).",
        size=10,
    )
    table(
        doc,
        ["Plan", "Mensuel", "Annuel", "Essai carte", "Sièges", "Plaidoiries / mois", "Crédits IA / mois"],
        [
            ["Découverte (free)", "0 MAD", "—", "aucun", "1 avocat · 0 assistant", "3", "50"],
            ["Solo", "299 MAD", "2 990 MAD", "14 jours", "1 avocat + 1 assistant", "35", "600"],
            ["Cabinet (recommandé)", "699 MAD", "6 990 MAD", "30 jours", "5 avocats + 3 assistants", "120", "2 000"],
            ["Entreprise", "1 290 MAD", "12 900 MAD", "30 jours", "15 avocats + 10 assistants", "350", "5 000"],
        ],
    )
    add_para(
        doc,
        "Annuel ≈ 10 mois payés pour 12 mois (réduction structurelle déjà codée dans les prix_yearly). "
        "À partir de Solo : export PDF, logo cabinet, invitations d’équipe. "
        "Cabinet / Entreprise : IA premium renforcée (dont clé de barre / analyses plus larges). "
        "Dossiers : limités à 5 sur Découverte ; illimités dès Solo.",
        size=10,
    )

    heading(doc, "2.1 Quotas IA déjà appliqués à chaque action", 2)
    add_para(
        doc,
        "Chaque action IA consomme des crédits (AI_ACTION_CREDITS). "
        "En parallèle, certaines actions ont un plafond d’occurrences / mois (AI_MONTHLY_LIMITS). "
        "Les deux garde-fous s’appliquent ensemble : il faut assez de crédits ET ne pas dépasser le compteur d’action.",
        size=10,
    )
    table(
        doc,
        ["Action dans l’app", "Coût en crédits", "Max / mois Découverte", "Solo", "Cabinet", "Entreprise"],
        [
            ["Rédaction plaidoirie (compose)", "10", "3 plaidoiries*", "35*", "120*", "350*"],
            ["Révision du texte", "5", "8", "60", "250", "600"],
            ["Analyse", "8", "1", "12", "50", "150"],
            ["Synthèse / résumé", "4", "3", "25", "100", "300"],
            ["Question sur pièces (ask)", "2", "10", "100", "400", "1 200"],
            ["Clé de barre (fiche orale audience)", "8", "1", "10", "50", "150"],
            ["Budget crédits IA total / mois", "—", "50", "600", "2 000", "5 000"],
        ],
    )
    add_para(
        doc,
        "* Une rédaction consomme 1 unité du quota « plaidoiries » et 10 crédits IA. "
        "Exemple Solo : à 35 plaidoiries × 10 = 350 crédits rien que pour la rédaction ; "
        "il reste des crédits pour révisions, analyses, etc., dans la limite de 600.",
        size=9,
        italic=True,
    )
    add_para(
        doc,
        "Anti-abus déjà actif : plafond journalier de crédits (Découverte 12 · Solo 50 · Cabinet 180 · Entreprise 300) "
        "et nombre de slots IA parallèles (1 / 1 / 2 / 3). "
        "Ces plafonds empêchent un drain OpenAI en quelques heures même si le quota mensuel n’est pas encore épuisé.",
        size=10,
    )

    add_para(
        doc,
        "Note — « Clé de barre » : ce n’est pas un code d’accès. C’est une fiche orale préparée par l’IA "
        "à partir des pièces du dossier (faits / moyens / demandes / points d’audience) pour aider l’avocat "
        "à plaider à la barre. Compte dans les quotas IA comme les autres actions.",
        size=9,
        italic=True,
    )

    # —— 3 ——
    heading(doc, "3. Justification des prix (Van Westendorp)", 1)
    add_para(
        doc,
        "Les prix publics ont été positionnés avec Van Westendorp (sensibilité prix) : "
        "l’optimum retenu est celui déployé, ni le plancher (qui dégrade la perception) ni le plafond (qui freine la conversion).",
        size=10,
    )
    table(
        doc,
        ["Plan", "Prix déployé", "Plancher étudié", "Plafond étudié", "Élasticité"],
        [
            ["Solo", "299 MAD / mois", "199", "399", "Élevée (−1,8) → essai 14 j + annuel"],
            ["Cabinet", "699 MAD / mois", "499", "899", "Moyenne (−1,2) → offre phare, essai 30 j"],
            ["Entreprise", "1 290 MAD / mois", "990", "1 790", "Faible (−0,7) → peu de discount public"],
        ],
    )

    # —— 4 ——
    heading(doc, "4. Cycle d’abonnement — comportement réel de l’app", 1)
    add_para(
        doc,
        "Voici exactement ce que fait LexHub aujourd’hui, dans l’ordre.",
        size=10,
    )
    for i, line in enumerate(
        [
            "Inscription : e-mail OTP → création du cabinet + compte admin. Abonnement Découverte démarré automatiquement.",
            "Code de parrainage optionnel à l’inscription : le filleul colle le code d’un cabinet déjà client (ou ouvre un lien du type …/login?tab=register&ref=CODE). LexHub mémorise alors « ce nouveau cabinet a été invité par tel cabinet ».",
            "L’admin ouvre Paramètres → Plan & usage (ou /tarifs) et choisit Solo / Cabinet / Entreprise.",
            "Essai payant avec carte (modèle type Coursera) : 14 jours Solo, 30 jours Cabinet / Entreprise. 0 MAD pendant l’essai.",
            "Si le filleul a un parrainage valide : +14 jours ajoutés automatiquement à la fin d’essai au démarrage de cet essai.",
            "Pendant l’essai : l’admin peut résilier → aucun débit à la fin.",
            "Fin d’essai + carte encore enregistrée + pas d’annulation → statut ACTIVE payant ; période mensuelle ou annuelle démarre.",
            "À ce moment, si le filleul avait un parrain : le parrain reçoit +30 jours sur current_period_end (une fois).",
            "Sans conversion / après fin d’abo / essai annulé → retour Découverte. Les données dossier restent.",
            "Résiliation en période payante : motif obligatoire (liste codée), effet en fin de période déjà payée.",
            "Échec de paiement : statut past_due ; après 7 jours de grâce, restriction des actions IA / plaidoirie.",
            "Paiement Maroc : CMI et PayZone sont branchés côté environnement (statut « prêt opérateur ») ; Stripe international « bientôt ».",
        ],
        1,
    ):
        add_para(doc, f"{i}. {line}", size=10)

    # —— 4.1 Essai = quotas du plan (pas illimité) ——
    heading(doc, "4.1 Pendant l’essai — le client n’est pas illimité", 2)
    add_para(
        doc,
        "Trois cabinets peuvent démarrer un essai le même jour (Solo, Cabinet ★ parfois appelé « Pro », "
        "Entreprise) : chacun est guidé par les quotas de son propre plan. L’essai suspend uniquement "
        "la facturation (0 MAD pendant la fenêtre) — ce n’est pas un accès illimité.",
        size=10,
    )
    add_para(
        doc,
        "Découverte n’a pas d’essai carte : freemium permanent (1 avocat · 3 plaidoiries/mois · 50 crédits IA/mois).",
        size=9,
        italic=True,
        color=MUTED,
    )
    add_para(doc, "Durée d’essai et quotas applicables dès le 1er jour", size=10, bold=True)
    table(
        doc,
        [
            "Plan",
            "Durée essai",
            "Utilisateurs",
            "Plaidoiries/mois",
            "Crédits IA/mois",
            "Plafond IA/jour",
            "Slots IA //",
        ],
        [
            ["Solo", "14 jours", "1 avocat + 1 assistant", "35", "600", "50", "1"],
            ["Cabinet ★", "30 jours", "5 avocats + 3 assistants", "120", "2 000", "180", "2"],
            ["Entreprise", "30 jours", "15 avocats + 10 assistants", "350", "5 000", "300", "3"],
        ],
    )
    add_para(doc, "Comment le client est guidé pendant l’essai", size=10, bold=True)
    bullet(doc, "Carte bancaire pour démarrer l’essai, mais 0 MAD débité pendant la durée d’essai.")
    bullet(
        doc,
        "Résiliation avant la fin d’essai → aucun débit. Sinon l’abonnement du plan démarre automatiquement à la fin.",
    )
    bullet(
        doc,
        "Mêmes quotas que le plan payant dès le jour 1 (crédits, plaidoiries, plafonds journaliers, sièges). "
        "Essai Entreprise ≠ essai Solo.",
    )
    bullet(
        doc,
        "Quota atteint : usage à la demande (~0,25 MAD/crédit) possible — OFF par défaut ; impossible sur Découverte.",
    )
    bullet(doc, "Parrainage : filleul Solo/Cabinet peut recevoir +14 jours d’essai via un code valide.")
    add_para(
        doc,
        "En résumé : aucun essai Solo / Cabinet / Entreprise n’est illimité — seule la facturation est suspendue "
        "pendant 14 ou 30 jours selon le plan.",
        size=10,
        bold=True,
    )

    # —— 5 ——
    heading(doc, "5. Usage à la demande — déjà implémenté", 1)
    add_para(
        doc,
        "Fonctionnalité livrée et visible dans Paramètres → Plan & usage. "
        "Principe : le forfait du plan est inclus dans l’abonnement. "
        "Au-delà, soit l’app bloque, soit l’admin a explicitement accepté de payer le surplus.",
        size=10,
    )

    heading(doc, "5.1 Comportement par défaut (production)", 2)
    bullet(doc, "Découverte : usage à la demande impossible. Quota atteint = message d’upgrade ou attendre le mois suivant.")
    bullet(doc, "Solo / Cabinet / Entreprise : interrupteur « Usage à la demande » OFF par défaut (on_demand_enabled = false).")
    bullet(doc, "OFF = arrêt net. Aucun crédit hors forfait, aucun MAD de surplus sur la facture.")
    bullet(doc, "Seul l’admin cabinet peut changer le réglage (PATCH /api/billing/on-demand).")

    heading(doc, "5.2 Quand l’admin active l’option", 2)
    bullet(doc, "Tarif fixe codé : 0,25 MAD par crédit IA hors forfait (AI_OVERAGE_MAD_PER_CREDIT).")
    bullet(doc, "Les crédits hors forfait sont comptés (ai_overage_credits) et convertis en estimation MAD affichée dans l’UI.")
    bullet(doc, "Trois modes UI : Désactivé · Activé sans plafond · Activé avec plafond mensuel en MAD (on_demand_mad_limit).")
    bullet(doc, "Si le plafond MAD est atteint : l’action suivante est refusée avec message clair (augmenter le plafond ou attendre).")
    bullet(doc, "Le surplus n’est jamais silencieux : l’admin le voit avant et sur la facture.")

    heading(doc, "5.3 Cas gérés dans l’application", 2)
    table(
        doc,
        ["Cas réel", "Réponse de LexHub aujourd’hui"],
        [
            [
                "Solo, 600/600 crédits, usage OFF",
                "L’app refuse l’action (écran / message d’erreur : quota du mois atteint). Elle propose : activer l’usage à la demande (~0,25 MAD par crédit en plus), passer à un plan supérieur, ou attendre le mois suivant. (Code technique HTTP 402 = paiement / quota requis.)",
            ],
            [
                "Solo, crédits épuisés, usage ON, sans plafond",
                "Action autorisée. Crédits comptés en overage → MAD estimés dans Plan & usage.",
            ],
            [
                "Cabinet, usage ON, plafond 100 MAD, surplus déjà ≈ 100 MAD",
                "Action refusée. Message : plafond d’usage à la demande atteint.",
            ],
            [
                "Découverte, 3/3 plaidoiries ou 50/50 crédits",
                "Stop. Pas d’overage possible. Proposition Solo / Cabinet.",
            ],
            [
                "Plaidoiries Solo 35/35, usage OFF",
                "Rédaction bloquée (même logique de garde-fou que les crédits).",
            ],
            [
                "Admin désactive l’usage après avoir généré du surplus",
                "Plus aucun nouveau surplus ; le surplus déjà compté reste visible pour facturation.",
            ],
        ],
    )

    # —— 6 ——
    heading(doc, "6. Parrainage cabinet → cabinet — déjà en production", 1)
    add_para(
        doc,
        "Le programme est live sur lexhub.space : section landing #parrainage, liens navbar, "
        "champ code à l’inscription, API GET /api/billing/referral, bloc dans Plan & usage, "
        "popup de rappel périodique pour avocat / admin cabinet. "
        "Migration base 034 (referral_code, cabinet_referrals) appliquée.",
        size=10,
    )
    add_para(
        doc,
        "En français simple : un cabinet invite un confrère. Le nouveau gagne du temps d’essai. "
        "L’ancien gagne un mois offert seulement quand le nouveau paie vraiment. "
        "Pas d’argent cash versé : ce sont des jours d’abonnement ajoutés automatiquement.",
        size=10,
    )

    heading(doc, "6.1 Parcours exact (déjà codé)", 2)
    for i, line in enumerate(
        [
            "Chaque cabinet obtient un code unique (ex. LH-CABINE-XXXX) via ensure_referral_code — créé au premier affichage Plan & usage / API referral.",
            "Le parrain copie le code ou le lien https://lexhub.space/login?tab=register&ref=CODE.",
            "Le filleul s’inscrit avec ce code. LexHub enregistre referred_by_cabinet_id + une ligne cabinet_referrals (status pending).",
            "Quand le filleul choisit Solo ou Cabinet et commence son essai (carte enregistrée, 0 MAD pendant l’essai) : LexHub lui ajoute automatiquement 14 jours d’essai en plus. Une seule fois par filleul.",
            "Quand le filleul passe en abonnement payant actif (fin d’essai convertie ou select_plan payant) : le parrain reçoit +30 jours sur current_period_end (une fois, referrer_reward_applied).",
            "Si le filleul annule avant la fin d’essai : pas de paiement → le parrain ne reçoit rien. C’est voulu.",
        ],
        1,
    ):
        add_para(doc, f"{i}. {line}", size=10)

    heading(doc, "6.2 Règles anti-abus déjà enforced", 2)
    table(
        doc,
        ["Règle", "Implémentation actuelle"],
        [
            ["Périmètre", "Cabinet → cabinet uniquement (code lié au cabinet, pas à un user isolé)."],
            ["Auto-parrainage", "Refusé si referrer.id == nouveau cabinet."],
            ["Code invalide", "HTTP 400 à l’inscription (« Code de parrainage invalide »)."],
            ["Bonus filleul", "+14 jours, une seule fois par filleul."],
            ["Récompense parrain", "+30 jours, une seule fois par filleul, au paiement réel."],
            ["Plafond annuel", "Max 12 lignes cabinet_referrals créées / an / parrain (MAX_ACTIVE_REFERRALS_PER_YEAR)."],
            ["UI parrain", "Code + lien + stats (invites / pending / rewarded) dans Plan & usage."],
            ["Rappel produit", "Popup LexProductModal ~40 s après connexion, cooldown localStorage ~5 jours."],
        ],
    )
    add_para(
        doc,
        "Intérêt économique déjà aligné : CAC bas (pas de cashback), récompense conditionnée au revenu, "
        "plafond annuel contre les fermes de comptes. C’est un levier de croissance déjà branché, pas un concept.",
        size=10,
    )

    # —— 7 ——
    heading(doc, "7. Fidélisation commerciale (procédure WEBVOICE)", 1)
    add_para(
        doc,
        "Les règles produit (quotas, essais, parrainage, usage à la demande) sont dans le code. "
        "Les gestes commerciaux ci-dessous sont la procédure humaine WEBVOICE pour retenir un client payant "
        "sans republier une grille promo sur /tarifs. Une offre max par client / 12 mois ; remises tracées en facturation.",
        size=10,
    )
    table(
        doc,
        ["Situation", "Offre proposée"],
        [
            ["Solo payant ≥ 6 mois, demande d’annulation", "1 mois offert OU −20 % sur le prochain trimestre"],
            ["Cabinet payant ≥ 6 mois, ≥ 3 sièges utilisés", "1 mois offert OU bonus sur passage annuel"],
            ["Entreprise ≥ 6 mois", "Jusqu’à −15 % sur 3 mois OU +2 sièges assistants 1 mois"],
            ["Motif « trop cher »", "Downgrade assisté (Cabinet → Solo) plutôt qu’un prix public baissé"],
            ["Motif « pause » (≥ 3 mois payants)", "Gel 30–60 jours sans perte de données"],
            ["Win-back ≤ 90 j après départ", "Réactivation avec 14 j d’essai du dernier plan payant"],
        ],
    )

    # —— 8 ——
    heading(doc, "8. Indicateurs à suivre (an 1)", 1)
    table(
        doc,
        ["Indicateur", "Cible indicative"],
        [
            ["Conversion essai → payant (Cabinet)", "≥ 25 %"],
            ["Part d’abonnements annuels", "≥ 30 % du MRR"],
            ["Churn mensuel payant", "< 5 %"],
            ["Cabinets ayant affiché / partagé leur code parrainage", "≥ 40 %"],
            ["Filleuls devenus payants / invitations", "≥ 20 %"],
            ["Acceptation offre de rétention", "≥ 35 % des demandes éligibles"],
            ["Upsell Solo → Cabinet (> 3 mois)", "≥ 15 %"],
        ],
    )

    # —— 9 ——
    heading(doc, "9. Synthèse pour l’encadrant", 1)
    add_para(
        doc,
        "Tout ce qui suit est déjà en production sur https://lexhub.space :",
        size=10,
        bold=True,
    )
    bullet(doc, "Grille 0 / 299 / 699 / 1 290 MAD + essais + quotas détaillés.")
    bullet(doc, "Essais carte 14 / 30 jours, retour Découverte, résiliation motif, grâce past_due 7 j.")
    bullet(doc, "Essai = quotas du plan choisi (pas illimité) ; 0 MAD jusqu’à conversion ou résiliation (§4.1).")
    bullet(doc, "Usage à la demande OFF par défaut, 0,25 MAD/crédit, plafonds MAD optionnels, UI claire.")
    bullet(doc, "Parrainage cabinet → cabinet : +14 j filleul / +30 j parrain au paiement, plafond 12/an, UI + landing + popup.")
    add_para(
        doc,
        "Priorités restantes côté business / ops (pas produit manquant) : finaliser l’activation opérateur CMI/PayZone, "
        "appliquer la procédure de rétention §7, et suivre les KPIs §8 — en particulier le funnel parrainage.",
        size=10,
    )

    add_para(
        doc,
        "Preuves techniques : backend/app/services/billing.py · referral.py · alembic 034_referral_parrainage · "
        "API /api/billing/{plans,entitlements,on-demand,referral} · UI Plan & usage · landing #parrainage. "
        "Production vérifiée septembre 2026.",
        size=9,
        italic=True,
        color=MUTED,
    )

    path = OUT / "01_Proposition_plan_rentabilisation_LexHub_v2.docx"
    for candidate in (
        path,
        OUT / "01_Proposition_plan_rentabilisation_LexHub_v3.docx",
        OUT / "01_Proposition_plan_rentabilisation_LexHub_v4.docx",
    ):
        try:
            doc.save(candidate)
            return candidate
        except PermissionError:
            continue
    raise PermissionError("Fermez les fichiers Word ouverts puis relancez le script.")


if __name__ == "__main__":
    print(build())
