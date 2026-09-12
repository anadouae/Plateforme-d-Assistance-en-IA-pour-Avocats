# -*- coding: utf-8 -*-
"""Pack encadrant entreprise : guide à jour + note BDD + copie kit famille."""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DESK = Path(r"C:\Users\daurinia\Desktop\documents de  rentabilisation et guide de test & use")
KIT = ROOT / "docs" / "kit-test-plaidoirie-famille"
NAVY = RGBColor(0x0B, 0x1A, 0x4D)
MUTED = RGBColor(0x4A, 0x55, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def para(doc, text, *, bold=False, size=11, after=6, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    font(r, size=size, bold=bold, color=color)
    return p


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(head)
        font(r, size=8, bold=True, color=WHITE)
        shade(cell, "0B1A4D")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            font(r, size=8)
            if ri % 2:
                shade(cell, "F1F5F9")
    doc.add_paragraph()


def setup(doc):
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)


def write_guide() -> Path:
    doc = Document()
    setup(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LexHub")
    font(r, size=20, bold=True, color=NAVY)
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run("Guide d’utilisation & de tests — production")
    font(r, size=13, bold=True, color=NAVY)
    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t3.add_run(
        f"https://lexhub.space · {date.today().strftime('%d/%m/%Y')} · "
        "Livrable encadrant entreprise (WebVoice) — stagiaire PFA"
    )
    font(r, size=9, italic=True, color=MUTED)

    para(
        doc,
        "LexHub est la plateforme du cabinet (produit WebVoice). "
        "Ce document décrit le comportement réel en production, pour tester et former.",
        size=10,
    )

    h(doc, "1. Comptes de test", 1)
    para(doc, "Mot de passe cabinet démo : Sanad2026!  ·  Admin plateforme : WebVoice2026!", size=10)
    table(
        doc,
        ["E-mail", "Rôle", "Quoi tester"],
        [
            ["demo@sanad.ma", "Admin cabinet", "Tout le produit + équipe + plan"],
            ["avocat@sanad.ma", "Avocat", "Dossiers, dictée, plaidoirie"],
            ["assistant@sanad.ma", "Assistant", "Pièces + plaidoiries validées (pas d’IA lourde)"],
            ["admin@webvoice.ma", "Admin plateforme", "Cabinets, plans, journaux — pas le fond des affaires"],
        ],
    )
    para(
        doc,
        "Si un bandeau « compte mémorisé » s’affiche : le navigateur a gardé une visite précédente "
        "(même après redémarrage). Cliquez « Se connecter avec un autre compte », puis saisissez l’e-mail. "
        "Ce n’est pas une session bloquée sur le serveur.",
        size=10,
    )

    h(doc, "2. Ce que fait vraiment le produit", 1)
    table(
        doc,
        ["Sujet", "Comportement prod"],
        [
            ["Plaidoirie", "5 étapes : dossier → dictée → transcription → rédaction IA → export"],
            ["Dictée", "Micro ou fichier audio · STT cloud · arabe"],
            ["Corriger", "Forme fusha, pas de résumé. Texte sans sens (ex. k;k;l;) = refus"],
            ["Nouvelle plaidoirie", "Choix : reprendre le brouillon local OU page vide"],
            ["Assistant", "Voit les plaidoiries validées. Aperçu + texte extrait. Pas de synthèse IA"],
            ["Équipe", "Invitation e-mail uniquement (pas de code cabinet à l’écran)"],
            ["IA dossier", "Synthèse / Q&A / clé de barre = avocat et admin cabinet"],
            ["Suppression", "Pièce = déposant · dossier = avocat responsable"],
            ["Plans", "Découverte 0 · Solo 299 · Cabinet 699 · Entreprise 1290 MAD / mois"],
            ["Paiement CMI", "Non branché en live — parcours préparé"],
            ["Clé de barre", "Fiche orale d’audience, pas un mot de passe"],
        ],
    )

    h(doc, "3. Créer un dossier (aides jaunes)", 1)
    para(
        doc,
        "En cliquant un champ, une carte jaune explique le sens, pourquoi c’est utile à la plaidoirie, "
        "et donne un exemple. Rien n’est obligatoire hors client, matière, degré, ville, tribunal, rôle.",
        size=10,
    )
    para(doc, "Kit famille : recopier docs/kit-test-plaidoirie-famille/00_FICHE_DOSSIER.txt", size=10)

    h(doc, "4. Parcours plaidoirie famille (à cocher)", 1)
    for line in [
        "Connexion demo@sanad.ma",
        "Dossiers → Nouveau → matière Famille → valeurs de la fiche",
        "Uploader les DOCX 01 à 03 (et 05 si présent)",
        "Plaidoiries → Nouvelle → Commencer une nouvelle",
        "Lier le dossier → micro → lire 04_Texte_dictée…",
        "Corriger → Rédaction IA → relire les faits (2000 DH, Yusuf, Rabat)",
        "Export Word + PDF",
        "Valider → se connecter assistant → la plaidoirie est visible",
    ]:
        p = doc.add_paragraph(style="List Number")
        for run in list(p.runs):
            run.text = ""
        r = p.add_run(line)
        font(r, size=10)

    h(doc, "5. Rôles — ce que l’assistant ne fait pas", 1)
    para(
        doc,
        "Pas de dictée, pas de rédaction IA, pas de synthèse de pièce, pas de clé de barre. "
        "Il dépose des pièces, lit le texte extrait, imprime les plaidoiries validées.",
        size=10,
    )

    dest = OUT / "02_Guide_utilisation_et_tests_LexHub_v7.docx"
    doc.save(dest)
    return dest


def write_db() -> Path:
    doc = Document()
    setup(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LexHub — Note base de données")
    font(r, size=16, bold=True, color=NAVY)
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run(
        f"PostgreSQL 16 · multi-tenant cabinet_id · {date.today().strftime('%d/%m/%Y')}\n"
        "Pour l’encadrant entreprise — lecture métier, sans mots de passe de production."
    )
    font(r, size=9, italic=True, color=MUTED)

    para(
        doc,
        "La base s’appelle encore « sanad » sur le serveur (héritage du nom de code). "
        "Les fichiers (PDF, Word, audio) sont dans MinIO ; PostgreSQL garde les métadonnées et le texte extrait. "
        "Adminer n’est pas public : accès via tunnel SSH sur le VPS uniquement.",
        size=10,
    )

    h(doc, "1. Principe d’isolation", 1)
    para(
        doc,
        "Chaque cabinet est un locataire. Les requêtes métier filtrent sur cabinet_id. "
        "L’admin plateforme voit les cabinets, pas le contenu des affaires.",
        size=10,
    )

    h(doc, "2. Tables principales", 1)
    table(
        doc,
        ["Table", "Rôle"],
        [
            ["cabinets", "Tenant : nom, ville, logo, parrainage, statut"],
            ["users", "Comptes + rôle + 2FA + rattachement cabinet"],
            ["cases", "Dossiers + identité juridique (tribunal, rôles, RG…)"],
            ["documents", "Métadonnées pièces + texte OCR + uploader"],
            ["plaidoirie_drafts", "Brouillons / validées, assignation, épinglage"],
            ["agenda_events", "Planning + lieu + rappels"],
            ["cabinet_notifications", "Cloche cabinet"],
            ["cabinet_invites", "Invitations e-mail (chemin prod)"],
            ["subscriptions / invoices", "Plans, essais, factures"],
            ["ai_credits / usage", "Quotas et consommation IA"],
            ["audit_logs", "Journal (action, acteur, IP)"],
            ["membership_requests", "Ancien flux code cabinet — non exposé UI"],
        ],
    )

    h(doc, "3. Ce que l’encadrant peut demander (sans secret)", 1)
    para(
        doc,
        "Un compte PostgreSQL en lecture seule, limité au schéma métier, "
        "peut être créé pour audit interne. Ne jamais livrer le mot de passe root "
        "ni le .env dans un PDF académique. Sauvegardes : dump PostgreSQL + volume MinIO "
        "(voir docs/BACKUP_RESTORE.md dans le dépôt).",
        size=10,
    )

    h(doc, "4. Hors base", 1)
    para(
        doc,
        "Redis : files Celery et cache court. MinIO : binaires. "
        "Caddy : TLS. Les jetons de rafraîchissement de session sont des cookies HttpOnly.",
        size=10,
    )

    dest = OUT / "03_Note_base_de_donnees_LexHub.docx"
    doc.save(dest)
    return dest


def copy_pack(guide: Path, dbnote: Path) -> Path:
    DESK.mkdir(parents=True, exist_ok=True)
    kit_dest = DESK / "kit-test-plaidoirie-famille"
    if kit_dest.exists():
        shutil.rmtree(kit_dest)
    shutil.copytree(
        KIT,
        kit_dest,
        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"),
    )
    shutil.copy2(guide, DESK / guide.name)
    shutil.copy2(dbnote, DESK / dbnote.name)
    rent = OUT / "01_Proposition_plan_rentabilisation_LexHub_v5.docx"
    if rent.exists():
        shutil.copy2(rent, DESK / rent.name)
    older = DESK / "Guide_utilisation_et_tests_LexHub_reorganise.docx"
    # keep existing rentabilisation already in folder
    readme = (
        "# Livraison encadrant entreprise — LexHub\n\n"
        f"Date : {date.today().isoformat()}\n"
        "Produit : https://lexhub.space (réalisé en PFA, stagiaire développeuse).\n\n"
        "| Fichier | Usage |\n"
        "|---|---|\n"
        "| 01_Proposition_plan_rentabilisation_LexHub_v5.docx | Stratégie / tarifs |\n"
        f"| {guide.name} | Utilisation & tests (prod à jour) |\n"
        f"| {dbnote.name} | Base de données (lecture métier) |\n"
        "| kit-test-plaidoirie-famille/ | Dossier famille + dictées arabes |\n"
        "| LexHub_Strategie_rentabilisation.docx | Version déjà dans ce dossier |\n"
    )
    (DESK / "README.txt").write_text(readme, encoding="utf-8")
    return DESK


def main() -> None:
    guide = write_guide()
    dbnote = write_db()
    dest = copy_pack(guide, dbnote)
    print("OK", guide)
    print("OK", dbnote)
    print("OK pack", dest)


if __name__ == "__main__":
    main()
