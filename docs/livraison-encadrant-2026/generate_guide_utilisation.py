# -*- coding: utf-8 -*-
"""
Guide d'utilisation LexHub — aligné STRICTEMENT sur la production https://lexhub.space
Source de vérité vérifiée : GET /api/billing/plans + code frontend/backend déployé (sept. 2026).

Régénérer :
  python docs/livraison-encadrant-2026/generate_guide_utilisation.py
→ écrase 02_Guide_utilisation_et_tests_LexHub.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent
LOGO = OUT / "logo-lexhub.png"
NAVY = RGBColor(0x0A, 0x12, 0x28)
GOLD = RGBColor(0xC9, 0xA8, 0x4A)
MUTED = RGBColor(0x4A, 0x55, 0x68)

# ── Grille EXACTE prod (API lexhub.space/api/billing/plans — vérifiée) ──
PLANS_PROD = [
    ["Découverte", "0", "—", "—", "1 avocat · 0 assistant", "5", "3", "50"],
    ["Solo", "299", "2 990", "14 j", "1 avocat + 1 assistant", "∞", "35", "600"],
    ["Cabinet ★", "699", "6 990", "30 j", "5 avocats + 3 assistants", "∞", "120", "2 000"],
    ["Entreprise", "1 290", "12 900", "30 j", "15 avocats + 10 assistants", "∞", "350", "5 000"],
]

AI_LIMITS = [
    ["Rédaction plaidoirie", "10", "3*", "35*", "120*", "350*"],
    ["Révision", "5", "8", "60", "250", "600"],
    ["Analyse", "8", "1", "12", "50", "150"],
    ["Synthèse", "4", "3", "25", "100", "300"],
    ["Question (ask)", "2", "10", "100", "400", "1 200"],
    ["Clé de barre (fiche orale)", "8", "1", "10", "50", "150"],
    ["Budget crédits / mois", "—", "50", "600", "2 000", "5 000"],
]


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=6, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, bold=True, size=8, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "0A1228")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_run_font(r, size=8)
            if ri % 2:
                shade(cell, "EEF1F8")
    doc.add_paragraph()
    return t


def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    for run in list(p.runs):
        run.text = ""
    r = p.add_run(text)
    set_run_font(r, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def numbered(doc, text, size=10):
    p = doc.add_paragraph(style="List Number")
    for run in list(p.runs):
        run.text = ""
    r = p.add_run(text)
    set_run_font(r, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def tip(doc, text):
    add_para(doc, "Astuce : " + text, size=9, italic=True, color=MUTED)


def warn(doc, text):
    add_para(doc, "Important : " + text, size=9, bold=True, color=RGBColor(0x9A, 0x34, 0x1C))


def build() -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.7)
        s.right_margin = Cm(1.7)

    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Cm(3.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LexHub — Guide d’utilisation & de test")
    set_run_font(r, size=20, bold=True, color=NAVY)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Aligné sur la production déployée · https://lexhub.space")
    set_run_font(r, size=11, italic=True, color=MUTED)

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run("Comptes de test · rôles · parcours · septembre 2026")
    set_run_font(r, size=10, color=GOLD)

    add_para(
        doc,
        "Ce document décrit UNIQUEMENT ce qui tourne aujourd’hui sur lexhub.space "
        "(API /api/billing/plans + interface réelle). "
        "Utilisez le sommaire pour aller à une section, puis revenez-y dès que vous changez de sujet.",
        size=10,
    )
    warn(
        doc,
        "Ne pas utiliser d’anciennes grilles (ex. 10/50/200/500 plaidoiries) : "
        "la prod est 3 / 35 / 120 / 350 plaidoiries et 50 / 600 / 2 000 / 5 000 crédits IA.",
    )

    # —— SOMMAIRE ——
    heading(doc, "Sommaire", 1)
    toc = [
        "1. Démarrage en 2 minutes",
        "2. Comptes de test",
        "3. Acteurs & rôles (matrice)",
        "4. Menus réels de l’application",
        "5. Pages publiques",
        "6. Connexion / langue / déconnexion",
        "7. Tableau de bord",
        "8. Plaidoirie — 5 étapes déployées",
        "9. Dossiers, pièces, clé de barre",
        "10. Documents · Planning · Activité · Équipe",
        "11. Paramètres · Plan & usage · parrainage · usage à la demande",
        "12. Plans & quotas EXACTS (prod)",
        "13. Admin plateforme WEBVOICE",
        "14. Scénarios de visite",
        "15. Problèmes fréquents",
        "16. Checklist « j’ai tout vu »",
    ]
    for line in toc:
        add_para(doc, line, size=11, space_after=2)

    # —— 1 ——
    heading(doc, "1. Démarrage en 2 minutes", 1)
    numbered(doc, "Ouvrir https://lexhub.space dans Google Chrome (micro plus fiable).")
    numbered(doc, "Connexion → compte demo@sanad.ma / Sanad2026! (§2).")
    numbered(doc, "Suivre le scénario A (§14) pour une visite cabinet complète.")
    tip(doc, "Hard refresh Ctrl+F5 si l’écran semble ancien (cache).")

    # —— 2 ——
    heading(doc, "2. Comptes de test", 1)
    add_para(
        doc,
        "Créés par le script seed_test_accounts.py. Toujours se déconnecter avant de changer de compte.",
        size=10,
    )
    table(
        doc,
        ["E-mail", "Mot de passe", "Rôle", "Cabinet", "Pour tester"],
        [
            ["demo@sanad.ma", "Sanad2026!", "Admin cabinet", "Démo Atlas (Casa)", "Tout : équipe, abo, plaidoirie, dossiers"],
            ["avocat@sanad.ma", "Sanad2026!", "Avocat", "Démo Atlas", "Dictée / dossiers — pas Équipe ni paiement"],
            ["assistant@sanad.ma", "Sanad2026!", "Assistant", "Démo Atlas", "Pièces + lecture ; menu « Plaidoiries validées »"],
            ["test@cabinet-rabat.ma", "Test2026!", "Admin cabinet", "Test Rabat", "Isolation : ne voit PAS Atlas"],
            ["admin@webvoice.ma", "WebVoice2026!", "Admin plateforme", "WEBVOICE", "/admin — pas les dossiers clients"],
        ],
    )

    # —— 3 ——
    heading(doc, "3. Acteurs & rôles (matrice)", 1)
    table(
        doc,
        ["Action", "Admin cabinet", "Avocat", "Assistant", "Admin plateforme"],
        [
            ["Espace après login", "Cabinet (/dashboard)", "Cabinet", "Cabinet", "/admin"],
            ["Menu Plaidoirie (dictée IA)", "Oui (= avocat + admin)", "Oui", "Non → « Plaidoiries validées »", "Non"],
            ["Créer / supprimer dossier", "Oui", "Oui", "Non", "Non"],
            ["Déposer pièces", "Oui", "Oui", "Oui", "Non"],
            ["Supprimer pièces", "Oui", "Oui", "Non", "Non"],
            ["Clé de barre IA", "Oui", "Oui", "Non", "Non"],
            ["Export Word/PDF", "Oui", "Oui", "Limité", "Non"],
            ["Planning + Activité", "Oui", "Oui", "Oui", "Non"],
            ["Équipe (invites)", "Oui (seul)", "Non", "Non", "Non"],
            ["Plan & usage / payer", "Oui (seul)", "Quotas lecture", "Quotas lecture", "Non"],
            ["Parrainage (code)", "Oui", "Voir / partager", "Non", "Non"],
            ["Zone sensible (quitter / supprimer son compte)", "Oui", "Oui", "Oui", "Non"],
            ["Suspendre / supprimer cabinet", "Non", "Non", "Non", "Oui"],
            ["Éditer plans plateforme", "Non", "Non", "Non", "Oui"],
        ],
    )
    bullet(
        doc,
        "Admin cabinet = avocat responsable du cabinet : fait aussi le produit (dictée, dossiers, clé de barre, export) "
        "+ paie, invite l’équipe, configure logo / abonnement.",
    )
    bullet(
        doc,
        "Avocat = même produit (dictée, dossiers, pièces, export), sans gérer l’abonnement ni l’équipe.",
    )
    bullet(
        doc,
        "Assistant = aide pièces / lecture / plaidoiries validées ; pas d’IA lourde ni suppression.",
    )
    bullet(
        doc,
        "Admin plateforme = ops LexHub (WEBVOICE) : cabinets, plans, logs — pas le travail quotidien d’un dossier client.",
    )

    # —— 4 ——
    heading(doc, "4. Menus réels de l’application", 1)
    add_para(doc, "Menu latéral gauche (compte cabinet) — libellés FR déployés :", size=10)
    table(
        doc,
        ["Menu (icône)", "URL", "Visible pour"],
        [
            ["Logo → Tableau de bord", "/dashboard", "Tous rôles cabinet"],
            ["Plaidoirie", "/redaction", "Admin + Avocat"],
            ["Plaidoiries validées", "/redaction", "Assistant (même URL, UI différente)"],
            ["Dossiers", "/cases", "Tous"],
            ["Documents", "/documents", "Tous"],
            ["Planning", "/agenda", "Tous"],
            ["Activité", "/activity", "Tous"],
            ["Équipe", "/team", "Admin cabinet seulement"],
        ],
    )
    add_para(
        doc,
        "En-tête (droite) : recherche (Ctrl+K), notifications (admin / assistant), "
        "raccourci abonnement (admin), aide, thème, paramètres (engrenage), langue FR/AR, menu compte.",
        size=10,
    )
    add_para(
        doc,
        "Paramètres : engrenage / menu compte → /settings "
        "(pas dans le menu latéral). Sections : Général · Cabinet (admin) · Plan & usage · Sécurité · Zone sensible.",
        size=10,
    )

    # —— 5 ——
    heading(doc, "5. Pages publiques", 1)
    bullet(doc, "Landing : / — produit, FAQ, section #parrainage")
    bullet(doc, "Tarifs : /tarifs — grille live")
    bullet(doc, "Fonctionnalités, Sécurité, À propos, Aide, Contact, Mentions, Confidentialité")
    bullet(doc, "Login / inscription : /login")

    # —— 6 ——
    heading(doc, "6. Connexion / langue / déconnexion", 1)
    numbered(doc, "/login → e-mail + mot de passe §2")
    numbered(doc, "Basculer FR ↔ AR (commutateur en-tête)")
    numbered(doc, "Déconnexion : menu compte (haut droite)")
    tip(doc, "Inscription réelle : OTP e-mail + infos cabinet + code parrainage optionnel. Pour la démo, utilisez les comptes seed.")

    # —— 7 ——
    heading(doc, "7. Tableau de bord", 1)
    add_para(
        doc,
        "Après login cabinet : /dashboard. Bannière d’abonnement possible (essai, Découverte…). "
        "Popups possibles : compléter profil, plans, parrainage — fermer ou suivre.",
        size=10,
    )

    # —— 8 ——
    heading(doc, "8. Plaidoirie — 5 étapes déployées", 1)
    add_para(
        doc,
        "Compte : avocat@sanad.ma ou demo@sanad.ma. Pipeline UI réel (plea-step-carousel) :",
        size=10,
        bold=True,
    )
    table(
        doc,
        ["Étape", "Nom FR (UI réelle)", "Ce que vous faites"],
        [
            ["1", "Dossier & pièces", "Lier le dossier · joindre / vérifier les pièces"],
            ["2", "Dictée", "Micro Chrome ou fichier audio — parler en arabe"],
            ["3", "Transcription", "Corriger / vérifier le texte STT"],
            ["4", "Rédaction IA", "Génération / révision — toujours à relire"],
            ["5", "Export", "Word · PDF · modèle (en-tête cabinet)"],
        ],
    )
    tip(doc, "Assistant : pas cette pipeline complète — accès « Plaidoiries validées » seulement.")

    # —— 9 ——
    heading(doc, "9. Dossiers, pièces, clé de barre", 1)
    heading(doc, "9.1 Dossiers", 2)
    numbered(doc, "Menu Dossiers → créer (admin/avocat) ou ouvrir.")
    numbered(doc, "Déposer des pièces ; attendre l’indexation avant l’IA.")

    heading(doc, "9.2 Clé de barre (fiche orale)", 2)
    add_para(
        doc,
        "Ce n’est PAS un mot de passe. C’est une trame orale (faits / moyens / demandes / points d’audience) "
        "générée depuis les pièces du dossier pour aider à plaider à l’audience.",
        size=10,
    )
    numbered(doc, "Ouvrir un dossier avec pièces indexées.")
    numbered(doc, "Bouton « Clé de barre IA » sur la fiche dossier.")
    numbered(doc, "Lire le texte à l’écran — responsabilité de l’avocat.")
    tip(doc, "Coûte 8 crédits IA + 1 unité du quota « clé de barre » du plan.")

    # —— 10 ——
    heading(doc, "10. Documents · Planning · Activité · Équipe", 1)
    bullet(doc, "Documents (/documents) : toutes les pièces du cabinet.")
    bullet(doc, "Planning (/agenda) : échéances / audiences.")
    bullet(doc, "Activité (/activity) : journal du cabinet.")
    bullet(doc, "Équipe (/team) : admin only — invite e-mail ou code join ; rôles avocat / assistant / admin.")

    # —— 11 ——
    heading(doc, "11. Paramètres · Plan & usage · parrainage · usage à la demande", 1)
    numbered(doc, "Ouvrir /settings → section Plan & usage (admin gère ; autres voient quotas).")
    numbered(
        doc,
        "Usage à la demande : OFF par défaut. Activé = ~0,25 MAD / crédit hors forfait "
        "(sans plafond ou avec plafond MAD). Jamais silencieux. Indisponible sur Découverte.",
    )
    numbered(
        doc,
        "Parrainage : copier code LH-… ou lien ?ref=CODE. "
        "Filleul : +14 jours d’essai quand il démarre Solo/Cabinet. "
        "Parrain : +30 jours (~1 mois) quand le filleul devient payant après l’essai. "
        "Max 12 / an · pas d’auto-parrainage.",
    )
    numbered(doc, "Cabinet : logo, nom AR (exports). Sécurité : mot de passe / 2FA. Zone sensible : fermer compte.")

    # —— 12 ——
    heading(doc, "12. Plans & quotas EXACTS (prod)", 1)
    add_para(
        doc,
        "Vérifié via https://lexhub.space/api/billing/plans — ne pas recopier d’anciens PDF.",
        size=9,
        italic=True,
        color=MUTED,
    )
    table(
        doc,
        ["Plan", "MAD/mois", "MAD/an", "Essai", "Sièges", "Dossiers", "Plaidoiries/mois", "Crédits IA/mois"],
        PLANS_PROD,
    )
    add_para(doc, "∞ = illimité (−1 dans l’API). Annuel ≈ 10 mois payés / 12.", size=9, italic=True)
    heading(doc, "12.1 Quotas IA par action (prod)", 2)
    table(
        doc,
        ["Action", "Crédits", "Découverte", "Solo", "Cabinet", "Entreprise"],
        AI_LIMITS,
    )
    add_para(
        doc,
        "* Plaidoiries = quota dédié + 10 crédits IA chacune. "
        "Plafonds journaliers crédits : 12 / 50 / 180 / 300. Slots parallèles IA : 1 / 1 / 2 / 3.",
        size=9,
        italic=True,
    )
    add_para(
        doc,
        "Paiement Maroc : CMI / PayZone (env prêt, activation opérateur). "
        "Essai carte : 0 MAD pendant l’essai ; résiliation avant la fin = pas de débit.",
        size=10,
    )

    # —— 13 ——
    heading(doc, "13. Admin plateforme WEBVOICE", 1)
    add_para(doc, "Compte : admin@webvoice.ma / WebVoice2026!", size=10, bold=True)
    table(
        doc,
        ["Menu /admin", "URL", "Actions"],
        [
            ["Tableau de bord", "/admin", "Métriques plateforme"],
            ["Cabinets", "/admin/cabinets", "Détail, Suspendre/Activer, Supprimer (taper supprimer)"],
            ["Plans & tarifs", "/admin/plans", "Grille abonnements"],
            ["Journal d’audit", "/admin/logs", "Logs (IP masquée en UI)"],
        ],
    )
    warn(doc, "Supprimer un cabinet = définitif (comptes, dossiers, pièces). Cabinet « WebVoice Platform » protégé.")

    # —— 14 ——
    heading(doc, "14. Scénarios de visite", 1)
    heading(doc, "A — Cabinet complet (~35 min) — demo@sanad.ma", 2)
    for step in [
        "Login → dashboard",
        "Dossier → pièce → Clé de barre IA",
        "Plaidoirie : étapes 1→5 (micro court) → export",
        "Planning → événement",
        "Équipe → voir membres",
        "Paramètres → Plan & usage → code parrainage + usage à la demande OFF",
        "Déconnexion",
    ]:
        numbered(doc, step)

    heading(doc, "B — Avocat vs assistant (~15 min)", 2)
    numbered(doc, "avocat@sanad.ma → pas Équipe / pas paiement")
    numbered(doc, "assistant@sanad.ma → tenter supprimer dossier (refusé) ; menu Plaidoiries validées")

    heading(doc, "C — Multi-tenant (~10 min)", 2)
    numbered(doc, "test@cabinet-rabat.ma / Test2026! → pas de dossiers Atlas")

    heading(doc, "D — Admin plateforme (~15 min)", 2)
    numbered(doc, "admin@webvoice.ma → /admin/cabinets → détail (sans supprimer sans intention)")

    heading(doc, "E — Public (~5 min)", 2)
    numbered(doc, "Sans compte : landing #parrainage + /tarifs")

    # —— 15 ——
    heading(doc, "15. Problèmes fréquents", 1)
    table(
        doc,
        ["Symptôme", "Cause", "Action"],
        [
            ["IA / clé de barre refusée", "Quota ou rôle assistant", "demo/avocat ou attendre / upgrader"],
            ["Message quota / « 402 »", "Forfait mois fini, usage OFF", "Admin : usage à la demande ou autre plan"],
            ["Pas de menu Équipe", "Pas admin cabinet", "demo@sanad.ma"],
            ["Arrivé sur /admin", "Compte webvoice", "Déconnexion → demo@…"],
            ["Micro KO", "Navigateur / permission", "Chrome + autoriser ; sinon fichier audio"],
            ["Écran « vieux »", "Cache", "Ctrl+F5"],
        ],
    )

    # —— 16 ——
    heading(doc, "16. Checklist « j’ai tout vu »", 1)
    for item in [
        "Landing + #parrainage + /tarifs",
        "demo : dashboard, dossier, clé de barre, plaidoirie 5 étapes, planning, équipe, Plan & usage",
        "avocat : limites",
        "assistant : limites + Plaidoiries validées",
        "Rabat : isolation",
        "admin plateforme : /admin/cabinets",
    ]:
        bullet(doc, "☐  " + item)

    add_para(
        doc,
        "Sources : GET https://lexhub.space/api/billing/plans · billing.py · referral.py · "
        "app-sidebar.tsx · plea-step-carousel.tsx · admin-shell.tsx · seed_test_accounts.py. "
        "Régénérer ce fichier : python docs/livraison-encadrant-2026/generate_guide_utilisation.py",
        size=8,
        italic=True,
        color=MUTED,
    )

    path = OUT / "02_Guide_utilisation_et_tests_LexHub.docx"
    for candidate in (
        path,
        OUT / "02_Guide_utilisation_et_tests_LexHub_PROD.docx",
        OUT / "02_Guide_utilisation_et_tests_LexHub_AJOUR.docx",
        OUT / "02_Guide_utilisation_et_tests_LexHub_v5.docx",
    ):
        try:
            doc.save(candidate)
            return candidate
        except PermissionError:
            continue
    raise PermissionError("Fermez le .docx Word ouvert puis relancez le script.")


if __name__ == "__main__":
    print(build())
