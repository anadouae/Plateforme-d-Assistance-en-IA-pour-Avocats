# -*- coding: utf-8 -*-
"""Livraison finale Hatim — plans + guide test + vérification rôles + design."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
NAVY = RGBColor(0x0A, 0x12, 0x28)
GOLD = RGBColor(0xFF, 0x6A, 0x00)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=6, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "0A1228")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_run_font(r, size=9)
            if ri % 2:
                shade(cell, "EEF1F8")
    doc.add_paragraph()
    return t


def add_image(doc, path: Path, width_cm=16):
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_para(doc, f"[Image manquante : {path.name}]", italic=True, size=9)


def build_plans():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("LexHub — Proposition commerciale")
    set_run_font(r, size=22, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Plans · tarifs MAD · descriptions · WEBVOICE / Hatim Benhaddou · août 2026")
    set_run_font(r, size=11, italic=True)

    add_image(doc, ASSETS / "lexhub-product-hero.png", 15)
    add_para(
        doc,
        "LexHub (https://lexhub.space) — LegalTech pour cabinets marocains : dictée arabe → "
        "fusḥā → structure de plaidoirie → export Word/PDF, dossiers, équipe, planning.",
        size=10,
    )

    heading(doc, "1. Vision pricing (pro, 2030)", 1)
    add_para(
        doc,
        "Modèle freemium + essai 14 jours du plan Cabinet (recommandé), inspiré Cursor / Clio / "
        "Notion : activation rapide, upgrade clair, pas d'essai 30 jours trop long pour un outil IA.",
    )

    heading(doc, "2. Grille tarifaire", 1)
    add_image(doc, ASSETS / "lexhub-plans-tarifs.png", 15)
    table(
        doc,
        ["Plan", "MAD / mois", "MAD / an*", "Sièges", "Plaidoiries / mois", "Positionnement"],
        [
            ["Découverte", "0", "0", "1 avocat", "5", "Freemium permanent"],
            ["Solo", "299", "≈ 2 990", "1 + 1 assistant", "50", "Avocat indépendant"],
            ["Cabinet ★", "699", "≈ 6 990", "5 + 3", "200", "Recommandé — équipe + clé de barre"],
            ["Entreprise", "1 890", "sur devis", "20 + 10", "Élevé / illimité soft", "Gros cabinets"],
        ],
    )
    add_para(doc, "* Annuel : affichage commercial (≈ 2 mois offres typiques SaaS) — mock paiement CMI/PayZone.", italic=True, size=9)

    heading(doc, "3. Descriptions par plan", 1)
    plans = [
        (
            "Découverte (0 MAD)",
            "Tester LexHub sans carte. 1 avocat, quotas bas. Ideal pour découvrir la dictée. "
            "Pas d'invitations équipe. Upsell naturel après essai ou saturation quota.",
        ),
        (
            "Solo (299 MAD/mois)",
            "Avocat solo + 1 assistant. Volume plaidoiries sérieux, exports, dossiers. "
            "Parfait barreau individuel qui délègue la mise en forme.",
        ),
        (
            "Cabinet (699 MAD/mois) — recommandé",
            "Cœur de marché : plusieurs avocats, assistants, clé de barre IA, branding cabinet, "
            "invites e-mail. Essai 14 j à l'inscription. Meilleur rapport valeur / prix pour un cabinet moyen.",
        ),
        (
            "Entreprise (1 890 MAD/mois)",
            "Gros volumes, sièges élevés, pilotage multi-équipes. Sur-mesure possible (SSO, facturation ICE).",
        ),
    ]
    for name, desc in plans:
        add_para(doc, name, bold=True, color=NAVY)
        add_para(doc, desc, size=10)

    heading(doc, "4. Essai & cycle de vie", 1)
    add_para(
        doc,
        "Inscription → essai Cabinet 14 j → e-mails J-7 / J-3 / J-1 → fin d'essai → Découverte "
        "sauf upgrade. Paiement électronique Maroc (CMI / PayZone / virement) en préparation ; "
        "sélection de plan déjà active (mode mock).",
    )

    heading(doc, "5. Ce qui n'est PAS vendu en marketing", 1)
    add_para(
        doc,
        "Masquage PII interne, détails infra : restent techniques. On vend : temps gagné, "
        "qualité fusḥā, collaboration cabinet, exports officiels.",
        size=10,
    )

    path = OUT / "01_Proposition_plans_tarifs_LexHub.docx"
    doc.save(path)
    return path


def build_guide():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Guide d'utilisation LexHub — tests complets")
    set_run_font(r, size=20, bold=True, color=NAVY)
    add_para(doc, "https://lexhub.space · Destinataire : Hatim / WEBVOICE · août 2026", italic=True, size=10)

    heading(doc, "1. Comptes de test", 1)
    table(
        doc,
        ["E-mail", "Mot de passe", "Rôle", "À tester"],
        [
            ["demo@sanad.ma", "Sanad2026!", "Admin cabinet", "Tout + équipe + abo"],
            ["avocat@sanad.ma", "Sanad2026!", "Avocat", "Dictée, dossiers, pas équipe"],
            ["assistant@sanad.ma", "Sanad2026!", "Assistant", "Plaidoiries validées, docs"],
            ["admin@webvoice.ma", "WebVoice2026!", "Admin plateforme", "/admin cabinets & plans"],
        ],
    )

    heading(doc, "2. Types de plaidoirie (point Hatim)", 1)
    add_para(
        doc,
        "Nouvelle plaidoirie → /redaction/new : choisir un TYPE (civil, pénal, famille…). "
        "Filtre par domaine dans le picker. Le type est envoyé à l'API compose : "
        "il oriente la mise en forme (rôles des parties, sections) SANS réécrire les arguments. "
        "Bibliothèque : filtre « Tous les types ».",
        size=10,
    )
    table(
        doc,
        ["Exemples de types", "Effet"],
        [
            ["civil-general / civil-defense", "مدعي / مدعى عليه"],
            ["penal-defense / partie-civile", "متهم vs طرف مدني"],
            ["famille, commercial, administratif…", "Canevas + postprocess adapté"],
            ["libre", "Page blanche + heuristique soft"],
        ],
    )

    heading(doc, "3. Parcours démo 12 min", 1)
    steps = [
        "Connexion demo@sanad.ma",
        "Plaidoirie → Nouvelle → choisir type (ex. famille) → dicter 1–2 min AR",
        "Fusḥā / contradictions / structure → valider → assigner assistant",
        "Compte assistant : voir plaidoirie validée + export",
        "Équipe → inviter (e-mail) ; Paramètres → Abonnement / tarifs",
        "Admin WebVoice : /admin/plans",
    ]
    for i, s in enumerate(steps, 1):
        add_para(doc, f"{i}. {s}", size=10)

    heading(doc, "4. Matrice des espaces par rôle", 1)
    table(
        doc,
        ["Espace", "Admin cab.", "Avocat", "Assistant", "Admin WV"],
        [
            ["Dashboard", "Oui", "Oui", "Oui (vue stagiaire)", "Non → /admin"],
            ["Plaidoirie / dictée", "Oui", "Oui", "Lecture validées", "Non"],
            ["Dossiers / docs / agenda", "Oui", "Oui", "Oui (docs sans IA)", "Non"],
            ["Équipe / invites", "Oui", "Non", "Non", "Non"],
            ["Abonnement (changer plan)", "Oui", "Voir + demander", "Voir + demander", "Plans globaux"],
            ["Admin plateforme", "Non", "Non", "Non", "Oui"],
        ],
    )

    heading(doc, "5. Invites (plus de code public)", 1)
    add_para(
        doc,
        "Admin → Équipe → Inviter par e-mail → lien /join/accept → mot de passe → compte actif. "
        "Quotas sièges selon le plan.",
        size=10,
    )

    heading(doc, "6. Checklist QA rapide", 1)
    for row in [
        "☐ Login 4 rôles",
        "☐ Nouvelle plaidoirie avec TYPE + filtre bibliothèque",
        "☐ Bannière essai / Découverte selon statut",
        "☐ Avocat voit « Demander upgrade » (pas changer plan seul)",
        "☐ Invite e-mail + acceptation",
        "☐ Export Word/PDF",
        "☐ Admin plans",
    ]:
        add_para(doc, row, size=10)

    path = OUT / "02_Guide_utilisation_et_comptes_test_LexHub.docx"
    doc.save(path)
    return path


def build_verif_design():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Vérification profonde + design LexHub")
    set_run_font(r, size=20, bold=True, color=NAVY)
    add_para(doc, "Audit UX / rôles / palette — livraison finale", italic=True, size=10)

    heading(doc, "1. Clarification « idées intelligentes » (pour Douae)", 1)
    add_para(doc, "Déjà en place (Live)", bold=True, color=NAVY)
    add_para(
        doc,
        "L'app affiche des conseils selon l'USAGE RÉEL du cabinet : "
        "pas de plaidoirie encore → « Dictez votre 1ʳᵉ » ; essai bientôt fini → upgrade ; "
        "quota 80 % → alerte ; admin sans 2FA → sécuriser. Ce n'est PAS du spam publicitaire.",
        size=10,
    )
    add_para(doc, "Backlog (plus tard, pas bloquant pour finir aujourd'hui)", bold=True, color=NAVY)
    add_para(
        doc,
        "NPS après validation, rappel audience sans plaidoirie, pack annuel -2 mois, "
        "tests A/B de textes upgrade. Ce sont des améliorations futures — le cœur produit est déjà là.",
        size=10,
    )

    heading(doc, "2. Palette — est-elle pro ?", 1)
    table(
        doc,
        ["Token", "Hex", "Rôle"],
        [
            ["Navy / ink", "#0A1228", "Identité, titres, CTA"],
            ["Royal", "#0B1A4D / #12307A", "Profondeur"],
            ["Accent gold→orange", "#FF6A00", "CTA secondaires, barre quota"],
            ["Fond cream cool", "#EEF1F8", "Fond app (pas cream chaud cliché)"],
            ["Card", "#FBFCFE", "Surfaces"],
        ],
    )
    add_para(doc, "Verdict", bold=True)
    add_para(
        doc,
        "Oui, la base est professionnelle (navy juridique + accent chaud + fond froid). "
        "Recommandation 2030 : 1) garder navy + orange #FF6A00 comme accent unique "
        "(éviter le vieux doré #C9A227 dans les écrans billing) ; "
        "2) plus d'espace blanc, moins de cards empilées ; "
        "3) typographie display déjà en place — la renforcer sur landing ; "
        "4) motion légère déjà présente — la garder sobre (type Linear/Cursor). "
        "Éviter purple/glow AI générique. Score actuel : ~8/10 LegalTech premium.",
        size=10,
    )

    heading(doc, "3. Vérification espaces (tous acteurs)", 1)
    add_para(
        doc,
        "Contrôles effectués côté code : guards rôle, banners selon entitlements, "
        "picker type plaidoirie remis en parcours, API types complète, CTAs upgrade "
        "différenciés admin vs avocat/assistant.",
        size=10,
    )
    table(
        doc,
        ["Point", "Statut"],
        [
            ["Type plaidoirie → filtre → API", "Corrigé / live"],
            ["Upgrade visible tous rôles cabinet", "Oui (CTA adapté)"],
            ["Assistant sans dictée", "OK"],
            ["Admin WV isolé /admin", "OK"],
            ["Invites e-mail only", "OK"],
            ["E-mails rétention P1–P3", "OK"],
            ["Paiement réel CMI", "Phase 2 (mock)"],
            ["Clé OpenAI", "En attente WEBVOICE"],
        ],
    )

    heading(doc, "4. Inspirations mondiales (IA 2030)", 1)
    add_para(
        doc,
        "Cursor (activation + e-mails cycle) · Clio (LegalTech clair) · Linear (densité pro) · "
        "Notion (essai + upgrade). LexHub combine dictée AR locale + SaaS cabinet — différenciation Maroc.",
        size=10,
    )

    path = OUT / "03_Verification_roles_design_clarifications_LexHub.docx"
    doc.save(path)
    return path


def build_readme():
    text = """# Livraison finale LexHub → Hatim / WEBVOICE

## Contenu de ce dossier

1. `01_Proposition_plans_tarifs_LexHub.docx` — plans, tarifs MAD, descriptions, images
2. `02_Guide_utilisation_et_comptes_test_LexHub.docx` — guide + comptes test + types plaidoirie
3. `03_Verification_roles_design_clarifications_LexHub.docx` — audit rôles + palette + backlog expliqué
4. `assets/` — visuels plans + hero produit
5. Copies utiles : guide / catalogue emails / étude marché (si présents)

## Prod

https://lexhub.space

## Comptes

voir guide 02
"""
    path = OUT / "README.md"
    path.write_text(text, encoding="utf-8")
    return path


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    paths = [build_plans(), build_guide(), build_verif_design(), build_readme()]
    for p in paths:
        print(p, p.stat().st_size if p.exists() else 0)


if __name__ == "__main__":
    main()
