# -*- coding: utf-8 -*-
"""Doc critique marché + clarifications pour Douae / Hatim."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent / "04_Critique_marche_securite_clarifications_LexHub.docx"
NAVY = RGBColor(0x0A, 0x12, 0x28)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = NAVY


def para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LexHub — Critique marché, sécurité & clarifications")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY
    para(doc, "Pour Douae + Hatim · août 2026 · https://lexhub.space", size=10)

    h(doc, "1. C’est quoi le « backlog NPS » ? (simple)", 1)
    para(doc, "NPS = demander aux clients « recommanderiez-vous LexHub ? » (note 0–10).")
    para(doc, "Promo annuel = offrir 2 mois gratuits si paiement annuel.")
    para(
        doc,
        "Ce n’était PAS une tâche pour toi aujourd’hui. C’était une idée pour dans 3–6 mois, "
        "comme une liste de courses future. Tu n’as RIEN à coder là-dessus pour livrer.",
        bold=True,
    )

    h(doc, "2. Critique honnête de l’app (marché)", 1)
    para(doc, "Forces", bold=True)
    para(doc, "• Différenciation réelle : dictée arabe → plaidoirie marocaine (peu de concurrents).")
    para(doc, "• Isolation multi-cabinet solide (dossiers / docs / plaidoiries).")
    para(doc, "• Rôles clairs admin / avocat / assistant.")
    para(doc, "• Essai 14 j + e-mails rétention type SaaS.")
    para(doc, "Faiblesses corrigées aujourd’hui", bold=True)
    para(doc, "• Auto-upgrade plan payant sans paiement → verrouillé si mock off.")
    para(doc, "• Compte désactivé qui se réactivait via reset MDP → corrigé.")
    para(doc, "• WebSocket acceptait refresh token → access only.")
    para(doc, "• OTP brute-force → limite 5 échecs / e-mail.")
    para(doc, "• Invite accept sans re-check sièges → corrigé.")
    para(doc, "• XSS facture HTML → échappé.")
    para(doc, "• Bouton « simuler échec paiement » retiré de l’UI client.")
    para(doc, "• Claim « hébergé Maroc » vs OVH France → reformulé honnête.")
    para(doc, "• CTAs tarifs alignés sur essai Cabinet 14 j.")
    para(doc, "Reste phase 2 (marché réel)", bold=True)
    para(doc, "• Paiement CMI/PayZone live.")
    para(doc, "• Clé OpenAI (qualité clé de barre).")
    para(doc, "• MFA obligatoire admin plateforme.")
    para(doc, "• Facture fiscale ICE/TVA.")

    h(doc, "3. Verdict", 1)
    para(
        doc,
        "Pour une démo / pilot WEBVOICE + PFA : l’app est finalisable et crédible. "
        "Pour « n°1 marché Maroc » commercial mass-market : il manque surtout le paiement réel "
        "et la preuve d’adoption (cabinets pilotes). Le produit cœur (plaidoirie) est le différenciateur — "
        "c’est le bon pari.",
    )

    doc.save(OUT)
    print(OUT, OUT.stat().st_size)


if __name__ == "__main__":
    main()
