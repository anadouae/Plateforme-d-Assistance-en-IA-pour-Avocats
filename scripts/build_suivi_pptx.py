#!/usr/bin/env python3
"""PowerPoint premier suivi PFA — Sanad (version corrigée)."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
MAQ = ROOT / "docs" / "maquettes"
DIAG = ROOT / "docs" / "diagrammes"
OUT = ROOT / "docs" / "Sanad_Premier_Suivi_Ecole.pptx"
DESK = Path.home() / "Desktop" / "Sanad_Premier_Suivi_Ecole.pptx"
DESK_V2 = Path.home() / "Desktop" / "Sanad_Premier_Suivi_PFA_v2.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x4A)
GOLD = RGBColor(0xD2, 0xB5, 0x89)
CREAM = RGBColor(0xF5, 0xF0, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x4A, 0x55, 0x68)
DARK = RGBColor(0x1A, 0x1A, 0x2E)


def set_run(run, size=18, bold=False, color=NAVY, font="Georgia"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_bar(slide, y=0, h=0.12, color=GOLD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), Inches(13.333), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def title_box(slide, text, size=28, color=NAVY):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.65))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    set_run(r, size=size, bold=True, color=color)


def bullets(slide, items, top=1.15, size=15, color=DARK, left=0.65, width=12):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)


def footer(slide, text="Sanad · Premier suivi PFA · ENSA Tétouan · WebVoice · Douae El Assal"):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    set_run(r, size=10, color=MUTED)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_image_fit(slide, path, left, top, max_w, max_h):
    if not Path(path).exists():
        return
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    ratio = w / h
    if ratio > max_w / max_h:
        dw, dh = max_w, max_w / ratio
    else:
        dh, dw = max_h, max_h * ratio
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(dw), Inches(dh))


def card(slide, left, top, w, h, fill=WHITE):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = GOLD
    return sh


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 TITRE
    s = blank(prs)
    add_bg(s, NAVY)
    add_bar(s, 0, 0.12, GOLD)
    add_bar(s, 7.38, 0.12, GOLD)
    b = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(1))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = "SANAD"
    set_run(r, size=54, bold=True, color=GOLD)
    b2 = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(1.2))
    r2 = b2.text_frame.paragraphs[0].add_run()
    r2.text = "Plateforme d’Assistance en Intelligence Artificielle pour Avocats"
    set_run(r2, size=22, color=WHITE)
    b3 = s.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(2))
    for i, line in enumerate(
        [
            "Premier suivi — Projet de Fin d’Année (PFA)",
            "ENSA Tétouan  ·  Entreprise d’accueil : WebVoice",
            "Stagiaire : Douae El Assal  ·  2025-2026",
        ]
    ):
        p = b3.text_frame.paragraphs[0] if i == 0 else b3.text_frame.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=16, color=CREAM)
        p.space_after = Pt(6)

    # 2 OBJECTIF
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Objectif de ce suivi")
    bullets(
        s,
        [
            "Présenter la problématique et l’idée Sanad de façon claire et convaincante",
            "Mettre en valeur le cœur produit : dicter → transcrire & corriger → modifier → PDF / DOC",
            "Présenter l’ensemble des fonctionnalités de la plateforme",
            "Synthétiser le benchmarking (marché marocain + aperçu international)",
            "Présenter les exigences et la première implémentation des interfaces",
            "Faire le point sur l’avancement réel et le serveur VPS ; valider la suite",
        ],
    )
    footer(s)

    # 3 PROBLEMATIQUE — reformulated (no "fichiers dispersés")
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Problématique")
    box = s.shapes.add_textbox(Inches(0.65), Inches(1.1), Inches(12), Inches(1.1))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = (
        "Même dans un cabinet organisé, le passage de la parole de l’avocat "
        "à un document de plaidoirie exploitable reste lent, manuel et peu intégré au dossier numérique."
    )
    set_run(r, size=17, bold=True, color=NAVY)
    bullets(
        s,
        [
            "La plaidoirie est souvent dictée ou notée, puis retravaillée / retapée hors d’un flux unique",
            "Les outils du marché couvrent soit la gestion de cabinet, soit une dictée isolée, soit une IA générale — rarement le parcours complet",
            "La collaboration avocat / assistant manque d’un espace commun avec droits clairs autour du même dossier",
            "Résultat : perte de temps sur des tâches techniques, au détriment du travail juridique à forte valeur",
            "Marché marocain (~13 000–16 000 avocats) : forte demande de digitalisation métier, encore inégale",
        ],
        top=2.4,
        size=14,
    )
    footer(s)

    # 4 SOLUTION
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Solution proposée — Sanad")
    bullets(
        s,
        [
            "Sanad = plateforme web SaaS multi-cabinets pour les cabinets d’avocats marocains",
            "Cœur produit : la plaidoirie dictée → transcrite & corrigée → modifiable → export PDF / DOC",
            "Un seul espace pour plaidoirie, dossiers, documents, synthèse, planning et équipe",
            "L’IA accélère transcription et correction — l’avocat reste maître du texte final",
            "Isolation par cabinet, rôles distincts (avocat / assistant / admin), traçabilité",
            "Message produit : « Deux gestes. Tout le reste suit. »",
        ],
        size=14,
    )
    footer(s)

    # 5 CŒUR PLAIDOIRIE — mise en valeur
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Cœur Sanad — parcours plaidoirie", size=26)
    intro = s.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(12), Inches(0.55))
    ir = intro.text_frame.paragraphs[0].add_run()
    ir.text = (
        "L’avocat dicte sa plaidoirie. Sanad la transcrit et la corrige. "
        "Il peut ensuite la modifier, puis l’exporterer en PDF ou DOC."
    )
    set_run(ir, size=15, bold=True, color=NAVY)

    steps = [
        ("1", "Dicter", "Parole / micro\nou saisie"),
        ("2", "Transcrire\n& corriger", "IA : texte propre\nsans inventer"),
        ("3", "Modifier", "Vérifier, ajouter,\néditer librement"),
        ("4", "Exporter", "PDF / Imprimer\nou fichier DOC"),
        ("5", "Classer", "Dans un dossier\nprécis du cabinet"),
    ]
    for i, (num, title, detail) in enumerate(steps):
        left = 0.45 + i * 2.55
        card(s, left, 1.85, 2.4, 3.55, NAVY if i % 2 == 0 else WHITE)
        tb = s.shapes.add_textbox(Inches(left + 0.15), Inches(2.05), Inches(2.1), Inches(3.2))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = num
        set_run(run, size=28, bold=True, color=GOLD if i % 2 == 0 else NAVY)
        p2 = tb.text_frame.add_paragraph()
        p2.space_before = Pt(10)
        run2 = p2.add_run()
        run2.text = title
        set_run(run2, size=16, bold=True, color=WHITE if i % 2 == 0 else NAVY)
        for j, line in enumerate(detail.split("\n")):
            p3 = tb.text_frame.add_paragraph()
            p3.space_before = Pt(6 if j == 0 else 2)
            run3 = p3.add_run()
            run3.text = line
            set_run(run3, size=12, color=CREAM if i % 2 == 0 else MUTED)

    foot_note = s.shapes.add_textbox(Inches(0.65), Inches(5.65), Inches(12), Inches(0.9))
    fn = foot_note.text_frame.paragraphs[0].add_run()
    fn.text = (
        "Point fort : ce n’est pas une dictée figée — l’avocat garde le contrôle "
        "(modification + export PDF/DOC + archivage dossier)."
    )
    set_run(fn, size=14, bold=True, color=DARK)
    footer(s)

    # 6 FONCTIONNALITES
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Fonctionnalités de la plateforme")
    pillars = [
        ("Plaidoirie (cœur)", "Dicter → corriger → modifier → PDF/DOC"),
        ("Dossiers", "Affaires, historique, rattachement"),
        ("Documents", "Upload, OCR, pièces du dossier"),
        ("Synthèse", "Résumé des pièces à la demande"),
        ("Planning", "Audiences & délais liés au dossier / plaidoirie"),
        ("Équipe", "Rôles avocat / assistant / admin"),
        ("Sécurité", "Multi-tenant, auth, chiffrement (en consolidation)"),
        ("Multilingue", "Interfaces en français, arabe et anglais"),
    ]
    for i, (t, d) in enumerate(pillars):
        col, row = i % 2, i // 2
        left, top = 0.55 + col * 6.35, 1.15 + row * 1.35
        fill = WHITE if i % 2 == 0 else NAVY
        card(s, left, top, 6.05, 1.2, fill)
        tb = s.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.22), Inches(5.6), Inches(0.85))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = t
        set_run(run, size=15, bold=True, color=NAVY if i % 2 == 0 else GOLD)
        p2 = tb.text_frame.add_paragraph()
        run2 = p2.add_run()
        run2.text = d
        set_run(run2, size=12, color=MUTED if i % 2 == 0 else CREAM)
    footer(s)

    # 7 EXIGENCES — clearer wording
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Exigences fonctionnelles (synthèse)")
    bullets(
        s,
        [
            "Accès & comptes : inscription cabinet, connexion sécurisée, rôles, invitations d’équipe",
            "Plaidoirie (priorité) : dictée → transcription & correction → modification par l’avocat → export PDF ou DOC → classement dossier",
            "Dossiers & documents : création/suivi des affaires, dépôt des pièces, OCR, synthèse",
            "Planning : audiences, délais et tâches rattachés aux dossiers (préparation de plaidoirie)",
            "Collaboration : avocat (dicter, modifier, exporter) / assistant (consulter, préparer pièces, exporter) / admin",
            "Sécurité & isolation : espaces séparés par cabinet (renforcement en cours)",
            "Expérience utilisateur : site public + espace cabinet connecté ; langues FR, AR, EN",
        ],
        size=14,
    )
    footer(s)

    # 8 BENCHMARK SCHEMA
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Benchmarking — marché marocain (schéma)", size=26)

    competitors = [
        ("SmartAvocat", "smartavocat.ma", "Gestion cabinet\nDossiers · bilingue", "Pas de flux\nplaidoirie vocale"),
        ("Avocato", "avocato.ma", "Gestion + Mahakim\nAudiences", "Pas dictée →\ndocument métier"),
        ("AvocatPro", "avocatpro.ma", "Gestion + Mahakim\nOrganisation", "Pas cœur\nplaidoirie"),
        ("Legaliq", "legaliq.ma", "IA juridique /\nfiscale · docs", "Pas SaaS cabinet\ncomplet"),
        ("Kanoun", "kanoun.com", "IA droit MA\n+ OHADA", "Pas multi-tenant\ncabinet"),
        ("Adala", "adala.ai", "Assistant IA\nQ&R · docs", "Pas équipe\ncabinet"),
        ("Dragon Legal", "diskod.ma", "Dictée desktop\n(Nuance)", "Pas SaaS ni\ndossier Sanad"),
    ]

    # header row labels
    hdr = s.shapes.add_textbox(Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.35))
    hr = hdr.text_frame.paragraphs[0].add_run()
    hr.text = "Plateforme (lien) → offre principale → limite face au besoin Sanad"
    set_run(hr, size=12, bold=True, color=MUTED)

    for i, (name, link, offer, limit) in enumerate(competitors):
        left = 0.35 + i * 1.85
        # name + link
        c1 = card(s, left, 1.4, 1.75, 1.55, NAVY)
        tb = s.shapes.add_textbox(Inches(left + 0.05), Inches(1.5), Inches(1.65), Inches(1.35))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = name
        set_run(run, size=10, bold=True, color=GOLD)
        p2 = tb.text_frame.add_paragraph()
        run2 = p2.add_run()
        run2.text = link
        set_run(run2, size=8, color=CREAM)
        # offer
        card(s, left, 3.1, 1.75, 1.5, WHITE)
        tb2 = s.shapes.add_textbox(Inches(left + 0.08), Inches(3.2), Inches(1.6), Inches(1.3))
        for j, line in enumerate(offer.split("\n")):
            p = tb2.text_frame.paragraphs[0] if j == 0 else tb2.text_frame.add_paragraph()
            run = p.add_run()
            run.text = line
            set_run(run, size=9, color=DARK)
        # limit
        card(s, left, 4.75, 1.75, 1.35, CREAM)
        tb3 = s.shapes.add_textbox(Inches(left + 0.08), Inches(4.85), Inches(1.6), Inches(1.15))
        for j, line in enumerate(limit.split("\n")):
            p = tb3.text_frame.paragraphs[0] if j == 0 else tb3.text_frame.add_paragraph()
            run = p.add_run()
            run.text = line
            set_run(run, size=9, color=MUTED)

    # legend
    leg = s.shapes.add_textbox(Inches(0.4), Inches(6.25), Inches(12.5), Inches(0.6))
    lr = leg.text_frame.paragraphs[0].add_run()
    lr.text = "Ligne 1 = plateforme + site  ·  Ligne 2 = services offerts  ·  Ligne 3 = ce qui manque par rapport à Sanad"
    set_run(lr, size=11, color=MUTED)
    footer(s)

    # 9 BENCHMARK RESULT / GAP clear
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Résultat du benchmark — où se place Sanad ?")
    bullets(
        s,
        [
            "Au Maroc aujourd’hui : soit des logiciels de gestion de cabinet (dossiers, Mahakim), soit des IA juridiques générales, soit une dictée desktop",
            "Ce qui manque : un SaaS qui relie dictée → transcription/correction → modification → PDF/DOC → dossier",
            "Sanad vise exactement cet espace : plateforme cabinet + cœur plaidoirie éditable & exportable + collaboration",
            "À l’international (aperçu) : dictée Word/OS ou LegalTech généralistes — rarement adaptés au cabinet marocain bilingue",
            "En une phrase : Sanad ne concurrence pas « une IA de plus », il comble le manque d’un flux métier intégré",
        ],
        size=14,
    )
    # Sanad box
    card(s, 0.65, 5.35, 12, 1.35, NAVY)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(5.55), Inches(11.5), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Sanad = Dicter → Corriger → Modifier → PDF/DOC + Dossiers + Équipe"
    set_run(run, size=16, bold=True, color=GOLD)
    p2 = tb.text_frame.add_paragraph()
    run2 = p2.add_run()
    run2.text = "Sources sites : smartavocat.ma · avocato.ma · avocatpro.ma · legaliq.ma · kanoun.com · adala.ai · diskod.ma"
    set_run(run2, size=11, color=CREAM)
    footer(s)

    # 10 ROLES avocat vs assistant
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Espaces distincts — Avocat & Assistant")
    card(s, 0.55, 1.2, 5.9, 5.2, NAVY)
    card(s, 6.85, 1.2, 5.9, 5.2, WHITE)
    for left, title, color, items in [
        (
            0.75,
            "Espace Avocat",
            GOLD,
            [
                "Dicte la plaidoirie (micro / texte)",
                "Voit le texte transcrit & corrigé",
                "Modifie, ajoute, vérifie librement",
                "Exporte en PDF ou DOC",
                "Classe dans un dossier précis",
            ],
        ),
        (
            7.05,
            "Espace Assistant",
            NAVY,
            [
                "Consulte les plaidoiries réalisées",
                "Peut imprimer / exporter PDF ou DOC",
                "Prépare et classe les pièces",
                "Suit échéances / planning du dossier",
                "Ne dicte pas à la place de l’avocat",
            ],
        ),
    ]:
        tb = s.shapes.add_textbox(Inches(left), Inches(1.45), Inches(5.4), Inches(4.7))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = title
        set_run(run, size=20, bold=True, color=color)
        for it in items:
            p2 = tb.text_frame.add_paragraph()
            p2.space_before = Pt(12)
            run2 = p2.add_run()
            run2.text = "•  " + it
            set_run(run2, size=14, color=CREAM if left < 3 else DARK)
    footer(s)

    # 11 UI public
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Première implémentation UI — site public", size=26)
    for (name, lab), left in zip(
        [("01-accueil.png", "Accueil"), ("02-connexion.png", "Connexion"), ("03-inscription.png", "Inscription")],
        [0.4, 4.55, 8.7],
    ):
        add_image_fit(s, MAQ / name, left, 1.15, 3.9, 5.0)
        lb = s.shapes.add_textbox(Inches(left), Inches(6.5), Inches(3.9), Inches(0.3))
        r = lb.text_frame.paragraphs[0].add_run()
        r.text = lab
        set_run(r, size=12, bold=True, color=NAVY)
        lb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer(s)

    # 12 UI app core
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "UI — cœur métier (plaidoirie mise en avant)", size=24)
    banner = s.shapes.add_textbox(Inches(0.4), Inches(0.95), Inches(12.5), Inches(0.35))
    br = banner.text_frame.paragraphs[0].add_run()
    br.text = "Écran Plaidoirie : Dicter → Transcrire & corriger → Modifier → Exporter PDF / DOC → Classer"
    set_run(br, size=12, bold=True, color=MUTED)
    for i, (name, lab) in enumerate(
        [
            ("10-tableau-de-bord.png", "Tableau de bord"),
            ("14-plaidoirie.png", "Plaidoirie · dicter / corriger / PDF-DOC"),
            ("11-dossiers.png", "Dossiers"),
            ("13-documents.png", "Documents"),
        ]
    ):
        left = 0.35 + i * 3.25
        add_image_fit(s, MAQ / name, left, 1.35, 3.1, 4.7)
        lb = s.shapes.add_textbox(Inches(left), Inches(6.55), Inches(3.1), Inches(0.35))
        r = lb.text_frame.paragraphs[0].add_run()
        r.text = lab
        set_run(r, size=10, bold=True, color=NAVY)
        lb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer(s)

    # 13 UI collab
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "UI — Planning (lié à la plaidoirie), Équipe, Paramètres", size=24)
    for i, (name, lab) in enumerate(
        [
            ("15-planning.png", "Planning · dossiers / échéances"),
            ("16-equipe.png", "Équipe · avatars & rôles"),
            ("17-parametres.png", "Paramètres cabinet"),
            ("06-securite.png", "Page Sécurité (public)"),
        ]
    ):
        left = 0.35 + i * 3.25
        add_image_fit(s, MAQ / name, left, 1.15, 3.1, 5.0)
        lb = s.shapes.add_textbox(Inches(left), Inches(6.55), Inches(3.1), Inches(0.3))
        r = lb.text_frame.paragraphs[0].add_run()
        r.text = lab
        set_run(r, size=10, bold=True, color=NAVY)
        lb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer(s)

    # 14 ARCHITECTURE (diagramme principal)
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Architecture principale", size=26)
    add_image_fit(s, DIAG / "01-architecture-principale.png", 0.55, 1.0, 12.2, 5.7)
    footer(s)

    # 15 Parcours + couches (schémas)
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Schémas — parcours plaidoirie & couches", size=24)
    add_image_fit(s, DIAG / "02-parcours-plaidoirie.png", 0.35, 1.05, 6.3, 5.5)
    add_image_fit(s, DIAG / "05-couches-techniques.png", 6.8, 1.05, 6.1, 5.5)
    footer(s)

    # 16 Modèle de données
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Base de données — vue d’ensemble", size=26)
    add_image_fit(s, DIAG / "04-modele-donnees.png", 0.55, 1.0, 12.2, 5.7)
    footer(s)

    # 17 Documentation livrée
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Documentation technique livrée")
    bullets(
        s,
        [
            "01-interfaces-application.md — description claire de tous les écrans",
            "02-document-conception.md — conception + diagrammes architecture / parcours / rôles",
            "03-base-de-donnees.md — tables, colonnes, MinIO, règles multi-tenant",
            "diagrammes/ — schémas PNG (architecture, plaidoirie, rôles, MCD, couches)",
            "maquettes/ — captures d’écran UI intégrées à cette présentation",
        ],
        size=15,
    )
    footer(s)

    # 18 AVANCEMENT — honest
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "État d’avancement (réaliste)")
    bullets(
        s,
        [
            "Statut : première implémentation des exigences principales — le projet n’est pas terminé",
            "En place : parcours plaidoirie (dictée, correction, édition, export PDF/DOC, classement dossier) + dossiers / docs / équipe",
            "En cours : tests, corrections, consolidation UX, fiabilisation transcription / correction",
            "À renforcer : sécurité avancée, hardening, qualité transcription, déploiement",
            "VPS Debian : clé SSH générée et transmise ; accès réseau encore en cours de validation avec l’entreprise",
            "Méthode : itérations avec l’encadrant entreprise — on ajuste au fur et à mesure",
        ],
        size=14,
    )
    footer(s)

    # 19 NEXT
    s = blank(prs)
    add_bg(s, CREAM)
    add_bar(s, 0, 0.1, NAVY)
    title_box(s, "Prochaines étapes")
    bullets(
        s,
        [
            "Fiabiliser le parcours : dicter → corriger → modifier → PDF/DOC → dossier",
            "Renforcer synthèse des pièces et lien dossier ↔ planning",
            "Déployer sur VPS dès que l’accès SSH est confirmé",
            "Préparer une démo courte du flux plaidoirie pour le prochain suivi",
            "Avancer documentation / rapport PFA",
            "Intégrer vos retours d’aujourd’hui",
        ],
    )
    footer(s)

    # 20 CONCLUSION
    s = blank(prs)
    add_bg(s, NAVY)
    add_bar(s, 0, 0.12, GOLD)
    add_bar(s, 7.38, 0.12, GOLD)
    b = s.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.6))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = "Conclusion"
    set_run(r, size=32, bold=True, color=GOLD)
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(3.6))
    for i, t in enumerate(
        [
            "Cœur Sanad : l’avocat dicte → le texte est transcrit et corrigé → il modifie → il exporte en PDF ou DOC.",
            "Ce n’est pas une IA isolée : c’est un flux métier intégré au dossier et à l’équipe.",
            "Le benchmark confirme ce manque sur le marché marocain — Sanad le vise explicitement.",
            "Première implémentation en place ; la consolidation et le déploiement continuent.",
        ]
    ):
        p = box2.text_frame.paragraphs[0] if i == 0 else box2.text_frame.add_paragraph()
        run = p.add_run()
        run.text = "•  " + t
        set_run(run, size=15, color=WHITE)
        p.space_after = Pt(12)
    box3 = s.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.7))
    r3 = box3.text_frame.paragraphs[0].add_run()
    r3.text = "Je suis prête à répondre à vos questions et à intégrer vos remarques."
    set_run(r3, size=16, bold=True, color=GOLD)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"OK → {OUT}")
    for target in (DESK_V2, DESK):
        try:
            prs.save(target)
            print(f"OK → {target}")
        except PermissionError:
            alt = target.with_name(target.stem + "_new.pptx")
            prs.save(alt)
            print(f"Locked → saved as {alt}")


if __name__ == "__main__":
    main()
