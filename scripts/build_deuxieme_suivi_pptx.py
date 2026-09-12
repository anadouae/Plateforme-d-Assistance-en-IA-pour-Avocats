#!/usr/bin/env python3
"""PPT 2e suivi PFA — professionnel, dense, non redondant."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[1]
MAQ = ROOT / "docs" / "maquettes"
DIAG = ROOT / "docs" / "diagrammes"
LOGO = ROOT / "assets" / "logo-sanad.png"
OUT = ROOT / "docs" / "Sanad_Deuxieme_Suivi_PFA.pptx"
DESK = Path.home() / "Desktop" / "Sanad_Deuxieme_Suivi_PFA.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x4A)
GOLD = RGBColor(0xC4, 0xA5, 0x6E)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x64, 0x75)
DARK = RGBColor(0x1E, 0x29, 0x3B)
SOFT = RGBColor(0xE8, 0xEE, 0xF5)
OK = RGBColor(0x1F, 0x6B, 0x4A)
TOTAL = 18


def set_run(run, size=16, bold=False, color=NAVY, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_bg(slide, color):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sp = sh._element
    slide.shapes._spTree.remove(sp)
    slide.shapes._spTree.insert(2, sp)


def accent_bar(slide):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(7.5)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = GOLD
    sh.line.fill.background()


def top_rule(slide):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.12), Inches(0), Inches(13.213), Inches(0.08)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    sh.line.fill.background()


def footer(slide, page):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.12), Inches(10.5), Inches(0.28))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = "Sanad · 2e suivi PFA · ENSA Tetouan · WebVoice · Douae El Assal"
    set_run(r, size=10, color=MUTED)
    num = slide.shapes.add_textbox(Inches(11.6), Inches(7.12), Inches(1.3), Inches(0.28))
    nr = num.text_frame.paragraphs[0].add_run()
    nr.text = f"{page} / {TOTAL}"
    set_run(nr, size=10, color=MUTED)
    num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.5))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    set_run(r, size=24, bold=True, color=NAVY, font="Georgia")
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.45), Inches(0.78), Inches(12.4), Inches(0.35))
        sr = sb.text_frame.paragraphs[0].add_run()
        sr.text = subtitle
        set_run(sr, size=12, color=MUTED)


def body_page(prs):
    s = blank(prs)
    add_bg(s, CREAM)
    accent_bar(s)
    top_rule(s)
    return s


def bullets(slide, items, left=0.5, top=1.2, width=12.2, size=14, color=DARK, gap=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = "▸  " + item
        set_run(run, size=size, color=color)


def card(slide, left, top, w, h, fill=WHITE, border=GOLD):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(1)
    return sh


def card_text(slide, left, top, w, h, title_t, lines, fill=WHITE):
    card(slide, left, top, w, h, fill)
    tb = slide.shapes.add_textbox(
        Inches(left + 0.18), Inches(top + 0.15), Inches(w - 0.35), Inches(h - 0.25)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title_t
    set_run(r, size=14, bold=True, color=NAVY if fill == WHITE else GOLD, font="Georgia")
    for line in lines:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = line
        set_run(r2, size=12, color=MUTED if fill == WHITE else CREAM)


def add_img(slide, path, left, top, max_w, max_h):
    p = Path(path)
    if not p.exists():
        return False
    try:
        from PIL import Image

        with Image.open(p) as im:
            w, h = im.size
        ratio = w / max(h, 1)
        if ratio > max_w / max_h:
            dw, dh = max_w, max_w / ratio
        else:
            dh, dw = max_h, max_h * ratio
        slide.shapes.add_picture(str(p), Inches(left), Inches(top), Inches(dw), Inches(dh))
        return True
    except Exception:
        slide.shapes.add_picture(str(p), Inches(left), Inches(top), Inches(max_w), Inches(max_h))
        return True


def draw_matrix(slide, rows, left=0.4, top=1.15, col_w=(3.2, 4.8, 4.5)):
    headers = ["Remarque (1er suivi)", "Reponse technique / produit", "Preuve dans Sanad"]
    x = left
    for i, h in enumerate(headers):
        card(slide, x, top, col_w[i] - 0.08, 0.42, NAVY, NAVY)
        tb = slide.shapes.add_textbox(
            Inches(x + 0.1), Inches(top + 0.08), Inches(col_w[i] - 0.25), Inches(0.3)
        )
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = h
        set_run(r, size=11, bold=True, color=GOLD)
        x += col_w[i]

    y = top + 0.5
    row_h = 0.78
    for idx, (a, b, c) in enumerate(rows):
        fill = WHITE if idx % 2 == 0 else SOFT
        x = left
        for i, text in enumerate((a, b, c)):
            card(
                slide,
                x,
                y,
                col_w[i] - 0.08,
                row_h - 0.06,
                fill,
                RGBColor(0xD0, 0xD6, 0xDE),
            )
            tb = slide.shapes.add_textbox(
                Inches(x + 0.1),
                Inches(y + 0.08),
                Inches(col_w[i] - 0.28),
                Inches(row_h - 0.2),
            )
            tf = tb.text_frame
            tf.word_wrap = True
            r = tf.paragraphs[0].add_run()
            r.text = text
            set_run(r, size=11, color=DARK)
            x += col_w[i]
        y += row_h


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 TITLE
    s = blank(prs)
    add_bg(s, NAVY)
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.9), Inches(13.333), Inches(0.6)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = GOLD
    sh.line.fill.background()
    if LOGO.exists():
        try:
            s.shapes.add_picture(str(LOGO), Inches(0.7), Inches(0.55), Inches(1.1), Inches(1.1))
        except Exception:
            pass
    b = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(11.8), Inches(0.9))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = "SANAD"
    set_run(r, size=56, bold=True, color=GOLD, font="Georgia")
    b2 = s.shapes.add_textbox(Inches(0.7), Inches(2.95), Inches(11.8), Inches(1.0))
    r2 = b2.text_frame.paragraphs[0].add_run()
    r2.text = "Deuxieme suivi PFA — Retours du premier suivi & livraisons"
    set_run(r2, size=22, color=WHITE, font="Georgia")
    b3 = s.shapes.add_textbox(Inches(0.7), Inches(4.3), Inches(11.8), Inches(1.8))
    for i, line in enumerate(
        [
            "Plateforme d'assistance IA pour cabinets d'avocats marocains",
            "ENSA Tetouan  ·  Entreprise d'accueil : WebVoice",
            "Stagiaire : Douae El Assal  ·  Aout 2026",
        ]
    ):
        p = b3.text_frame.paragraphs[0] if i == 0 else b3.text_frame.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = line
        set_run(run, size=15, color=CREAM)
    ft = s.shapes.add_textbox(Inches(0.7), Inches(7.02), Inches(12), Inches(0.35))
    fr = ft.text_frame.paragraphs[0].add_run()
    fr.text = "Structure : remarque → action → preuve  ·  une idee par slide"
    set_run(fr, size=11, bold=True, color=NAVY)

    # 2 OBJECTIF
    s = body_page(prs)
    title(
        s,
        "Objectif de ce deuxieme suivi",
        "Ce que nous validons aujourd'hui — pas une reprise du cadrage initial",
    )
    card_text(
        s,
        0.45,
        1.25,
        6.0,
        5.2,
        "Au programme",
        [
            "1. Montrer comment chaque remarque du 1er suivi a ete traitee",
            "2. Presenter les livraisons concretes (fonctionnel + technique)",
            "3. Demontrer le parcours plaidoirie arabe (coeur PFA)",
            "4. Expliquer l'architecture resiliente (STT / LLM / Docker)",
            "5. Positionner Sanad via une grille benchmark OK / Non OK",
            "6. Fixer les priorites jusqu'au prochain jalon",
        ],
    )
    card_text(
        s,
        6.7,
        1.25,
        6.1,
        5.2,
        "Hors scope de cette seance",
        [
            "Ne pas re-presenter toute la problematique marche",
            "Ne pas repeter le meme parcours sur 5 slides",
            "Ne pas empiler des listes de features identiques",
            "Focus : remarque → action → preuve",
        ],
        fill=NAVY,
    )
    footer(s, 2)

    # 3 ANCRAGE
    s = body_page(prs)
    title(
        s,
        "Ancrage — coeur produit (rappel unique)",
        "Valide au 1er suivi · fil rouge sans repetition ulterieure",
    )
    card(s, 0.8, 2.0, 11.7, 2.8, NAVY, GOLD)
    tb = s.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.9), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (
        "Transformer la voix de l'avocat (arabe) en plaidoirie structuree, "
        "modifiable et exportable (Word / PDF), classee dans le dossier du cabinet — "
        "sans dependance unique a un outil cloud."
    )
    set_run(r, size=20, bold=True, color=WHITE, font="Georgia")
    note = s.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.2))
    nr = note.text_frame.paragraphs[0].add_run()
    nr.text = (
        "Les modules cabinet (dossiers, documents, agenda, equipe) soutiennent ce flux ; "
        "ils ne le remplacent pas."
    )
    set_run(nr, size=14, color=MUTED)
    footer(s, 3)

    # 4 MATRICE
    s = body_page(prs)
    title(
        s,
        "Matrice des remarques — reponses & preuves",
        "Slide centrale : une ligne = une remarque (pas de doublon ensuite)",
    )
    draw_matrix(
        s,
        [
            (
                "STT : dependance fragile ?",
                "Whisper local + upload audio + Chrome secours + saisie manuelle",
                "API /dictation/transcribe · segments 55 s",
            ),
            (
                "LLM : risque fournisseur ?",
                "Multi-fournisseurs + Ollama local + mode degrade (edition/export)",
                "polish / compose / synthese avec fallback",
            ),
            (
                "Agenda ↔ affaires",
                "Evenements, taches et notes rattaches a un case_id",
                "Module Agenda + filtre par dossier",
            ),
            (
                "Domaines + recherche",
                "Domaine a la creation + filtre / recherche liste",
                "Affaires · case_type (civil, penal…)",
            ),
            (
                "Benchmark OK / Non OK",
                "Grille criteres cochables centree voix→dossier",
                "docs/benchmark-ok-nok.md",
            ),
            (
                "Presentation redondante",
                "Structure remarque→preuve ; 1 idee / slide",
                "Ce support",
            ),
            (
                "Gestion cabinet",
                "Auth, roles, equipe, templates, multi-tenant",
                "Login / Team / Admin",
            ),
        ],
        top=1.05,
        col_w=(3.1, 5.0, 4.5),
    )
    footer(s, 4)

    # 5 STT
    s = body_page(prs)
    title(
        s,
        "Transcription vocale — architecture de resilience",
        "Reponse detaillee a la question STT (une seule fois)",
    )
    card_text(
        s,
        0.45,
        1.2,
        4.0,
        5.3,
        "Niveaux de secours",
        [
            "A. Whisper local (faster-whisper) — arabe force",
            "B. Upload fichier audio (mp3 / wav / m4a / webm)",
            "C. Chrome Web Speech (secours navigateur)",
            "D. Saisie / collage manuel — toujours possible",
            "",
            "Longue dictee : decoupe ~55 s → texte progressif",
        ],
    )
    card_text(
        s,
        4.65,
        1.2,
        4.0,
        5.3,
        "Choix produit Maroc",
        [
            "Langue STT = arabe (avocats marocains)",
            "FR UI optionnel ; EN retire du switch",
            "Prompt juridique Whisper + repair arabe",
            "Badge « Whisper pret » = sante STT",
            "",
            "Phrase cle : accelerateur, jamais blocage unique.",
        ],
        fill=NAVY,
    )
    card_text(
        s,
        8.85,
        1.2,
        4.0,
        5.3,
        "Limite assumee",
        [
            "Modele small sur CPU = compromis qualite / cout",
            "Ameliorable (modele plus grand / GPU)",
            "OCR de mauvaise qualite = risque separe (scans)",
        ],
    )
    footer(s, 5)

    # 6 LLM + schema Whisper vs Ollama (pedagogique)
    s = body_page(prs)
    title(
        s,
        "Whisper ≠ Ollama — deux machines differentes",
        "A retenir pour l'oral : oreilles vs plume",
    )
    card_text(
        s,
        0.45,
        1.15,
        6.1,
        4.4,
        "WHISPER (local) = oreilles",
        [
            "Entree : AUDIO (micro ou fichier)",
            "Sortie : TEXTE brut arabe",
            "Ou : conteneur backend (faster-whisper)",
            "Ne corrige PAS le droit",
            "Ne fait PAS la synthese PDF",
        ],
    )
    card_text(
        s,
        6.8,
        1.15,
        6.0,
        4.4,
        "OLLAMA (local) = plume",
        [
            "Entree : TEXTE deja ecrit",
            "Sortie : texte corrige / structure",
            "Usages : polish, compose, synthese",
            "Secours cloud : Groq / Kimi…",
            "Ne lit PAS l'audio directement",
        ],
        fill=NAVY,
    )
    note = s.shapes.add_textbox(Inches(0.45), Inches(5.75), Inches(12.4), Inches(1.0))
    nr = note.text_frame.paragraphs[0].add_run()
    nr.text = (
        "Chaine : Audio → Whisper → texte brut → Ollama/LLM → document. "
        "Si LLM tombe : edition + export Word/PDF restent possibles (mode degrade)."
    )
    set_run(nr, size=13, bold=True, color=DARK)
    footer(s, 6)

    # 7 UX
    s = body_page(prs)
    title(s, "UX plaidoirie — parcours en deux ecrans", "Reponse a la saturation UI · preuve visuelle")
    bullets(
        s,
        [
            "Ecran 1 — /redaction : choix depart (vide ou racine modele)",
            "Ecran 2 — /redaction/[id] : dictee arabe + document officiel",
            "Moins de textes d'aide : interface professionnelle",
            "Login : email + mot de passe (lien magique discret)",
        ],
        left=0.45,
        top=1.2,
        width=5.8,
        size=14,
        gap=10,
    )
    # Capture actuelle (workspace) — fallback ancien fichier si absent
    if not add_img(s, MAQ / "14b-plaidoirie-workspace.png", 6.5, 1.15, 6.3, 5.5):
        add_img(s, MAQ / "14-plaidoirie.png", 6.5, 1.15, 6.3, 5.5)
    footer(s, 7)

    # 8 DOMAINES
    s = body_page(prs)
    title(s, "Dossiers par domaine + recherche", "Reponse directe a la remarque classement / filtre")
    bullets(
        s,
        [
            "A la creation : choix du domaine (civil, penal, commercial, famille…)",
            "Sur la liste : filtre domaine + recherche texte",
            "Interet metier : plaidoiries souvent similaires par domaine → reutilisation",
        ],
        left=0.45,
        top=1.2,
        width=5.8,
        size=14,
        gap=12,
    )
    add_img(s, MAQ / "11-dossiers.png", 6.5, 1.15, 6.3, 5.5)
    footer(s, 8)

    # 9 AGENDA
    s = body_page(prs)
    title(s, "Agenda lie aux affaires", "Reponse directe a la remarque sync planning ↔ dossiers")
    bullets(
        s,
        [
            "Evenements, taches et notes portent un case_id optionnel",
            "Filtre agenda par dossier + lien vers l'affaire",
            "Usage : audience / delai / preparation plaidoirie rattaches au dossier",
        ],
        left=0.45,
        top=1.2,
        width=5.8,
        size=14,
        gap=12,
    )
    add_img(s, MAQ / "15-planning.png", 6.5, 1.15, 6.3, 5.5)
    footer(s, 9)

    # 10 SYNTHESE
    s = body_page(prs)
    title(s, "Synthese des pieces", "Livraison cabinet autour du coeur plaidoirie")
    bullets(
        s,
        [
            "Upload → OCR / extraction → synthese structuree",
            "Blocs : parties, dates, montants, risques + resume",
            "Accessible depuis la fiche document (avocat / admin)",
            "Assistant : consultation limitee selon role",
        ],
        left=0.45,
        top=1.2,
        width=5.8,
        size=14,
        gap=12,
    )
    add_img(s, MAQ / "13-documents.png", 6.5, 1.15, 6.3, 5.5)
    footer(s, 10)

    # 11 ARCHITECTURE
    s = body_page(prs)
    title(s, "Architecture technique", "Vue d'ensemble unique — pas de re-listing features")
    if not add_img(s, DIAG / "sanad-architecture-systeme.png", 0.4, 1.1, 7.4, 5.5):
        add_img(s, DIAG / "01-architecture-principale.png", 0.4, 1.1, 7.4, 5.5)
    card_text(
        s,
        8.0,
        1.2,
        4.8,
        5.3,
        "Couches",
        [
            "Frontend : Next.js (RTL / arabe)",
            "API : FastAPI + JWT",
            "Donnees : PostgreSQL (pas SQLite prod)",
            "Jobs : Celery + Redis",
            "Fichiers : MinIO + chiffrement",
            "STT : Whisper dans backend",
            "LLM : Ollama (+ fallbacks)",
            "Multi-tenant par cabinet",
        ],
        fill=NAVY,
    )
    footer(s, 11)

    # 12 APIs
    s = body_page(prs)
    title(s, "APIs metier exposees", "Reference technique courte pour questions orales")
    groups = [
        ("Auth", "/api/auth/* — login, OTP email, magic link, refresh, Google"),
        ("Dictee", "/api/dictation/transcribe · polish · compose · stt-status"),
        ("Dossiers", "/api/cases — CRUD + domaine"),
        ("Documents", "/api/documents · /synthese · ask"),
        ("Workspace", "/api/workspace — agenda / taches / notes"),
        ("Cabinet", "/api/team · /plaidoirie-templates · /settings"),
        ("Admin", "/api/admin — supervision WebVoice"),
    ]
    for i, (t, d) in enumerate(groups):
        col, row = i % 2, i // 2
        left, top = 0.45 + col * 6.4, 1.2 + row * 1.15
        if row > 3:
            break
        card(s, left, top, 6.15, 1.0, WHITE)
        tb = s.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.18), Inches(5.75), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = t
        set_run(r, size=13, bold=True, color=NAVY, font="Georgia")
        p2 = tb.text_frame.add_paragraph()
        r2 = p2.add_run()
        r2.text = d
        set_run(r2, size=11, color=MUTED)
    footer(s, 12)

    # 13 DOCKER
    s = body_page(prs)
    title(s, "Images Docker — role & publication", "Pourquoi ce n'est pas « juste du packaging »")
    imgs = [
        ("sanad-backend", "API FastAPI + Whisper + OCR", "Coeur metier & STT"),
        ("sanad-frontend", "Next.js build production", "Interface avocat / admin"),
        ("sanad-celery", "Workers asynchrones", "OCR, indexation pieces"),
        ("postgres / redis / minio", "Images officielles", "Donnees, queue, fichiers"),
        ("mailhog / ollama", "Dev & LLM local", "Emails test + modeles"),
    ]
    for i, (name, what, why) in enumerate(imgs):
        y = 1.15 + i * 0.95
        card(s, 0.45, y, 12.4, 0.85, WHITE if i % 2 == 0 else SOFT)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(y + 0.15), Inches(11.9), Inches(0.6))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = name
        set_run(r, size=14, bold=True, color=NAVY)
        p2 = tb.text_frame.add_paragraph()
        r2 = p2.add_run()
        r2.text = f"{what}  ·  {why}"
        set_run(r2, size=12, color=MUTED)
    note = s.shapes.add_textbox(Inches(0.45), Inches(6.05), Inches(12.4), Inches(0.8))
    nr = note.text_frame.paragraphs[0].add_run()
    nr.text = (
        "Publication Docker Hub (anadouae1/sanad-*) : meme binaire sur PC / VPS / demo — "
        "reproductibilite, deploiement rapide, preuve de maturite DevOps. "
        "Code source : GitHub Plateforme-d-Assistance-en-IA-pour-Avocats."
    )
    set_run(nr, size=12, color=DARK)
    footer(s, 13)

    # 14 DEPLOY
    s = body_page(prs)
    title(s, "Deploiement Docker Compose", "Schema — infra reelle du projet")
    add_img(s, DIAG / "sanad-deploiement.png", 1.5, 1.1, 10.3, 5.5)
    footer(s, 14)

    # 15 BENCHMARK
    s = body_page(prs)
    title(
        s,
        "Benchmark — grille OK / Non OK",
        "Format demande au 1er suivi · criteres centres coeur produit",
    )
    bullets(
        s,
        [
            "Criteres : dictee→texte, correction/edition, export PDF/DOC, classement dossier, roles, domaines, agenda lie, mode degrade",
            "Gestionnaires (SmartAvocat, Avocato…) : OK dossiers — Non OK flux plaidoirie integre",
            "IA juridiques (Legaliq, Kanoun…) : utiles sur le fond — Non OK SaaS cabinet complet",
            "Dragon Legal : OK dictee desktop — Non OK parcours web jusqu'au dossier",
            "Sanad : OK sur le flux integre voix → document → dossier → equipe",
            "Detail tabule : docs/benchmark-ok-nok.md (a projeter si question)",
        ],
        top=1.25,
        size=14,
        gap=10,
    )
    footer(s, 15)

    # 16 BILAN
    s = body_page(prs)
    title(
        s,
        "Bilan des livraisons depuis le 1er suivi",
        "Checklist unique — pas de reprise des slides precedentes",
    )
    left_items = [
        "Whisper arabe + fichiers audio + segments longs",
        "UX plaidoirie 2 ecrans + templates cabinet",
        "Domaines dossiers + filtre/recherche",
        "Agenda / taches / notes ↔ case_id",
    ]
    right_items = [
        "Synthese pieces (prompt anti-invention)",
        "Auth pro (OTP, magic link, roles)",
        "UI arabe defaut / FR optionnel",
        "GitHub + images Docker Hub publiees",
    ]
    for i, t in enumerate(left_items):
        card(s, 0.45, 1.25 + i * 1.2, 6.1, 1.05, WHITE)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.5 + i * 1.2), Inches(5.6), Inches(0.6))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = "✓  " + t
        set_run(r, size=14, bold=True, color=OK)
    for i, t in enumerate(right_items):
        card(s, 6.8, 1.25 + i * 1.2, 6.1, 1.05, WHITE)
        tb = s.shapes.add_textbox(Inches(7.05), Inches(1.5 + i * 1.2), Inches(5.6), Inches(0.6))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = "✓  " + t
        set_run(r, size=14, bold=True, color=OK)
    footer(s, 16)

    # 17 DEMO + SUITE
    s = body_page(prs)
    title(s, "Demo (5–7 min) & suite", "Scenario unique + prochaines priorites")
    card_text(
        s,
        0.45,
        1.2,
        6.2,
        5.3,
        "Scenario de demo",
        [
            "1. Login demo@sanad.ma",
            "2. Affaire + domaine + filtre",
            "3. Plaidoirie : modele → audio arabe → corriger → Word/PDF",
            "4. Piece → synthese",
            "5. Agenda lie au dossier",
            "6. (option) assistant / admin WebVoice",
        ],
    )
    card_text(
        s,
        6.9,
        1.2,
        5.9,
        5.3,
        "Suite jusqu'au prochain jalon",
        [
            "Qualite STT juridique (modele / glossaire)",
            "OCR scans difficiles",
            "SMTP production (Resend / SES)",
            "Durcissement VPS / monitoring",
            "Veille analyse globale dossier = backlog",
        ],
        fill=NAVY,
    )
    footer(s, 17)

    # 18 CLOSE
    s = blank(prs)
    add_bg(s, NAVY)
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = GOLD
    sh.line.fill.background()
    b = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.2))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = "Merci — Questions & discussion"
    set_run(r, size=36, bold=True, color=GOLD, font="Georgia")
    b2 = s.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(11.7), Inches(2.0))
    for i, line in enumerate(
        [
            "Message final : chaque remarque du 1er suivi a une reponse d'architecture",
            "ou de produit — pas seulement un nouvel ecran.",
            "",
            "Douae El Assal  ·  ENSA Tetouan  ·  WebVoice  ·  Sanad",
        ]
    ):
        p = b2.text_frame.paragraphs[0] if i == 0 else b2.text_frame.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=15, color=CREAM)
    footer_box = s.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.3))
    fr = footer_box.text_frame.paragraphs[0].add_run()
    fr.text = f"{TOTAL} / {TOTAL}"
    set_run(fr, size=11, color=GOLD)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    try:
        prs.save(str(DESK))
        print(f"Wrote {OUT}")
        print(f"Also {DESK}")
    except Exception as e:
        print(f"Wrote {OUT} (desktop copy failed: {e})")


if __name__ == "__main__":
    main()
