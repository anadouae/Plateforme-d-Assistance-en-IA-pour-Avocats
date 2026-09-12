import logging
import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz
from docx import Document as DocxDocument
from PIL import Image

logger = logging.getLogger(__name__)

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".tif",
    ".webm",
    ".wav",
    ".mp3",
    ".m4a",
    ".ogg",
    ".mp4",
    ".mpeg",
    ".mpga",
}
AUDIO_EXTENSIONS = {".webm", ".wav", ".mp3", ".m4a", ".ogg", ".mp4", ".mpeg", ".mpga"}
ALLOWED_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "text/txt",
    "text/markdown",
    "text/x-markdown",
    "image/png",
    "image/jpeg",
    "image/tiff",
    "audio/webm",
    "audio/wav",
    "audio/x-wav",
    "audio/wave",
    "audio/mpeg",
    "audio/mp3",
    "audio/mp4",
    "audio/x-m4a",
    "audio/m4a",
    "audio/ogg",
    "audio/aac",
    "video/webm",  # MediaRecorder / certains navigateurs
    "video/mp4",
}

OCR_LANG = "fra+ara"


def _ocr_settings():
    from app.core.config import settings

    return settings


def _max_pdf_ocr_pages() -> int:
    return int(_ocr_settings().MAX_PDF_OCR_PAGES or 80)


def _max_tesseract_pages() -> int:
    return int(_ocr_settings().MAX_PDF_TESSERACT_PAGES or 40)


def _ocr_dpi() -> int:
    return int(_ocr_settings().OCR_DPI or 300)


def _ocr_quality_ok(text: str, *, min_len: int = 40) -> bool:
    """Résultat OCR assez complet pour synthétiser (qualité, pas un aperçu)."""
    if not text or len(text.strip()) < min_len:
        return False
    return not is_scrambled_arabic(text)

_ARABIC_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]")
# Mots / particules souvent en DÉBUT de phrase arabe (ordre logique)
_AR_STARTERS = {
    "في", "من", "على", "إلى", "الى", "إن", "أن", "هذا", "هذه", "ذلك", "تلك",
    "وقد", "فقد", "تحت", "بعد", "قبل", "مع", "عن", "بين", "كل", "أي", "لا",
    "ما", "هل", "تم", "يتم", "يمكن", "حيث", "إذ", "اذا", "إذا", "لكن", "غير",
    "مقدمة", "المبحث", "المطلب", "الفقرة", "خاتمة", "أولا", "ثانيا", "ثالثا",
    "ملخص", "بحث", "الفصل", "الباب", "المادة", "القانون", "قانون",
}
# Mots souvent en FIN visuelle quand la ligne est inversée (donc début logique)
_AR_TRAILERS_WHEN_REVERSED = {
    "في", "من", "على", "إلى", "الى", "و", "أو", "ثم",
}


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_audio_document(filename: str, mime_type: str = "") -> bool:
    """Pièce audio archivée (extension dictée Whisper) — distinct de PDF/image."""
    ext = get_extension(filename or "")
    if ext in AUDIO_EXTENSIONS:
        return True
    mt = (mime_type or "").lower().strip()
    if mt.startswith("audio/"):
        return True
    # Certains navigateurs envoient video/webm|mp4 pour de l'audio encapsulé
    if mt in {"video/webm", "video/mp4"} and ext in AUDIO_EXTENSIONS:
        return True
    return False


def document_kind(filename: str, mime_type: str = "") -> str:
    return "audio" if is_audio_document(filename, mime_type) else "file"


def arabic_char_ratio(text: str) -> float:
    if not text:
        return 0.0
    letters = [c for c in text if c.isalpha() or "\u0600" <= c <= "\u06FF"]
    if not letters:
        return 0.0
    ar = sum(1 for c in letters if _ARABIC_RE.match(c))
    return ar / len(letters)


def _strip_word(w: str) -> str:
    return re.sub(r"^[^\w\u0600-\u06FF]+|[^\w\u0600-\u06FF]+$", "", w or "")


_AR_VERBS_START = {
    "يقوم", "يتم", "يمكن", "يعتمد", "يهدف", "يشهد", "ساهم", "يساهم", "يتعلق",
    "تنص", "ينص", "تبرز", "يبرز", "توجد", "يوجد", "تشهد", "واكب", "اعتمد",
}


def _line_reading_score(words: list[str]) -> int:
    """Score heuristique : ordre logique arabe vs ordre visuel inversé."""
    if not words:
        return 0
    clean = [_strip_word(w) for w in words]
    clean = [w for w in clean if w]
    if not clean:
        return 0
    n = len(clean)
    score = 0
    first = clean[0]
    last = clean[-1]
    if first in _AR_STARTERS or first in _AR_VERBS_START:
        score += 4
    if last in _AR_STARTERS and first not in _AR_STARTERS:
        score -= 3
    if last in _AR_TRAILERS_WHEN_REVERSED and n >= 3:
        score -= 2
    # Prépositions / verbes mieux en début de ligne qu'en fin
    for i, w in enumerate(clean):
        pos = i / max(n - 1, 1)  # 0 = début, 1 = fin
        if w in ("في", "من", "على", "إلى", "الى", "مع", "عن", "بين"):
            score += 2 if pos < 0.35 else (-2 if pos > 0.65 else 0)
        if w in _AR_VERBS_START:
            score += 3 if pos < 0.35 else (-3 if pos > 0.65 else 0)
    for a, b in zip(clean, clean[1:]):
        if a in ("في", "من", "على", "إلى", "الى", "مع", "عن", "بين") and len(b) > 2:
            score += 1
        if b in ("في", "من", "على") and len(a) > 3:
            score -= 1
    return score


def repair_arabic_line(line: str) -> str:
    """Corrige l'ordre des mots arabes souvent inversé par PyMuPDF / OCR LTR."""
    raw = (line or "").strip()
    if not raw or arabic_char_ratio(raw) < 0.35:
        return line
    words = raw.split()
    if len(words) < 3:
        return line
    fwd = _line_reading_score(words)
    rev_words = list(reversed(words))
    rev = _line_reading_score(rev_words)
    if rev > fwd:
        return " ".join(rev_words)
    return line


def repair_arabic_extracted_text(text: str) -> str:
    """Répare le texte extrait (lignes arabes mot-inversées). Sûr pour FR/EN."""
    if not text or arabic_char_ratio(text) < 0.12:
        return text
    parts: list[str] = []
    for line in text.splitlines():
        if line.strip():
            parts.append(repair_arabic_line(line))
        else:
            parts.append(line)
    return "\n".join(parts)


def safe_repair_arabic_text(text: str) -> str:
    """Applique la réparation seulement si le score de lecture s'améliore."""
    if not text:
        return text
    fixed = repair_arabic_extracted_text(text)
    if fixed == text:
        return text
    if _arabic_quality_score(fixed) >= _arabic_quality_score(text):
        return fixed
    return text


def is_scrambled_arabic(text: str) -> bool:
    """
    Détecte une couche PDF/OCR arabe illisible (BIDI cassé / mots coupés).
    Doit rester CONSERVATEUR : un texte Word juridique fluide ne doit PAS être flaggé.
    """
    if not text or arabic_char_ratio(text) < 0.2:
        return False

    # Fragments typiques des PDF arabes mal encodés (éviter les vrais mots: عقارية، قضائية…)
    if re.search(
        r"(العقا\b|المعامالت|المعامﻼت|صصُخ|زّويمي|قّالموث|الّوفع)",
        text,
    ):
        return True

    words = [_strip_word(w) for w in text.split()]
    words = [w for w in words if w]
    if len(words) >= 25:
        # Trop de « mots » d'une lettre = OCR cassé
        tiny = sum(1 for w in words if len(w) <= 1)
        if tiny / len(words) >= 0.22:
            return True

    lines = [ln.strip() for ln in text.splitlines() if len(ln.split()) >= 4]
    if lines:
        bad_start = sum(1 for ln in lines if ln[:1] in "،.؛!?ًٌٍَُِّْ")
        if bad_start / len(lines) >= 0.35:
            return True
        legal = sum(
            1
            for m in (
                "القانون",
                "الظهير",
                "مدونة",
                "المحكمة",
                "المقاولة",
                "المادة",
                "الفصل",
                "المشرع",
                "المسطرة",
                "التنفيذ",
            )
            if m in text
        )
        # Long mémoire / PDF juridique majoritairement lisible → ne pas crier au BIDI
        if len(words) >= 200 and arabic_char_ratio(text) >= 0.4 and legal >= 2:
            return False
        # Si inverser les mots améliore FORTEMENT le score → texte inversé
        orig = sum(_line_reading_score(ln.split()) for ln in lines)
        rev = sum(
            _line_reading_score(list(reversed(ln.split()))) for ln in lines
        )
        if rev > orig + max(10, len(lines) * 3):
            # Sur un long document, exiger un écart plus net (quelques lignes inversées ≠ tout le PDF)
            if len(text) >= 12000 and rev < orig * 1.35:
                return False
            return True
        if len(words) >= 40 and arabic_char_ratio(text) >= 0.45 and legal >= 1:
            return False
        return False

    # Texte court sans structure de lignes : seuil bas uniquement
    return len(words) < 30 and _arabic_quality_score(text) < 5


def looks_like_native_arabic_prose(text: str) -> bool:
    """Heuristique positive : prose arabe cohérente (Word / texte propre)."""
    if not text or arabic_char_ratio(text) < 0.4:
        return False
    words = text.split()
    if len(words) < 25:
        return False
    if is_scrambled_arabic(text):
        return False
    legal = sum(
        1
        for m in ("القانون", "الظهير", "مدونة", "المحكمة", "المقاولة", "المادة", "الفصل")
        if m in text
    )
    return legal >= 1 or len(words) >= 60


def _arabic_quality_score(text: str) -> int:
    """Plus le score est haut, plus l'ordre de lecture semble correct."""
    if not text:
        return 0
    score = 0
    for line in text.splitlines():
        words = line.split()
        if len(words) >= 3 and arabic_char_ratio(line) > 0.35:
            score += _line_reading_score(words)
    # Bonus si des titres arabes classiques apparaissent au début d'une ligne
    for marker in ("مقدمة", "المبحث", "ملخص", "خاتمة", "المطلب"):
        if re.search(rf"(?m)^{marker}", text):
            score += 5
        if re.search(rf"(?m){marker}$", text):
            score -= 3
    return score


def _page_text_reading_order(page) -> str:
    """Extraction dict + tri RTL pour lignes arabes (évite get_text() LTR)."""
    try:
        data = page.get_text("dict")
    except Exception:
        return page.get_text() or ""

    lines_out: list[str] = []
    blocks = sorted(
        data.get("blocks", []),
        key=lambda b: (round(b.get("bbox", [0, 0, 0, 0])[1], 1), b.get("bbox", [0, 0, 0, 0])[0]),
    )
    for block in blocks:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans") or []
            if not spans:
                continue
            span_bits = []
            for s in spans:
                t = s.get("text") or ""
                bbox = s.get("bbox") or [0, 0, 0, 0]
                span_bits.append((float(bbox[0]), t))
            probe = "".join(t for _, t in span_bits)
            if arabic_char_ratio(probe) > 0.3:
                span_bits.sort(key=lambda x: x[0], reverse=True)
            else:
                span_bits.sort(key=lambda x: x[0])
            line_text = "".join(t for _, t in span_bits)
            # Normaliser espaces
            line_text = re.sub(r"[ \t]+", " ", line_text).strip()
            if line_text:
                lines_out.append(repair_arabic_line(line_text))
    return "\n".join(lines_out)


def _ocr_image(image: Image.Image, *, fast: bool = False) -> tuple[str, int]:
    import pytesseract

    if fast:
        configs = [("ara+fra", "--psm 6")]
    else:
        configs = [
            ("ara+fra", "--psm 6"),
            ("fra+ara", "--psm 6"),
            (OCR_LANG, "--psm 4"),
        ]
    best_text = ""
    best_score = -10**9
    for lang, cfg in configs:
        try:
            text = pytesseract.image_to_string(image, lang=lang, config=cfg)
            cleaned = repair_arabic_extracted_text(text.strip())
            q = _arabic_quality_score(cleaned) + len(cleaned) // 40
            if q > best_score and cleaned:
                best_score = q
                best_text = cleaned
        except Exception as exc:
            logger.debug("OCR lang=%s failed: %s", lang, exc)
    if not best_text:
        try:
            best_text = pytesseract.image_to_string(image, lang=OCR_LANG).strip()
            best_text = repair_arabic_extracted_text(best_text)
        except Exception as exc:
            logger.warning("OCR failed: %s", exc)
            return "", 35
    score = min(95, max(35, int(len(best_text) / 8))) if best_text else 35
    return best_text, score


def _ocr_image_timed(image: Image.Image, *, fast: bool = False) -> tuple[str, int]:
    from app.core.config import settings
    from app.core.timeouts import run_with_timeout

    per_page = min(90.0, float(settings.OCR_TESSERACT_TIMEOUT_SEC or 480) / 4)
    return run_with_timeout(
        lambda: _ocr_image(image, fast=fast),
        per_page,
        label="ocr_image",
    )


def _ocr_pdf_pymupdf(content: bytes, *, arabic: bool = True) -> tuple[str, int | None]:
    """OCR intégré PyMuPDF (souvent meilleur que la couche texte BIDI cassée)."""
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        parts: list[str] = []
        lang = "ara" if arabic else "fra+ara"
        dpi = _ocr_dpi()
        n_pages = min(len(doc), _max_pdf_ocr_pages())
        for i in range(n_pages):
            page = doc[i]
            try:
                tp = page.get_textpage_ocr(language=lang, dpi=dpi, full=True)
                t = page.get_text(textpage=tp) or ""
            except Exception:
                t = page.get_text() or ""
            if t.strip():
                parts.append(t.strip())
        doc.close()
        if not parts:
            return "", 0
        combined = safe_repair_arabic_text("\n\n".join(parts))
        score = min(92, max(50, 40 + len(combined) // 40))
        return combined, score
    except Exception as exc:
        logger.warning("PyMuPDF OCR failed: %s", exc)
        return "", 0


def _ocr_pdf_pymupdf_timed(content: bytes, *, arabic: bool = True) -> tuple[str, int | None]:
    from app.core.config import settings
    from app.core.timeouts import run_with_timeout

    try:
        return run_with_timeout(
            lambda: _ocr_pdf_pymupdf(content, arabic=arabic),
            float(settings.OCR_PYMUPDF_TIMEOUT_SEC or 90),
            label="ocr_pdf_pymupdf",
        )
    except Exception as exc:
        from app.core.timeouts import OperationTimeout

        if isinstance(exc, OperationTimeout):
            raise
        logger.warning("PyMuPDF OCR timed wrapper failed: %s", exc)
        return "", 0


def _ocr_pdf_scanned(content: bytes) -> tuple[str, int | None]:
    from app.core.config import settings
    from app.core.timeouts import run_with_timeout

    max_pages = _max_tesseract_pages()
    dpi = _ocr_dpi()

    def _run() -> tuple[str, int | None]:
        from pdf2image import convert_from_bytes

        images = convert_from_bytes(
            content, dpi=dpi, first_page=1, last_page=max_pages,
        )
        parts: list[str] = []
        scores: list[int] = []
        for img in images:
            page_text, page_score = _ocr_image(img, fast=False)
            if page_text:
                parts.append(page_text)
                scores.append(page_score)
        if not parts:
            return "", 0
        combined = safe_repair_arabic_text("\n\n".join(parts))
        avg = int(sum(scores) / len(scores)) if scores else 0
        return combined, avg

    try:
        return run_with_timeout(
            _run,
            float(settings.OCR_TESSERACT_TIMEOUT_SEC or 120),
            label="ocr_pdf_scanned",
        )
    except Exception as exc:
        from app.core.timeouts import OperationTimeout

        if isinstance(exc, OperationTimeout):
            raise
        logger.warning("OCR PDF scanned failed: %s", exc)
        return "", 0


def _pick_best_text(*candidates: str) -> str:
    scored = []
    for t in candidates:
        if not t or not t.strip():
            continue
        # Pénaliser fort le texte brouillé
        pen = 40 if is_scrambled_arabic(t) else 0
        scored.append((_arabic_quality_score(t) + min(len(t), 5000) // 100 - pen, t))
    if not scored:
        return ""
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[0][1]


def extract_from_pdf(content: bytes) -> tuple[str, int | None]:
    doc = fitz.open(stream=content, filetype="pdf")
    page_count = len(doc)
    naive_parts: list[str] = []
    ordered_parts: list[str] = []
    for page in doc:
        try:
            naive_parts.append(page.get_text("text", sort=True) or page.get_text() or "")
        except TypeError:
            naive_parts.append(page.get_text() or "")
        ordered_parts.append(_page_text_reading_order(page))
    doc.close()

    naive = "\n".join(naive_parts).strip()
    ordered = "\n".join(ordered_parts).strip()
    naive_fixed = safe_repair_arabic_text(naive)
    ordered_fixed = safe_repair_arabic_text(ordered)
    text = _pick_best_text(ordered_fixed, naive_fixed, ordered, naive)
    ar_heavy = arabic_char_ratio(text) > 0.25 or arabic_char_ratio(naive) > 0.25

    # PDF numérique avec couche texte propre → extraction native (qualité max)
    if looks_like_native_arabic_prose(text) and len(text) >= 80:
        return text, 100
    if text and len(text) > 80 and not is_scrambled_arabic(text) and not ar_heavy:
        return text, 100

    thin_scan = page_count >= 2 and len(text) < page_count * 40
    must_ocr = (
        (not text or len(text) < 50)
        or (ar_heavy and is_scrambled_arabic(text))
        or thin_scan
    )

    if must_ocr:
        from app.services.cloud_ocr import ocr_document_cloud

        logger.info("PDF scan/arabe → OCR cloud (qualité max)")
        cloud_text, cloud_score = ocr_document_cloud(
            content, filename="document.pdf", mime="application/pdf"
        )
        # Faire confiance à Mistral/Gemini : ne pas relancer Tesseract (8+ min)
        if cloud_text and len(cloud_text.strip()) >= 80:
            return cloud_text, cloud_score or 96

        logger.info("OCR cloud indisponible → PyMuPDF puis Tesseract haute qualité")
        ocr_a, score_a = "", 0
        try:
            ocr_a, score_a = _ocr_pdf_pymupdf_timed(content, arabic=ar_heavy)
        except Exception as exc:
            from app.core.timeouts import OperationTimeout

            if isinstance(exc, OperationTimeout):
                logger.warning("PyMuPDF OCR timeout — Tesseract haute qualité")
            else:
                logger.warning("PyMuPDF OCR error: %s", exc)

        ocr_b, score_b = "", 0
        try:
            ocr_b, score_b = _ocr_pdf_scanned(content)
        except Exception as exc:
            from app.core.timeouts import OperationTimeout

            if isinstance(exc, OperationTimeout):
                logger.warning("Tesseract OCR timeout — meilleur texte obtenu")
            else:
                logger.warning("Tesseract OCR error: %s", exc)

        best = _pick_best_text(cloud_text, ocr_a, ocr_b, text)
        # Ne jamais préférer une couche PDF arabe inversée si un OCR existe
        if ar_heavy and text and is_scrambled_arabic(text):
            best = _pick_best_text(cloud_text, ocr_a, ocr_b) or best
        if best and _ocr_quality_ok(best, min_len=40):
            sc = cloud_score if best == cloud_text else (
                score_a if best == ocr_a else (score_b if best == ocr_b else 70)
            )
            if best == text and ar_heavy and is_scrambled_arabic(text):
                sc = min(sc or 55, 45)
            return best, sc or 75
        if best and len(best) > len(text or ""):
            return best, max(cloud_score or 0, score_a or 0, score_b or 0, 55)
        # Couche native arabe brouillée sans OCR utilisable → score bas (pas 100)
        if text and len(text) > 50:
            if ar_heavy and is_scrambled_arabic(text):
                return text, 40
            return text, 55
        if best:
            return best, max(cloud_score or 0, score_a or 0, score_b or 0, 40)
        return best or text, 40

    if len(text) > 50:
        return text, 100

    from app.services.cloud_ocr import ocr_document_cloud

    cloud_text, cloud_score = ocr_document_cloud(
        content, filename="document.pdf", mime="application/pdf"
    )
    if cloud_text and len(cloud_text.strip()) >= 80:
        return cloud_text, cloud_score or 96

    ocr_text, ocr_score = _ocr_pdf_scanned(content)
    if len(ocr_text) > 20:
        return ocr_text, ocr_score

    return text, (85 if text else 0)


def _docx_xml_text(el: ET.Element) -> str:
    bits: list[str] = []
    for node in el.iter(f"{_W_NS}t"):
        if node.text:
            bits.append(node.text)
        if node.tail:
            bits.append(node.tail)
    return "".join(bits).strip()


def _extract_docx_footnotes(content: bytes) -> list[str]:
    """Lit word/footnotes.xml (python-docx ignore souvent les notes)."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as zf:
            if "word/footnotes.xml" not in zf.namelist():
                return []
            root = ET.fromstring(zf.read("word/footnotes.xml"))
    except Exception:
        return []
    notes: list[str] = []
    for fn in root.findall(f"{_W_NS}footnote"):
        fid = fn.get(f"{_W_NS}id") or fn.get("id") or ""
        if fid in {"-1", "0"}:  # séparateurs Word
            continue
        text = _docx_xml_text(fn)
        if text and len(text) > 1:
            notes.append(f"[{fid}] {text}" if fid else text)
    return notes


def extract_from_docx(content: bytes) -> tuple[str, int | None]:
    """
    Extraction native Word (pas d'OCR).
    Paragraphes + tableaux + notes de bas de page.
    Pas d'inversion BIDI (le DOCX numérique est déjà dans le bon ordre).
    """
    from app.services.arabic_text_fix import repair_arabic_stt

    # Garde anti zip-bomb (DOCX = ZIP)
    try:
        with zipfile.ZipFile(BytesIO(content)) as zf:
            infos = zf.infolist()
            if len(infos) > 4000:
                raise ValueError("DOCX trop complexe (trop d’entrées ZIP).")
            uncompressed = sum(max(0, int(i.file_size or 0)) for i in infos)
            if uncompressed > 180 * 1024 * 1024:
                raise ValueError("DOCX décompressé trop volumineux.")
            compressed = max(1, len(content))
            if uncompressed / compressed > 200 and uncompressed > 20 * 1024 * 1024:
                raise ValueError("DOCX compression anormale (zip-bomb).")
    except zipfile.BadZipFile as exc:
        raise ValueError("Fichier Word invalide (ZIP).") from exc

    doc = DocxDocument(BytesIO(content))
    parts: list[str] = []

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                ct = (cell.text or "").strip()
                if ct and ct not in cells:
                    cells.append(ct)
            if cells:
                parts.append(" | ".join(cells))

    footnotes = _extract_docx_footnotes(content)
    if footnotes:
        parts.append("")  # séparateur
        parts.append("— Notes —" if arabic_char_ratio("\n".join(parts)) < 0.2 else "— الهوامش —")
        parts.extend(footnotes)

    text = "\n".join(parts).strip()
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    # Lexique ة / nettoyage STT-like uniquement — JAMAIS reverse de mots
    if text and arabic_char_ratio(text) >= 0.12:
        text = repair_arabic_stt(text)

    return text, None


def is_native_office_document(filename: str, mime_type: str = "") -> bool:
    name = (filename or "").lower()
    mime = (mime_type or "").lower()
    return name.endswith(".docx") or "wordprocessingml" in mime


def extract_from_image(content: bytes) -> tuple[str, int | None]:
    from app.services.cloud_ocr import ocr_document_cloud

    mime = "image/png"
    try:
        img = Image.open(BytesIO(content))
        fmt = (img.format or "PNG").upper()
        mime = {
            "JPEG": "image/jpeg",
            "JPG": "image/jpeg",
            "PNG": "image/png",
            "TIFF": "image/tiff",
            "TIF": "image/tiff",
        }.get(fmt, "image/png")
    except Exception:
        img = None

    cloud_text, cloud_score = ocr_document_cloud(
        content, filename="scan.png", mime=mime
    )
    if cloud_text and len(cloud_text.strip()) >= 40:
        return cloud_text, cloud_score or 96

    try:
        image = img if img is not None else Image.open(BytesIO(content))
        return _ocr_image_timed(image, fast=False)
    except Exception:
        return "", 35


def extract_from_audio(content: bytes, filename: str) -> tuple[str, int | None]:
    """
    Transcription Whisper (même moteur que la dictée).
    Retourne (texte, ocr_score=None) — pas de score OCR pour l'audio.
    Langue : arabe d’abord (cabinet marocain), puis français si vide.
    Post-traitement : réparation STT + polish fusḥā (comme la dictée).
    """
    from app.core.config import settings
    from app.core.timeouts import OperationTimeout
    from app.services.arabic_text_fix import repair_arabic_stt
    from app.services.whisper_stt import transcribe_audio_bytes

    if not settings.WHISPER_ENABLED:
        return "", None

    name = filename or "audio.webm"
    text, provider = transcribe_audio_bytes(content, filename=name, language="ar")
    used_lang = "ar"
    if not (text or "").strip():
        # Secours FR si l’audio est clairement français
        text, provider = transcribe_audio_bytes(content, filename=name, language="fr")
        used_lang = "fr"

    if provider == "whisper:timeout":
        raise OperationTimeout(
            f"Transcription dépassée ({settings.WHISPER_TIMEOUT_SEC:.0f}s)."
        )

    text = (text or "").strip()
    if not text:
        return "", None

    # Lexique / artefacts STT
    if used_lang == "ar" or arabic_char_ratio(text) >= 0.25:
        text = repair_arabic_stt(text)
        # Polish conservateur (interdit d'inventer villes / noms à partir du phonétique)
        try:
            from app.services.dictation import polish_dictation

            sample = text if len(text) <= 6000 else text[:6000]
            polished, prov = polish_dictation(sample, language="ar", source="audio")
            if polished and len(polished.strip()) >= max(30, int(len(sample) * 0.35)):
                if len(text) <= 6000:
                    text = polished.strip()
                else:
                    text = (polished.strip() + "\n\n" + text[6000:]).strip()
                logger.info("Audio AR polished via %s (%d chars)", prov, len(text))
        except Exception as exc:
            logger.warning("Audio AR polish skipped: %s", exc)

    return text, None


def extract_text(content: bytes, filename: str, mime_type: str) -> tuple[str, int | None]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    dotted = f".{ext}" if ext else get_extension(filename)

    if is_audio_document(filename, mime_type) or dotted in AUDIO_EXTENSIONS:
        return extract_from_audio(content, filename)

    if ext == "pdf" or mime_type == "application/pdf":
        return extract_from_pdf(content)
    if ext == "docx" or "wordprocessingml" in mime_type:
        return extract_from_docx(content)
    if ext in {"txt", "md", "markdown"} or mime_type in {
        "text/plain",
        "text/txt",
        "text/markdown",
        "text/x-markdown",
    }:
        for enc in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
            try:
                return content.decode(enc).strip(), 100
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace").strip(), 80
    if ext in ("png", "jpg", "jpeg", "tiff", "tif") or mime_type.startswith("image/"):
        return extract_from_image(content)
    return "", None


def extract_parties(text: str) -> list[str]:
    patterns = [
        r"(?:entre|parties?)\s*:?\s*(.+?)(?:\n|$)",
        r"(?:locataire|bailleur|vendeur|acheteur|employeur|salarié)\s*:?\s*([^\n,;]+)",
    ]
    found = []
    noise = re.compile(
        r"^(prise|entre|et le stagiaire|cette obligation|d['’]information)\b|"
        r"obligation demeure|expiration du stage",
        re.IGNORECASE,
    )
    for pattern in patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            val = re.sub(r"\s+", " ", match.group(1)).strip(" .;,-")
            if len(val) < 4 or len(val) > 70:
                continue
            if noise.search(val):
                continue
            if val.upper().rstrip(".") in {"PRISE", "ENTRE", "ET"}:
                continue
            if val not in found:
                found.append(val)
    return found[:5]


def extract_dates(text: str) -> list[str]:
    """Dates calendaires — exclut les numéros de lois (1.74.447, 39.08, etc.)."""
    raw = re.findall(
        r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b|\b\d{4}[/.-]\d{1,2}[/.-]\d{1,2}\b",
        text,
    )
    out: list[str] = []
    for d in raw:
        # Exclure motifs type numéro de dahir / loi (1.74.447, 14.07)
        if re.fullmatch(r"\d\.\d{2}\.\d{3}", d.replace("/", ".").replace("-", ".")):
            continue
        # Exclure années isolées collées en faux format
        parts = re.split(r"[/.-]", d)
        if len(parts) == 3:
            try:
                a, b, c = (int(x) for x in parts)
            except ValueError:
                continue
            # jour/mois/année ou année/mois/jour plausibles
            if max(a, b) > 31 and c < 1000:
                continue
            year = c if c > 31 else (a if a > 31 else None)
            if year is not None and (year < 1900 or year > 2100) and year > 99:
                continue
            if year is not None and year < 100:
                year = 1900 + year if year >= 50 else 2000 + year
            if year is not None and not (1950 <= year <= 2035):
                continue
        if d not in out:
            out.append(d)
    return out[:10]


def extract_amounts(text: str) -> list[str]:
    amounts = re.findall(
        r"\b\d{1,3}(?:[ \u00a0]\d{3})*(?:[.,]\d+)?\s*(?:MAD|DH|Dhs|dirhams?|€|EUR|USD|\$)\b",
        text,
        flags=re.IGNORECASE,
    )
    return list(dict.fromkeys(amounts))[:10]


def _looks_like_broken_arabic_snippet(snippet: str) -> bool:
    """Rejette les bouts BIDI inversés / mots coupés (بنود المخاطر pourries)."""
    s = (snippet or "").strip()
    if len(s) < 25:
        return True
    if "$^{" in s or "\\" in s or s.startswith("#"):
        return True
    # Questions / titres de plan ≠ clauses à risque
    if "؟" in s or s.endswith("?") or "المبحث" in s or "الفقرة" in s:
        return True
    ar = sum(1 for c in s if "\u0600" <= c <= "\u06FF")
    if ar / max(1, len(s)) < 0.35:
        return False
    words = s.split()
    if not words:
        return True
    tiny = sum(1 for w in words if len(re.sub(r"[^\u0600-\u06FF]", "", w)) <= 1)
    if tiny / len(words) >= 0.15:
        return True
    # Mot coupé en fin (المشتر، مبا، ق…)
    last_ar = re.sub(r"[^\u0600-\u06FF]", "", words[-1])
    if 1 <= len(last_ar) <= 3 and last_ar not in {"في", "من", "إلى", "على", "عن", "أو", "لا", "ما", "قد", "كل"}:
        # fin tronquée fréquente sur fragments PDF
        if not s.endswith((".", "۔", ":", "؛")):
            return True
    if len(words[0]) <= 2 and words[0][:1] in "ةهيوكفبلمن":
        return True
    orig = _line_reading_score(words)
    rev = _line_reading_score(list(reversed(words)))
    if rev > orig + 3:
        return True
    # Indices forts d'arabe inversé (finale de phrase au début)
    if re.search(r"^(ومنعاً|صيانةً|برمتها|به الإخلال|على ويترتب)", s):
        return True
    if "به الإخلال على ويترتب" in s or "بطلان به الإخلال" in s:
        return True
    return False


def extract_risk_clauses(text: str) -> list[str]:
    keywords = (
        r"résiliation|résilier|pénalité|amende|responsabilité|garantie|caution|"
        r"indemnité|clause\s+pénale|force\s+majeure|exclusivité|"
        r"فسخ|غرامة|مسؤولية|ضمان|تعويض|بطلان|حجز|مزاد"
    )
    found = []
    for m in re.finditer(rf"([^\n.!?…]{{0,80}}(?:{keywords})[^\n.!?…]{{0,80}})", text, re.IGNORECASE):
        snippet = re.sub(r"\s+", " ", m.group(1)).strip(" ،,.-")
        if _looks_like_broken_arabic_snippet(snippet):
            continue
        if 25 < len(snippet) < 180 and snippet not in found:
            found.append(snippet)
    return found[:8]


def get_document_context_excerpt(text: str, question: str, max_len: int = 8000) -> str:
    """Extrait pertinent autour des mots de la question (sinon début du doc)."""
    if not text:
        return ""
    text = repair_arabic_extracted_text(text)
    if len(text) <= max_len:
        return text
    q_tokens = [t for t in re.findall(r"[\w\u0600-\u06FF]{3,}", (question or "").lower())]
    if not q_tokens:
        return text[:max_len]
    best_i = 0
    best_hits = -1
    step = max(200, max_len // 4)
    for i in range(0, max(1, len(text) - max_len), step):
        chunk = text[i : i + max_len].lower()
        hits = sum(1 for t in q_tokens if t in chunk)
        if hits > best_hits:
            best_hits = hits
            best_i = i
    return text[best_i : best_i + max_len]


def find_relevant_excerpt(text: str, question: str, max_len: int = 1200) -> str:
    return get_document_context_excerpt(text, question, max_len=max_len)
